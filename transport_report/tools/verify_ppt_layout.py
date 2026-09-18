"""Compare six-slide report packages. Run: python -m tools.verify_ppt_layout A B."""
from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.reporting.pptx_report import load_template, TemplateStructureChanged


def _shape_signature(shape):
    signature = [shape.name, shape.shape_id, shape.shape_type, shape.left, shape.top, shape.width, shape.height, shape.rotation]
    if shape.has_table:
        table = shape.table
        signature += [tuple(row.height for row in table.rows), tuple(column.width for column in table.columns), tuple((cell.is_spanned, cell.span_width, cell.span_height, cell.margin_left, cell.margin_right, cell.margin_top, cell.margin_bottom, cell._tc.tcPr.xml) for row in table.rows for cell in row.cells)]
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        signature += [shape.crop_left, shape.crop_right, shape.crop_top, shape.crop_bottom]
        if shape.name != 'report.chart':
            signature.append(shape.image.blob)
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        signature.append(tuple(_shape_signature(child) for child in shape.shapes))
    xml = deepcopy(shape._element)
    for body in xml.xpath('.//a:txBody | .//p:txBody'):
        body.getparent().remove(body)
    if shape.name == 'report.chart':
        for blip in xml.xpath('.//a:blip'):
            blip.attrib.pop('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed', None)
    signature.append(xml.xml)
    return tuple(signature)


def verify_ppt_layout(template: str | Path, output: str | Path) -> list[str]:
    """Return differences. Text and the named chart images may change only.

    Includes package CRC, preserved masters/layouts/themes, all shape order and
    geometry, table grid/merge/border/fill/margins, and picture crops/counts.
    Run/paragraph formatting is additionally exercised by integration tests.
    """
    try:
        original, generated = load_template(template), load_template(output)
    except TemplateStructureChanged as error:
        return [str(error)]
    errors = []
    for number, (before, after) in enumerate(zip(original.slides, generated.slides), 1):
        if len(before.shapes) != len(after.shapes):
            errors.append(f'slide {number}: shape count changed')
        if sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in before.shapes) != sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in after.shapes):
            errors.append(f'slide {number}: picture count changed')
        for old, new in zip(before.shapes, after.shapes):
            if _shape_signature(old) != _shape_signature(new):
                errors.append(f'slide {number}: layout changed for {old.name}')
    with ZipFile(template) as before, ZipFile(output) as after:
        if after.testzip() is not None:
            errors.append('output ZIP CRC check failed')
        for name in before.namelist():
            if name.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/theme/')):
                if name not in after.namelist() or before.read(name) != after.read(name):
                    errors.append(f'package template part changed: {name}')
    return errors


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('template', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    errors = verify_ppt_layout(args.template, args.output)
    for error in errors:
        print(error)
    if not errors:
        print('Six-slide template layout verified.')
    return bool(errors)


if __name__ == '__main__':
    raise SystemExit(main())
