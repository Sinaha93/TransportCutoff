from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath
from posixpath import join as posix_join
from posixpath import normpath as posix_normpath
from xml.etree.ElementTree import ParseError, iterparse
from zipfile import BadZipFile, LargeZipFile, ZipFile

from openpyxl import load_workbook


REGULAR_SHEET = "화성운반비내역(8월)"
_REGULAR_SHEET = re.compile(r"^화성운반비내역\((?P<month>[1-9]|1[0-2])월\)$")
SUBCONTRACT_SHEETS = (
    "용차1-OK로지웰",
    "용차2-대원로지스틱",
    "용차3-정동물류",
)
_MONTH = re.compile(r"^[0-9]{4}-(0[1-9]|1[0-2])$")
_REGULAR_COLUMN_B_HEADERS = {"날짜", "운반지역", "운반지", "목적지", "납품처"}
_REGULAR_TOTAL_COUNT_HEADERS = {"계", "총회수", "총횟수"}
_REGULAR_UNIT_HEADERS = {"단가", "운반단가", "운임단가"}
_REGULAR_SUBTOTAL_HEADERS = {"소계", "금액", "운반비"}
_DATE_HEADERS = {"일자", "날짜", "오더일자", "운반일", "운송일"}
_DESTINATION_HEADERS = {"운반지역", "운반지", "목적지", "납품처", "도착지"}
_TONNAGE_HEADERS = {"톤수", "차량톤수", "차종", "차종(t)", "차종(톤)"}
_AMOUNT_HEADERS = {"금액", "운반비", "운임", "기본요금"}
_NUMBER_HEADERS = {"no", "번호"}
_TOTAL_LABELS = {"합계", "총계", "계"}
_SUBCONTRACT_TOTAL_LABEL_COLUMNS = (1, 7)
_SUBCONTRACT_LAST_FOOTPRINT_COLUMN = 10
_SQLITE_MAX = 2**63 - 1
# Layout bounds are intentionally well above the real workbook's 116 rows/54 columns.
MAX_WORKSHEET_ROWS = 2_000
MAX_WORKSHEET_COLUMNS = 128
MAX_PARSED_ENTRIES = 20_000
# Pre-load limits cover every worksheet, not only the four imported sheets. The
# real source has 6 sheets, at most 3,970 cell records/sheet, 7,535 total cells,
# 116 rows, 54 columns, and a largest worksheet XML of 101,774 bytes.
# Expanded merged-range areas are charged against the same cell envelopes because
# OpenPyXL creates a MergedCell object for each covered coordinate.
MAX_PREFLIGHT_WORKSHEETS = 256
MAX_PREFLIGHT_WORKSHEET_XML_BYTES = 8 * 1024 * 1024
MAX_PREFLIGHT_RELATED_XML_BYTES = 8 * 1024 * 1024
MAX_PREFLIGHT_WORKSHEET_CELLS = 50_000
MAX_PREFLIGHT_TOTAL_CELLS = 100_000
MAX_PREFLIGHT_ROWS = 20_000
MAX_PREFLIGHT_COLUMNS = 512
# OpenPyXL creates one ColumnDimension per raw <col> record; it does not expand
# min..max spans. The real source has up to 704 records and valid style spans
# through Excel's final column, so bound records separately from data columns.
MAX_PREFLIGHT_COLUMN_DIMENSION_RECORDS = 2_048
MAX_EXCEL_COLUMNS = 16_384
# Conservative source-cell limits: 10,000 trips/day, 4 decimals, 16 canonical chars.
MAX_TRIP_COUNT = Decimal("10000")
MAX_TRIP_COUNT_SCALE = 4
MAX_TRIP_COUNT_CANONICAL_LENGTH = 16
_CELL_REFERENCE = re.compile(r"^(?P<column>[A-Za-z]+)(?P<row>[1-9][0-9]*)$")
_DOCUMENT_RELATIONSHIP_ID = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
)
_COMMENTS_RELATIONSHIP = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
)
_OPENPYXL_EAGER_SHEET_RELATIONSHIPS = {
    _COMMENTS_RELATIONSHIP,
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/table",
}
_HYPERLINK_RELATIONSHIP = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
)


class TransportWorkbookError(Exception):
    """Base error for workbooks that cannot be imported safely."""


class WorkbookStructureError(TransportWorkbookError):
    """Raised when a required workbook structure or value is missing."""


class SubtotalMismatchError(TransportWorkbookError):
    """Raised when a regular row does not reconcile to its cached subtotal."""


class TotalCountMismatchError(TransportWorkbookError):
    """Raised when a regular row does not reconcile to its cached trip count."""


class SubcontractTotalMismatchError(TransportWorkbookError):
    """Raised when subcontract details do not reconcile to the displayed total."""


@dataclass(frozen=True, slots=True)
class ParsedTransportEntry:
    report_month: str
    destination_alias: str
    source_sheet: str
    source_row: int
    source_date: date
    day: int
    transport_type: str
    trip_count: Decimal
    unit_rate_won: int | None
    cost_won: int
    vehicle_type: str | None = None
    vehicle_driver_group: str | None = None
    source_note: str | None = None


class HwaseongWorkbookParser:
    def parse(
        self, workbook_path: str | Path, report_month: str
    ) -> list[ParsedTransportEntry]:
        if not isinstance(report_month, str) or _MONTH.fullmatch(report_month) is None:
            raise WorkbookStructureError("report_month must use YYYY-MM")

        path = Path(workbook_path)
        _preflight_workbook_resources(path)
        workbook = None
        formula_workbook = None
        try:
            workbook = load_workbook(
                path,
                data_only=True,
                read_only=False,
                keep_vba=False,
                keep_links=False,
            )
            formula_workbook = load_workbook(
                path,
                data_only=False,
                read_only=False,
                keep_vba=False,
                keep_links=False,
            )
        except Exception as error:
            if workbook is not None:
                workbook.close()
            raise WorkbookStructureError(f"Workbook could not be opened: {error}") from error

        try:
            regular_sheets = [
                name for name in workbook.sheetnames if _REGULAR_SHEET.fullmatch(name)
            ]
            missing = [name for name in SUBCONTRACT_SHEETS if name not in workbook.sheetnames]
            if not regular_sheets:
                missing.insert(0, "화성운반비내역(<month>월)")
            if missing:
                raise WorkbookStructureError(
                    "Workbook is missing required sheet(s): " + ", ".join(missing)
                )
            if len(regular_sheets) != 1:
                raise WorkbookStructureError(
                    "Workbook must contain exactly one regular transport sheet"
                )

            regular_sheet = workbook[regular_sheets[0]]
            _validate_regular_sheet_month(regular_sheet.title, report_month)
            for sheet_name in (regular_sheets[0], *SUBCONTRACT_SHEETS):
                _validate_worksheet_bounds(workbook[sheet_name])
            rows = self._parse_regular(
                regular_sheet,
                formula_workbook[regular_sheets[0]],
                report_month,
                MAX_PARSED_ENTRIES,
            )
            for sheet_name in SUBCONTRACT_SHEETS:
                rows.extend(
                    self._parse_subcontract(
                        workbook[sheet_name],
                        formula_workbook[sheet_name],
                        report_month,
                        MAX_PARSED_ENTRIES - len(rows),
                    )
                )
            return rows
        finally:
            workbook.close()
            formula_workbook.close()

    def _parse_regular(
        self, sheet, formula_sheet, report_month: str, entry_limit: int
    ) -> list[ParsedTransportEntry]:
        header_row = _find_regular_header(sheet, formula_sheet)
        if header_row is None:
            raise WorkbookStructureError(
                f"Required header was not found in sheet {sheet.title}"
            )

        rows: list[ParsedTransportEntry] = []
        vehicle_driver_group: str | None = None
        for row_number in range(header_row + 1, sheet.max_row + 1):
            if _is_regular_header(sheet, formula_sheet, row_number):
                continue

            group_value = _optional_text(sheet.cell(row_number, 1).value)
            if group_value is not None:
                vehicle_driver_group = group_value

            destination_alias = _destination_alias(
                sheet.cell(row_number, 2).value,
                f"{sheet.title} row {row_number} destination",
            )
            day_values = [
                sheet.cell(row_number, column).value for column in range(3, 34)
            ]
            day_formula_values = [
                formula_sheet.cell(row_number, column).value
                for column in range(3, 34)
            ]
            cached_count_value = sheet.cell(row_number, 34).value
            unit_value = sheet.cell(row_number, 36).value
            cached_subtotal_value = sheet.cell(row_number, 37).value
            for column, (cached_value, formula_value) in enumerate(
                zip(day_values, day_formula_values, strict=True), start=3
            ):
                if not _has_value(cached_value) and _is_formula(formula_value):
                    _required_cached_value(
                        cached_value,
                        formula_value,
                        sheet.title,
                        row_number,
                        formula_sheet.cell(row_number, column).coordinate,
                        "daily trip count",
                    )
            unit_formula_value = formula_sheet.cell(row_number, 36).value
            if not _has_value(unit_value) and _is_formula(unit_formula_value):
                _required_cached_value(
                    unit_value,
                    unit_formula_value,
                    sheet.title,
                    row_number,
                    f"AJ{row_number}",
                    "unit cost",
                )
            count_formula_value = formula_sheet.cell(row_number, 34).value
            if not _has_value(cached_count_value) and _is_formula(
                count_formula_value
            ):
                _required_cached_value(
                    cached_count_value,
                    count_formula_value,
                    sheet.title,
                    row_number,
                    f"AH{row_number}",
                    "cached total count",
                )
            subtotal_formula_value = formula_sheet.cell(row_number, 37).value
            if not _has_value(cached_subtotal_value) and _is_formula(
                subtotal_formula_value
            ):
                _required_cached_value(
                    cached_subtotal_value,
                    subtotal_formula_value,
                    sheet.title,
                    row_number,
                    f"AK{row_number}",
                    "cached subtotal",
                )
            if _has_regular_footer_label(sheet, formula_sheet, row_number):
                cached_footer_marker = sheet.cell(row_number, 38).value
                formula_footer_marker = formula_sheet.cell(row_number, 38).value
                if not _has_value(cached_footer_marker) and _is_formula(
                    formula_footer_marker
                ):
                    _required_cached_value(
                        cached_footer_marker,
                        formula_footer_marker,
                        sheet.title,
                        row_number,
                        f"AL{row_number}",
                        "footer marker",
                    )
            if _is_regular_footer_or_note(sheet, formula_sheet, row_number):
                continue
            if _is_regular_blank_or_group_row(sheet, formula_sheet, row_number):
                continue
            if destination_alias is None:
                raise WorkbookStructureError(
                    f"{sheet.title} row {row_number} destination is required"
                )
            if _is_regular_inactive_row(sheet, formula_sheet, row_number):
                continue
            unit_value = _required_cached_value(
                unit_value,
                unit_formula_value,
                sheet.title,
                row_number,
                f"AJ{row_number}",
                "unit cost",
            )
            if isinstance(unit_value, bool) or not isinstance(
                unit_value, (int, float, Decimal)
            ):
                raise WorkbookStructureError(
                    f"{sheet.title} row {row_number} unit cost must be numeric"
                )

            unit_rate = _integer_won(
                unit_value,
                f"{sheet.title} row {row_number} unit cost",
            )
            row_entries: list[ParsedTransportEntry] = []
            calculated_count = Decimal(0)
            for day, raw_count in enumerate(day_values, start=1):
                if raw_count in (None, ""):
                    continue
                trip_count = _trip_count_decimal(
                    raw_count,
                    f"{sheet.title} row {row_number} day {day} trip count",
                )
                if trip_count == 0:
                    continue
                source_date = _validate_calendar_day(
                    report_month, day, sheet.title, row_number
                )
                calculated_count += trip_count
                cost_won = _integer_won(
                    trip_count * unit_rate,
                    f"{sheet.title} row {row_number} day {day} calculated cost",
                )
                _append_parsed_entry(
                    row_entries,
                    ParsedTransportEntry(
                        report_month=report_month,
                        destination_alias=destination_alias,
                        source_sheet=sheet.title,
                        source_row=row_number,
                        source_date=source_date,
                        day=day,
                        transport_type="regular",
                        trip_count=trip_count,
                        unit_rate_won=unit_rate,
                        cost_won=cost_won,
                        vehicle_driver_group=vehicle_driver_group,
                    ),
                    entry_limit - len(rows),
                )

            cached_count = _nonnegative_decimal(
                _required_cached_value(
                    cached_count_value,
                    count_formula_value,
                    sheet.title,
                    row_number,
                    f"AH{row_number}",
                    "cached total count",
                ),
                f"{sheet.title} row {row_number} cached total count",
            )
            if calculated_count != cached_count:
                raise TotalCountMismatchError(
                    f"Total count mismatch in {sheet.title} row {row_number}: "
                    f"calculated {calculated_count}, cached {cached_count}"
                )
            cached_subtotal = _integer_won(
                _required_cached_value(
                    cached_subtotal_value,
                    subtotal_formula_value,
                    sheet.title,
                    row_number,
                    f"AK{row_number}",
                    "cached subtotal",
                ),
                f"{sheet.title} row {row_number} cached subtotal",
            )
            calculated_subtotal = sum(entry.cost_won for entry in row_entries)
            if calculated_subtotal != cached_subtotal:
                raise SubtotalMismatchError(
                    f"Subtotal mismatch in {sheet.title} row {row_number}: "
                    f"calculated {calculated_subtotal}, cached {cached_subtotal}"
                )
            rows.extend(row_entries)
        return rows

    def _parse_subcontract(
        self, sheet, formula_sheet, report_month: str, entry_limit: int
    ) -> list[ParsedTransportEntry]:
        header = _find_subcontract_header(sheet, formula_sheet)
        if header is None:
            raise WorkbookStructureError(
                f"Required header was not found in sheet {sheet.title}"
            )
        (
            header_row,
            date_column,
            destination_column,
            tonnage_column,
            number_column,
            last_footprint_column,
        ) = header

        rows: list[ParsedTransportEntry] = []
        calculated_total = 0
        displayed_total_found = False
        for row_number in range(header_row + 1, sheet.max_row + 1):
            if _is_displayed_total_row(sheet, formula_sheet, row_number):
                if displayed_total_found:
                    raise WorkbookStructureError(
                        f"{sheet.title} row {row_number} contains a second displayed total"
                    )
                cached_total = _integer_won(
                    _required_cached_value(
                        sheet.cell(row_number, 9).value,
                        formula_sheet.cell(row_number, 9).value,
                        sheet.title,
                        row_number,
                        f"I{row_number}",
                        "displayed cached total",
                    ),
                    f"{sheet.title} row {row_number} cached total",
                )
                if cached_total != calculated_total:
                    raise SubcontractTotalMismatchError(
                        f"Total mismatch in {sheet.title} row {row_number}: "
                        f"calculated {calculated_total}, cached {cached_total}"
                )
                displayed_total_found = True
                continue
            footprint_values = [
                sheet.cell(row_number, column).value
                for column in range(1, last_footprint_column + 1)
            ]
            formula_footprint_values = [
                formula_sheet.cell(row_number, column).value
                for column in range(1, last_footprint_column + 1)
            ]
            if not any(
                _has_value(value)
                for value in (*footprint_values, *formula_footprint_values)
            ):
                continue
            if _is_subcontract_template_row(
                sheet,
                formula_sheet,
                row_number,
                number_column,
                last_footprint_column,
            ):
                continue
            if displayed_total_found:
                if _is_subcontract_formula_footer(
                    sheet, formula_sheet, row_number, last_footprint_column
                ):
                    continue
                raise WorkbookStructureError(
                    f"{sheet.title} row {row_number} contains detail after "
                    "the displayed total"
                )
            raw_date = sheet.cell(row_number, date_column).value
            transport_date = _transport_date(
                raw_date, f"{sheet.title} row {row_number} date"
            )
            destination_alias = _destination_alias(
                sheet.cell(row_number, destination_column).value,
                f"{sheet.title} row {row_number} destination",
            )
            raw_tonnage = sheet.cell(row_number, tonnage_column).value
            raw_amount = sheet.cell(row_number, 9).value
            if destination_alias is None:
                raise WorkbookStructureError(
                    f"{sheet.title} row {row_number} destination is required"
                )
            if not _is_in_subcontract_billing_window(transport_date, report_month):
                raise WorkbookStructureError(
                    f"{sheet.title} row {row_number} date is outside the "
                    f"{report_month} billing window"
                )
            vehicle_type = _optional_text(raw_tonnage)
            if vehicle_type is None:
                raise WorkbookStructureError(
                    f"{sheet.title} row {row_number} vehicle is required"
                )
            amount = _integer_won(
                _required_cached_value(
                    raw_amount,
                    formula_sheet.cell(row_number, 9).value,
                    sheet.title,
                    row_number,
                    f"I{row_number}",
                    "amount",
                ),
                f"{sheet.title} row {row_number} amount",
            )
            calculated_total += amount
            _append_parsed_entry(
                rows,
                ParsedTransportEntry(
                    report_month=report_month,
                    destination_alias=destination_alias,
                    source_sheet=sheet.title,
                    source_row=row_number,
                    source_date=transport_date,
                    day=transport_date.day,
                    transport_type="nonregular",
                    trip_count=Decimal(1),
                    unit_rate_won=None,
                    cost_won=amount,
                    vehicle_type=vehicle_type,
                ),
                entry_limit,
            )
        if not displayed_total_found:
            raise WorkbookStructureError(
                f"{sheet.title} displayed cached total is required"
            )
        return rows


def _validate_worksheet_bounds(sheet) -> None:
    if sheet.max_row > MAX_WORKSHEET_ROWS:
        raise WorkbookStructureError(
            f"{sheet.title} exceeds worksheet row limit of {MAX_WORKSHEET_ROWS} "
            f"(found {sheet.max_row})"
        )
    if sheet.max_column > MAX_WORKSHEET_COLUMNS:
        raise WorkbookStructureError(
            f"{sheet.title} exceeds worksheet column limit of "
            f"{MAX_WORKSHEET_COLUMNS} (found {sheet.max_column})"
        )


def _preflight_workbook_resources(path: Path) -> None:
    try:
        with ZipFile(path, "r") as archive:
            sheet_refs = _read_workbook_sheet_refs(archive)
            worksheet_members, chartsheet_members = _resolve_sheet_members(
                archive, sheet_refs
            )
            for sheet_name, member_name in chartsheet_members:
                _preflight_chartsheet_xml(archive, sheet_name, member_name)
            total_cells = 0
            total_materialized_cells = 0
            for sheet_name, member_name in worksheet_members:
                cell_count, materialized_cell_count = _preflight_worksheet_xml(
                    archive, sheet_name, member_name
                )
                comment_count = _preflight_worksheet_relationships(
                    archive, sheet_name, member_name
                )
                materialized_cell_count += comment_count
                _validate_materialized_cell_count(
                    sheet_name, materialized_cell_count
                )
                total_cells += cell_count
                if total_cells > MAX_PREFLIGHT_TOTAL_CELLS:
                    raise WorkbookStructureError(
                        "Workbook exceeds workbook cell-record limit of "
                        f"{MAX_PREFLIGHT_TOTAL_CELLS}"
                    )
                total_materialized_cells += materialized_cell_count
                if total_materialized_cells > MAX_PREFLIGHT_TOTAL_CELLS:
                    raise WorkbookStructureError(
                        "Workbook exceeds workbook cell materialization limit of "
                        f"{MAX_PREFLIGHT_TOTAL_CELLS}"
                    )
    except WorkbookStructureError:
        raise
    except (BadZipFile, LargeZipFile, OSError, ValueError) as error:
        raise WorkbookStructureError(
            f"Workbook resource preflight failed: {error}"
        ) from error


def _read_workbook_sheet_refs(archive: ZipFile) -> list[tuple[str, str]]:
    refs: list[tuple[str, str]] = []
    try:
        with archive.open("xl/workbook.xml") as source:
            for _, element in iterparse(source, events=("end",)):
                if _xml_local_name(element.tag) == "sheet":
                    relationship_id = element.attrib.get(_DOCUMENT_RELATIONSHIP_ID)
                    if not relationship_id:
                        raise WorkbookStructureError(
                            "Workbook sheet is missing a relationship id"
                        )
                    refs.append(
                        (element.attrib.get("name") or "<unnamed>", relationship_id)
                    )
                    if len(refs) > MAX_PREFLIGHT_WORKSHEETS:
                        raise WorkbookStructureError(
                            "Workbook exceeds worksheet count limit of "
                            f"{MAX_PREFLIGHT_WORKSHEETS}"
                        )
                element.clear()
    except KeyError as error:
        raise WorkbookStructureError("Workbook XML is missing") from error
    except ParseError as error:
        raise WorkbookStructureError("Workbook XML is malformed") from error
    return refs


def _resolve_sheet_members(
    archive: ZipFile, sheet_refs: list[tuple[str, str]]
) -> tuple[list[tuple[str, str]], list[tuple[str, str]]]:
    wanted_ids = {relationship_id for _, relationship_id in sheet_refs}
    relationships: dict[str, tuple[str, str, str | None]] = {}
    try:
        with archive.open("xl/_rels/workbook.xml.rels") as source:
            for _, element in iterparse(source, events=("end",)):
                if _xml_local_name(element.tag) == "Relationship":
                    relationship_id = element.attrib.get("Id")
                    if relationship_id in wanted_ids:
                        if relationship_id in relationships:
                            raise WorkbookStructureError(
                                "Workbook relationships contain a duplicate id"
                            )
                        relationships[relationship_id] = (
                            element.attrib.get("Type", ""),
                            element.attrib.get("Target", ""),
                            element.attrib.get("TargetMode"),
                        )
                element.clear()
    except KeyError as error:
        raise WorkbookStructureError("Workbook relationships XML is missing") from error
    except ParseError as error:
        raise WorkbookStructureError(
            "Workbook relationships XML is malformed"
        ) from error

    worksheets: list[tuple[str, str]] = []
    chartsheets: list[tuple[str, str]] = []
    for sheet_name, relationship_id in sheet_refs:
        relationship = relationships.get(relationship_id)
        if relationship is None:
            raise WorkbookStructureError(
                f"Workbook sheet {sheet_name} has no relationship"
            )
        relationship_type, target, target_mode = relationship
        # OpenPyXL treats every workbook sheet relationship except a chartsheet
        # as a worksheet and eagerly parses its target.
        if "chartsheet" in relationship_type:
            if target_mode == "External":
                raise WorkbookStructureError(
                    f"Workbook sheet {sheet_name} has an external chartsheet relationship"
                )
            chartsheets.append(
                (sheet_name, _safe_sheet_member(sheet_name, target, "chartsheets"))
            )
            continue
        if target_mode == "External":
            raise WorkbookStructureError(
                f"Workbook sheet {sheet_name} has an external worksheet relationship"
            )
        worksheets.append(
            (sheet_name, _safe_sheet_member(sheet_name, target, "worksheets"))
        )
    return worksheets, chartsheets


def _safe_sheet_member(sheet_name: str, target: str, sheet_kind: str) -> str:
    if (
        not target
        or "\\" in target
        or "\x00" in target
        or ":" in target
        or "?" in target
        or "#" in target
    ):
        raise WorkbookStructureError(
            f"Workbook sheet {sheet_name} has an unsafe {sheet_kind[:-1]} "
            "relationship target"
        )
    target_path = PurePosixPath(target)
    if ".." in target_path.parts:
        raise WorkbookStructureError(
            f"Workbook sheet {sheet_name} has an unsafe {sheet_kind[:-1]} "
            "relationship target"
        )
    candidate = target[1:] if target.startswith("/") else posix_join("xl", target)
    member_name = posix_normpath(candidate)
    parts = PurePosixPath(member_name).parts
    if len(parts) < 3 or parts[:2] != ("xl", sheet_kind):
        raise WorkbookStructureError(
            f"Workbook sheet {sheet_name} has an unsafe {sheet_kind[:-1]} "
            "relationship target"
        )
    return member_name


class _SizeLimitedReader:
    def __init__(self, source, limit: int, description: str) -> None:
        self._source = source
        self._limit = limit
        self._description = description
        self._bytes_read = 0

    def read(self, size: int = -1) -> bytes:
        chunk = self._source.read(size)
        self._bytes_read += len(chunk)
        if self._bytes_read > self._limit:
            raise WorkbookStructureError(
                f"{self._description} exceeds XML size limit of {self._limit}"
            )
        return chunk


def _preflight_chartsheet_xml(
    archive: ZipFile, sheet_name: str, member_name: str
) -> None:
    member = _required_xml_member(
        archive, member_name, f"Chartsheet XML is missing for {sheet_name}"
    )
    if member.file_size > MAX_PREFLIGHT_RELATED_XML_BYTES:
        raise WorkbookStructureError(
            f"{sheet_name} exceeds chartsheet XML size limit of "
            f"{MAX_PREFLIGHT_RELATED_XML_BYTES}"
        )
    try:
        with archive.open(member) as source:
            limited_source = _SizeLimitedReader(
                source,
                MAX_PREFLIGHT_RELATED_XML_BYTES,
                f"{sheet_name} chartsheet",
            )
            for _, element in iterparse(limited_source, events=("end",)):
                element.clear()
    except ParseError as error:
        raise WorkbookStructureError(
            f"Chartsheet XML is malformed for {sheet_name}"
        ) from error


def _preflight_worksheet_relationships(
    archive: ZipFile, sheet_name: str, worksheet_member: str
) -> int:
    worksheet_path = PurePosixPath(worksheet_member)
    relationships_member = str(
        worksheet_path.parent
        / "_rels"
        / f"{worksheet_path.name}.rels"
    )
    try:
        member = archive.getinfo(relationships_member)
    except KeyError:
        return 0
    if member.is_dir():
        raise WorkbookStructureError(
            f"Worksheet relationships XML is invalid for {sheet_name}"
        )
    if member.file_size > MAX_PREFLIGHT_RELATED_XML_BYTES:
        raise WorkbookStructureError(
            f"{sheet_name} exceeds worksheet relationships XML size limit of "
            f"{MAX_PREFLIGHT_RELATED_XML_BYTES}"
        )

    relationships: list[tuple[str, str, str | None]] = []
    relationship_ids: set[str] = set()
    try:
        with archive.open(member) as source:
            limited_source = _SizeLimitedReader(
                source,
                MAX_PREFLIGHT_RELATED_XML_BYTES,
                f"{sheet_name} worksheet relationships",
            )
            for _, element in iterparse(limited_source, events=("end",)):
                if _xml_local_name(element.tag) == "Relationship":
                    relationship_id = element.attrib.get("Id")
                    if not relationship_id or relationship_id in relationship_ids:
                        raise WorkbookStructureError(
                            f"Worksheet relationships are invalid for {sheet_name}"
                        )
                    relationship_ids.add(relationship_id)
                    relationships.append(
                        (
                            element.attrib.get("Type", ""),
                            element.attrib.get("Target", ""),
                            element.attrib.get("TargetMode"),
                        )
                    )
                element.clear()
    except ParseError as error:
        raise WorkbookStructureError(
            f"Worksheet relationships XML is malformed for {sheet_name}"
        ) from error

    comment_count = 0
    for relationship_type, target, target_mode in relationships:
        if target_mode == "External":
            if relationship_type != _HYPERLINK_RELATIONSHIP:
                raise WorkbookStructureError(
                    f"{sheet_name} has an unsupported external worksheet relationship"
                )
            continue

        target_member = _safe_worksheet_related_member(
            sheet_name, worksheet_member, target
        )
        if relationship_type not in _OPENPYXL_EAGER_SHEET_RELATIONSHIPS:
            continue
        if relationship_type == _COMMENTS_RELATIONSHIP:
            parts = PurePosixPath(target_member).parts
            if len(parts) < 3 or parts[:2] != ("xl", "comments"):
                raise WorkbookStructureError(
                    f"{sheet_name} has an unsafe comment relationship target"
                )
            related_member = _required_xml_member(
                archive,
                target_member,
                f"Comment XML is missing for {sheet_name}",
            )
            comment_count += _preflight_comment_xml(
                archive, sheet_name, related_member
            )
            continue
        _required_xml_member(
            archive,
            target_member,
            f"Worksheet related XML is missing for {sheet_name}",
        )
        # Tables, drawings, and pivot definitions do not create worksheet cells.
        # Their direct parts and targets are checked here; the import service's
        # ZIP member and aggregate limits bound their remaining object graphs.
    return comment_count


def _safe_worksheet_related_member(
    sheet_name: str, worksheet_member: str, target: str
) -> str:
    if (
        not target
        or "\\" in target
        or "\x00" in target
        or ":" in target
        or "?" in target
        or "#" in target
    ):
        raise WorkbookStructureError(
            f"{sheet_name} has an unsafe worksheet relationship target"
        )
    parts = [] if target.startswith("/") else list(
        PurePosixPath(worksheet_member).parent.parts
    )
    for part in PurePosixPath(target).parts:
        if part in ("/", "."):
            continue
        if part == "..":
            if len(parts) <= 1:
                raise WorkbookStructureError(
                    f"{sheet_name} has an unsafe worksheet relationship target"
                )
            parts.pop()
            continue
        parts.append(part)
    if len(parts) < 2 or parts[0] != "xl":
        raise WorkbookStructureError(
            f"{sheet_name} has an unsafe worksheet relationship target"
        )
    return "/".join(parts)


def _required_xml_member(archive: ZipFile, member_name: str, message: str):
    try:
        member = archive.getinfo(member_name)
    except KeyError as error:
        raise WorkbookStructureError(message) from error
    if member.is_dir():
        raise WorkbookStructureError(message)
    return member


def _preflight_comment_xml(archive: ZipFile, sheet_name: str, member) -> int:
    if member.file_size > MAX_PREFLIGHT_RELATED_XML_BYTES:
        raise WorkbookStructureError(
            f"{sheet_name} exceeds comment XML size limit of "
            f"{MAX_PREFLIGHT_RELATED_XML_BYTES}"
        )
    comment_count = 0
    try:
        with archive.open(member) as source:
            limited_source = _SizeLimitedReader(
                source,
                MAX_PREFLIGHT_RELATED_XML_BYTES,
                f"{sheet_name} comment",
            )
            for _, element in iterparse(limited_source, events=("end",)):
                if _xml_local_name(element.tag) == "comment":
                    comment_count += 1
                    _validate_materialized_cell_count(sheet_name, comment_count)
                    _validate_preflight_cell_reference(
                        element.attrib.get("ref"), sheet_name
                    )
                element.clear()
    except ParseError as error:
        raise WorkbookStructureError(
            f"Comment XML is malformed for {sheet_name}"
        ) from error
    return comment_count


def _preflight_worksheet_xml(
    archive: ZipFile, sheet_name: str, member_name: str
) -> tuple[int, int]:
    try:
        member = archive.getinfo(member_name)
    except KeyError as error:
        raise WorkbookStructureError(
            f"Worksheet XML is missing for {sheet_name}"
        ) from error
    if member.is_dir():
        raise WorkbookStructureError(f"Worksheet XML is invalid for {sheet_name}")
    if member.file_size > MAX_PREFLIGHT_WORKSHEET_XML_BYTES:
        raise WorkbookStructureError(
            f"{sheet_name} exceeds worksheet XML size limit of "
            f"{MAX_PREFLIGHT_WORKSHEET_XML_BYTES}"
        )

    cell_count = 0
    materialized_cell_count = 0
    row_record_count = 0
    column_dimension_count = 0
    inferred_row_number = 0
    try:
        with archive.open(member) as source:
            for _, element in iterparse(source, events=("end",)):
                local_name = _xml_local_name(element.tag)
                if local_name == "c":
                    cell_count += 1
                    if cell_count > MAX_PREFLIGHT_WORKSHEET_CELLS:
                        raise WorkbookStructureError(
                            f"{sheet_name} exceeds worksheet cell-record limit of "
                            f"{MAX_PREFLIGHT_WORKSHEET_CELLS}"
                        )
                    materialized_cell_count += 1
                    _validate_materialized_cell_count(
                        sheet_name, materialized_cell_count
                    )
                    _validate_preflight_cell_reference(
                        element.attrib.get("r"), sheet_name
                    )
                elif local_name == "dimension":
                    reference = element.attrib.get("ref")
                    if reference:
                        _preflight_range_area(reference, sheet_name)
                elif local_name == "mergeCell":
                    materialized_cell_count += _preflight_range_area(
                        element.attrib.get("ref"), sheet_name
                    )
                    _validate_materialized_cell_count(
                        sheet_name, materialized_cell_count
                    )
                elif local_name == "hyperlink":
                    materialized_cell_count += _preflight_range_area(
                        element.attrib.get("ref"), sheet_name
                    )
                    _validate_materialized_cell_count(
                        sheet_name, materialized_cell_count
                    )
                elif local_name == "row":
                    row_record_count += 1
                    if row_record_count > MAX_PREFLIGHT_ROWS:
                        raise WorkbookStructureError(
                            f"{sheet_name} exceeds preflight row limit of "
                            f"{MAX_PREFLIGHT_ROWS} (worksheet row-record limit of "
                            f"{MAX_PREFLIGHT_ROWS})"
                        )
                    row_reference = element.attrib.get("r")
                    if row_reference is None:
                        row_number = inferred_row_number + 1
                    else:
                        try:
                            row_number = int(row_reference)
                        except ValueError as error:
                            raise WorkbookStructureError(
                                f"{sheet_name} contains an invalid row reference"
                            ) from error
                    inferred_row_number = row_number
                    if not 1 <= row_number <= MAX_PREFLIGHT_ROWS:
                        raise WorkbookStructureError(
                            f"{sheet_name} exceeds preflight row limit of "
                            f"{MAX_PREFLIGHT_ROWS}"
                        )
                elif local_name == "col":
                    column_dimension_count += 1
                    if (
                        column_dimension_count
                        > MAX_PREFLIGHT_COLUMN_DIMENSION_RECORDS
                    ):
                        raise WorkbookStructureError(
                            f"{sheet_name} exceeds worksheet column-dimension "
                            "record limit of "
                            f"{MAX_PREFLIGHT_COLUMN_DIMENSION_RECORDS}"
                        )
                    _validate_preflight_column_dimension(element.attrib, sheet_name)
                element.clear()
    except ParseError as error:
        raise WorkbookStructureError(
            f"Worksheet XML is malformed for {sheet_name}"
        ) from error
    return cell_count, materialized_cell_count


def _validate_preflight_column_dimension(
    attributes: dict[str, str], sheet_name: str
) -> None:
    minimum = _parse_bounded_decimal(attributes.get("min"), MAX_EXCEL_COLUMNS)
    maximum = _parse_bounded_decimal(attributes.get("max"), MAX_EXCEL_COLUMNS)
    if minimum is None or maximum is None or maximum < minimum:
        raise WorkbookStructureError(
            f"{sheet_name} contains an invalid column dimension"
        )


def _parse_bounded_decimal(value: str | None, maximum: int) -> int | None:
    if (
        value is None
        or not value
        or len(value) > len(str(maximum))
        or not value.isascii()
        or not value.isdecimal()
    ):
        return None
    parsed = int(value)
    return parsed if 1 <= parsed <= maximum else None


def _validate_materialized_cell_count(sheet_name: str, count: int) -> None:
    if count > MAX_PREFLIGHT_WORKSHEET_CELLS:
        raise WorkbookStructureError(
            f"{sheet_name} exceeds worksheet cell materialization limit of "
            f"{MAX_PREFLIGHT_WORKSHEET_CELLS}"
        )


def _preflight_range_area(reference: str | None, sheet_name: str) -> int:
    coordinates = (reference or "").split(":")
    if len(coordinates) not in (1, 2):
        raise WorkbookStructureError(
            f"{sheet_name} contains an invalid cell range reference"
        )
    start_row, start_column = _validate_preflight_cell_reference(
        coordinates[0], sheet_name
    )
    if len(coordinates) == 1:
        return 1
    end_row, end_column = _validate_preflight_cell_reference(
        coordinates[1], sheet_name
    )
    if end_row < start_row or end_column < start_column:
        raise WorkbookStructureError(
            f"{sheet_name} contains an invalid cell range reference"
        )
    return (end_row - start_row + 1) * (end_column - start_column + 1)


def _validate_preflight_cell_reference(
    reference: str | None, sheet_name: str
) -> tuple[int, int]:
    match = _CELL_REFERENCE.fullmatch(reference or "")
    if match is None:
        raise WorkbookStructureError(
            f"{sheet_name} contains an invalid cell reference"
        )
    row_number = int(match.group("row"))
    column_number = 0
    for character in match.group("column").upper():
        column_number = column_number * 26 + ord(character) - ord("A") + 1
        if column_number > MAX_PREFLIGHT_COLUMNS:
            break
    if row_number > MAX_PREFLIGHT_ROWS:
        raise WorkbookStructureError(
            f"{sheet_name} exceeds preflight row limit of {MAX_PREFLIGHT_ROWS}"
        )
    if column_number > MAX_PREFLIGHT_COLUMNS:
        raise WorkbookStructureError(
            f"{sheet_name} exceeds preflight column limit of {MAX_PREFLIGHT_COLUMNS}"
        )
    return row_number, column_number


def _xml_local_name(name: str) -> str:
    return name.rsplit("}", 1)[-1]


def _append_parsed_entry(
    rows: list[ParsedTransportEntry], entry: ParsedTransportEntry, limit: int
) -> None:
    if len(rows) >= limit:
        raise WorkbookStructureError(
            f"Workbook parsed entry limit of {MAX_PARSED_ENTRIES} exceeded"
        )
    rows.append(entry)


def _find_regular_header(sheet, formula_sheet) -> int | None:
    for row_number in range(1, min(sheet.max_row, 50) + 1):
        if _is_regular_header_view(sheet, row_number):
            if not _is_regular_header(sheet, formula_sheet, row_number):
                return None
            return row_number
    return None


def _is_regular_header(sheet, formula_sheet, row_number: int) -> bool:
    if not all(
        _is_regular_header_view(workbook_sheet, row_number)
        for workbook_sheet in (sheet, formula_sheet)
    ):
        return False
    if any(
        _has_value(workbook_sheet.cell(row_number, 35).value)
        for workbook_sheet in (sheet, formula_sheet)
    ):
        return False
    return all(
        sheet.cell(row_number, column).value
        == formula_sheet.cell(row_number, column).value
        and not _is_formula(formula_sheet.cell(row_number, column).value)
        for column in (1, 38)
    )


def _is_regular_header_view(sheet, row_number: int) -> bool:
    destination = _header_text(sheet.cell(row_number, 2).value)
    total_count = _header_text(sheet.cell(row_number, 34).value)
    unit = _header_text(sheet.cell(row_number, 36).value)
    subtotal = _header_text(sheet.cell(row_number, 37).value)
    day_values = [sheet.cell(row_number, column).value for column in range(3, 34)]
    return (
        destination in _REGULAR_COLUMN_B_HEADERS
        and total_count in _REGULAR_TOTAL_COUNT_HEADERS
        and unit in _REGULAR_UNIT_HEADERS
        and subtotal in _REGULAR_SUBTOTAL_HEADERS
        and day_values == list(range(1, 32))
    )


def _find_subcontract_header(
    sheet, formula_sheet
) -> tuple[int, int, int, int, int | None, int] | None:
    for row_number in range(1, min(sheet.max_row, 50) + 1):
        header = _subcontract_header_view(sheet, row_number)
        if header is None:
            continue
        if (
            _subcontract_header_view(formula_sheet, row_number) != header
            or any(
                _is_formula(formula_sheet.cell(row_number, column).value)
                for column in range(1, _SUBCONTRACT_LAST_FOOTPRINT_COLUMN + 1)
            )
        ):
            return None
        return (row_number, *header, _SUBCONTRACT_LAST_FOOTPRINT_COLUMN)
    return None


def _subcontract_header_view(
    sheet, row_number: int
) -> tuple[int, int, int, int | None] | None:
    columns: dict[str, int] = {}
    for column in range(1, sheet.max_column + 1):
        value = _header_text(sheet.cell(row_number, column).value)
        if value in _DATE_HEADERS:
            columns.setdefault("date", column)
        if value in _DESTINATION_HEADERS:
            columns.setdefault("destination", column)
        if value in _TONNAGE_HEADERS:
            columns.setdefault("tonnage", column)
        if value.casefold() in _NUMBER_HEADERS:
            columns.setdefault("number", column)
    amount_header = _header_text(sheet.cell(row_number, 9).value)
    if (
        {"date", "destination", "tonnage"} <= columns.keys()
        and amount_header in _AMOUNT_HEADERS
    ):
        return (
            columns["date"],
            columns["destination"],
            columns["tonnage"],
            columns.get("number"),
        )
    return None


def _header_text(value: object) -> str:
    return "" if value is None else str(value).strip().replace(" ", "")


def _has_value(value: object) -> bool:
    return value is not None and value != ""


def _is_blank_or_numeric_zero(value: object) -> bool:
    if not _has_value(value):
        return True
    if isinstance(value, bool):
        return False
    try:
        numeric_value = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return False
    return numeric_value.is_finite() and numeric_value == 0


def _is_displayed_total_row(sheet, formula_sheet, row_number: int) -> bool:
    label_columns = [
        column
        for column in _SUBCONTRACT_TOTAL_LABEL_COLUMNS
        if _header_text(sheet.cell(row_number, column).value) in _TOTAL_LABELS
    ]
    if len(label_columns) != 1:
        return False
    allowed_columns = {label_columns[0], 9}
    return any(
        _has_value(workbook_sheet.cell(row_number, 9).value)
        for workbook_sheet in (sheet, formula_sheet)
    ) and not any(
        _has_value(workbook_sheet.cell(row_number, column).value)
        for workbook_sheet in (sheet, formula_sheet)
        for column in range(1, _SUBCONTRACT_LAST_FOOTPRINT_COLUMN + 1)
        if column not in allowed_columns
    )


def _is_regular_footer_or_note(sheet, formula_sheet, row_number: int) -> bool:
    if not _has_regular_footer_label(sheet, formula_sheet, row_number):
        return False
    cached_marker = sheet.cell(row_number, 38).value
    formula_marker = formula_sheet.cell(row_number, 38).value
    return (
        not isinstance(cached_marker, bool)
        and isinstance(cached_marker, (int, float, Decimal))
        and (
            _is_formula(formula_marker)
            or (
                not isinstance(formula_marker, bool)
                and isinstance(formula_marker, (int, float, Decimal))
            )
        )
    )


def _has_regular_footer_label(sheet, formula_sheet, row_number: int) -> bool:
    if any(
        _has_value(workbook_sheet.cell(row_number, column).value)
        for workbook_sheet in (sheet, formula_sheet)
        for column in range(1, 37)
    ):
        return False
    cached_label = sheet.cell(row_number, 37).value
    formula_label = formula_sheet.cell(row_number, 37).value
    if not all(
        isinstance(value, str) and bool(value.strip()) and not _is_formula(value)
        for value in (cached_label, formula_label)
    ):
        return False
    return True


def _is_regular_blank_or_group_row(sheet, formula_sheet, row_number: int) -> bool:
    cached_group = sheet.cell(row_number, 1).value
    formula_group = formula_sheet.cell(row_number, 1).value
    if _has_value(formula_group) and not _has_value(cached_group):
        return False
    return not any(
        _has_value(workbook_sheet.cell(row_number, column).value)
        for workbook_sheet in (sheet, formula_sheet)
        for column in range(2, 39)
    )


def _is_regular_inactive_row(sheet, formula_sheet, row_number: int) -> bool:
    cached_group = sheet.cell(row_number, 1).value
    formula_group = formula_sheet.cell(row_number, 1).value
    if _has_value(formula_group) and not _has_value(cached_group):
        return False
    if any(
        _has_value(workbook_sheet.cell(row_number, column).value)
        for workbook_sheet in (sheet, formula_sheet)
        for column in (35, 38)
    ):
        return False
    return (
        not _has_value(sheet.cell(row_number, 36).value)
        and all(
            _is_blank_or_numeric_zero(sheet.cell(row_number, column).value)
            for column in range(3, 34)
        )
        and _is_blank_or_numeric_zero(sheet.cell(row_number, 34).value)
        and _is_blank_or_numeric_zero(sheet.cell(row_number, 37).value)
    )


def _is_subcontract_template_row(
    sheet,
    formula_sheet,
    row_number: int,
    number_column: int | None,
    last_footprint_column: int,
) -> bool:
    if number_column is None:
        return False
    return _has_value(sheet.cell(row_number, number_column).value) and not any(
        _has_value(workbook_sheet.cell(row_number, column).value)
        for workbook_sheet in (sheet, formula_sheet)
        for column in range(1, last_footprint_column + 1)
        if column != number_column
    )


def _is_subcontract_formula_footer(
    sheet, formula_sheet, row_number: int, last_footprint_column: int
) -> bool:
    return _is_formula(formula_sheet.cell(row_number, 9).value) and not any(
        _has_value(workbook_sheet.cell(row_number, column).value)
        for workbook_sheet in (sheet, formula_sheet)
        for column in range(1, last_footprint_column + 1)
        if column != 9
    )


def _required_cached_value(
    cached_value: object,
    formula_value: object,
    sheet_name: str,
    row_number: int,
    coordinate: str,
    label: str,
) -> object:
    if _has_value(cached_value):
        return cached_value
    if isinstance(formula_value, str) and formula_value.startswith("="):
        raise WorkbookStructureError(
            f"{sheet_name} row {row_number} cell {coordinate} cached value is "
            "missing; recalculate and save in Excel"
        )
    raise WorkbookStructureError(
        f"{sheet_name} row {row_number} {label} is required"
    )


def _is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _optional_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _destination_alias(value: object, label: str) -> str | None:
    if value is None:
        return None
    if not isinstance(value, str):
        raise WorkbookStructureError(f"{label} must be text")
    text = value.strip()
    return text or None


def _validate_regular_sheet_month(sheet_name: str, report_month: str) -> None:
    match = _REGULAR_SHEET.fullmatch(sheet_name)
    if match is None:
        raise WorkbookStructureError(f"Invalid regular sheet name: {sheet_name}")
    if int(match.group("month")) != int(report_month[-2:]):
        raise WorkbookStructureError(
            f"Regular sheet month does not match report_month {report_month}"
        )


def _validate_calendar_day(
    report_month: str, day: int, sheet_name: str, row_number: int
) -> date:
    year, month = (int(part) for part in report_month.split("-"))
    try:
        return date(year, month, day)
    except ValueError as error:
        raise WorkbookStructureError(
            f"{sheet_name} row {row_number} has invalid calendar date "
            f"{report_month}-{day:02d}"
        ) from error


def _nonnegative_decimal(value: object, label: str) -> Decimal:
    if isinstance(value, bool):
        raise WorkbookStructureError(f"{label} must be a nonnegative number")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError) as error:
        raise WorkbookStructureError(f"{label} must be a nonnegative number") from error
    if not result.is_finite() or result < 0:
        raise WorkbookStructureError(f"{label} must be a nonnegative number")
    return result


def _trip_count_decimal(value: object, label: str) -> Decimal:
    if isinstance(value, bool):
        raise WorkbookStructureError(f"{label} must be a bounded nonnegative number")
    raw_text = str(value).strip()
    if "e" in raw_text.casefold():
        raise WorkbookStructureError(f"{label} must be a bounded nonnegative number")
    try:
        result = Decimal(raw_text)
    except (InvalidOperation, ValueError) as error:
        raise WorkbookStructureError(
            f"{label} must be a bounded nonnegative number"
        ) from error
    if (
        not result.is_finite()
        or result < 0
        or result > MAX_TRIP_COUNT
        or result.as_tuple().exponent < -MAX_TRIP_COUNT_SCALE
    ):
        raise WorkbookStructureError(f"{label} must be a bounded nonnegative number")
    canonical_trip_count(result, label)
    return result


def canonical_trip_count(value: Decimal, label: str = "trip count") -> str:
    result = _trip_count_decimal_without_canonical_check(value, label)
    if result == 0:
        return "0"
    text = format(result, "f")
    canonical = text.rstrip("0").rstrip(".") if "." in text else text
    if len(canonical) > MAX_TRIP_COUNT_CANONICAL_LENGTH:
        raise WorkbookStructureError(f"{label} must be a bounded nonnegative number")
    return canonical


def _trip_count_decimal_without_canonical_check(
    value: Decimal, label: str
) -> Decimal:
    if not isinstance(value, Decimal):
        raise WorkbookStructureError(f"{label} must be a bounded nonnegative number")
    if (
        not value.is_finite()
        or value < 0
        or value > MAX_TRIP_COUNT
        or value.as_tuple().exponent < -MAX_TRIP_COUNT_SCALE
    ):
        raise WorkbookStructureError(f"{label} must be a bounded nonnegative number")
    return value


def _integer_won(value: object, label: str) -> int:
    if value is None or isinstance(value, bool):
        raise WorkbookStructureError(f"{label} must be an integer won amount")
    try:
        decimal_value = Decimal(str(value))
    except (InvalidOperation, ValueError) as error:
        raise WorkbookStructureError(f"{label} must be an integer won amount") from error
    if (
        not decimal_value.is_finite()
        or decimal_value < 0
        or decimal_value != decimal_value.to_integral_value()
        or decimal_value > _SQLITE_MAX
    ):
        raise WorkbookStructureError(f"{label} must be an integer won amount")
    return int(decimal_value)


def _transport_date(value: object, label: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        normalized = value.strip().replace(".", "-").replace("/", "-")
        try:
            return date.fromisoformat(normalized.rstrip("-"))
        except ValueError as error:
            raise WorkbookStructureError(f"{label} must be a valid date") from error
    raise WorkbookStructureError(f"{label} must be a valid date")


def _is_in_subcontract_billing_window(
    transport_date: date, report_month: str
) -> bool:
    year, month = (int(part) for part in report_month.split("-"))
    report_month_start = date(year, month, 1)
    previous_month_day = report_month_start - timedelta(days=1)
    return (transport_date.year, transport_date.month) in {
        (year, month),
        (previous_month_day.year, previous_month_day.month),
    }
