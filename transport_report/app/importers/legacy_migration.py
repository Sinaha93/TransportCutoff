from __future__ import annotations

import hashlib
import json
import math
import re
import sqlite3
import unicodedata
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.db import Database
from app.domain.calculations import (
    MAX_QUANTITY_EA,
    MAX_QUANTITY_SCALE,
    MAX_SQLITE_INTEGER,
    HistoricalAverages,
    HistoricalValueKind,
    calculate_destination,
    calculate_group,
    calculate_total,
    historical_averages,
)
from app.domain.models import Destination, GroupMember


LEGACY_SOURCE_TYPE = "legacy_workbook"
_HISTORY_SHEET = "누적 데이터"
_PLAN_SHEET = "26년 월계획"
_HISTORY_HEADERS = ("연도", "월", "연도+월", "납품처", "실적수량", "실적운반비")
_PLAN_QUANTITY_HEADER = "수량"
_PLAN_COST_HEADER = "운반비"
_MONTH = re.compile(r"^[0-9]{4}-(0[1-9]|1[0-2])$")


class LegacyMigrationError(Exception):
    """Base error for the guarded legacy-workbook migration."""


class LegacyMigrationConfirmationError(LegacyMigrationError):
    """Raised when commit is attempted without explicit UI confirmation."""


class LegacyMigrationRevisionError(LegacyMigrationError):
    """Raised when source, masters, or target rows changed after dry-run."""


class LegacyMigrationBlockedError(LegacyMigrationError):
    """Raised when structure, aliases, reconciliation, or conflicts block commit."""


@dataclass(frozen=True, slots=True)
class MigrationIssue:
    code: str
    message: str
    source_locator: str
    expected: str | None = None
    actual: str | None = None
    blocking: bool = True


@dataclass(frozen=True, slots=True)
class UnknownAlias:
    alias: str
    source_type: str
    source_locator: str


@dataclass(frozen=True, slots=True)
class StagedPlanRecord:
    report_month: str
    destination_id: int
    destination_alias: str
    quantity_ea: Decimal
    cost_won: int
    source_sheet: str
    source_row: int


@dataclass(frozen=True, slots=True)
class StagedActualRecord:
    report_month: str
    destination_id: int
    destination_alias: str
    quantity_ea: Decimal | None
    cost_won: int | None
    source_sheet: str
    source_row: int


@dataclass(frozen=True, slots=True)
class MigrationSummary:
    plan_records: int
    actual_records: int
    current_plan_records: int
    current_actual_records: int
    months: tuple[str, ...]
    rows_to_insert: int
    rows_unchanged: int


@dataclass(frozen=True, slots=True)
class MonthLockSnapshot:
    report_month: str
    is_locked: bool


@dataclass(frozen=True, slots=True)
class TargetRowSnapshot:
    table: str
    report_month: str
    destination_id: int
    existing_value: tuple[str, ...] | None
    incoming_value: tuple[str, ...]
    action: str


@dataclass(frozen=True, slots=True)
class DatabaseSnapshot:
    month_locks: tuple[MonthLockSnapshot, ...]
    targets: tuple[TargetRowSnapshot, ...]


@dataclass(frozen=True, slots=True)
class ReconciledMetrics:
    planned_quantity: Decimal | None
    planned_cost_won: int | None
    actual_quantity: Decimal | None
    actual_cost_won: int | None


@dataclass(frozen=True, slots=True)
class GroupReconciliation:
    group_id: int
    name: str
    result: ReconciledMetrics


@dataclass(frozen=True, slots=True)
class ReconciliationControl:
    name: str
    source_locator: str
    expected: Decimal | int | None
    actual: Decimal | int | None
    matches: bool


@dataclass(frozen=True, slots=True)
class MigrationReconciliation:
    current_total: ReconciledMetrics
    groups: tuple[GroupReconciliation, ...]
    actual_quantity_history: HistoricalAverages
    actual_cost_history: HistoricalAverages
    controls: tuple[ReconciliationControl, ...]


@dataclass(frozen=True, slots=True)
class MigrationDryRun:
    source_path: Path
    source_sha256: str
    master_revision: str
    database_snapshot: DatabaseSnapshot
    database_revision: str
    revision: str
    report_month: str
    plans: tuple[StagedPlanRecord, ...]
    actuals: tuple[StagedActualRecord, ...]
    summary: MigrationSummary
    unknown_aliases: tuple[UnknownAlias, ...]
    issues: tuple[MigrationIssue, ...]
    reconciliation: MigrationReconciliation

    @property
    def can_commit(self) -> bool:
        return not self.unknown_aliases and not any(
            issue.blocking for issue in self.issues
        )


@dataclass(frozen=True, slots=True)
class MigrationCommitResult:
    source_sha256: str
    revision: str
    inserted: int
    unchanged: int


@dataclass(frozen=True, slots=True)
class _SourceControl:
    name: str
    locator: str
    value: Decimal | int | None


@dataclass(frozen=True, slots=True)
class _ParsedSource:
    plans: tuple[StagedPlanRecord, ...]
    actuals: tuple[StagedActualRecord, ...]
    unknown_aliases: tuple[UnknownAlias, ...]
    issues: tuple[MigrationIssue, ...]
    controls: tuple[_SourceControl, ...]


class LegacyMigrationService:
    """Stage, reconcile, and atomically commit the one-time legacy workbook."""

    def __init__(self, database: Database) -> None:
        self.database = database

    def dry_run(
        self, source_path: Path | str, *, report_month: str = "2026-08"
    ) -> MigrationDryRun:
        source = Path(source_path).resolve()
        _require_report_month(report_month)
        if not source.is_file():
            raise LegacyMigrationError(f"원본 파일을 찾을 수 없습니다: {source}")
        source_sha256 = _sha256_file(source)

        with _read_only_connection(self.database) as connection:
            destinations = _load_destinations(connection)
            groups = _load_groups(connection)
            aliases, resolution_issues = _load_resolution(connection)
            master_revision = _master_revision(connection)
            parsed = _parse_source(
                source,
                report_month=report_month,
                aliases=aliases,
                initial_issues=resolution_issues,
            )
            database_issues, database_snapshot, insert_count, unchanged_count = (
                _inspect_database_state(
                    connection,
                    parsed.plans,
                    parsed.actuals,
                )
            )
        if _sha256_file(source) != source_sha256:
            raise LegacyMigrationRevisionError(
                "드라이런 도중 원본 파일이 변경되었습니다. 다시 드라이런하세요."
            )

        reconciliation, reconciliation_issues = _reconcile(
            report_month=report_month,
            plans=parsed.plans,
            actuals=parsed.actuals,
            destinations=destinations,
            groups=groups,
            source_controls=parsed.controls,
        )
        issues = tuple(
            sorted(
                (*parsed.issues, *database_issues, *reconciliation_issues),
                key=lambda item: (item.source_locator, item.code, item.message),
            )
        )
        months = tuple(
            sorted(
                {
                    *(record.report_month for record in parsed.plans),
                    *(record.report_month for record in parsed.actuals),
                }
            )
        )
        summary = MigrationSummary(
            plan_records=len(parsed.plans),
            actual_records=len(parsed.actuals),
            current_plan_records=sum(
                record.report_month == report_month for record in parsed.plans
            ),
            current_actual_records=sum(
                record.report_month == report_month for record in parsed.actuals
            ),
            months=months,
            rows_to_insert=insert_count,
            rows_unchanged=unchanged_count,
        )
        database_revision = _database_revision(database_snapshot)
        revision = _dry_run_revision(
            source_sha256=source_sha256,
            master_revision=master_revision,
            report_month=report_month,
            plans=parsed.plans,
            actuals=parsed.actuals,
            database_revision=database_revision,
            unknown_aliases=parsed.unknown_aliases,
            issues=issues,
        )
        return MigrationDryRun(
            source_path=source,
            source_sha256=source_sha256,
            master_revision=master_revision,
            database_snapshot=database_snapshot,
            database_revision=database_revision,
            revision=revision,
            report_month=report_month,
            plans=parsed.plans,
            actuals=parsed.actuals,
            summary=summary,
            unknown_aliases=parsed.unknown_aliases,
            issues=issues,
            reconciliation=reconciliation,
        )

    def commit(
        self,
        source_path: Path | str,
        *,
        report_month: str,
        confirmed: bool,
        expected_source_sha256: str,
        expected_master_revision: str,
        expected_database_revision: str,
        expected_revision: str,
    ) -> MigrationCommitResult:
        if confirmed is not True:
            raise LegacyMigrationConfirmationError(
                "마이그레이션 커밋에는 화면의 명시적인 확인이 필요합니다."
            )
        current = self.dry_run(source_path, report_month=report_month)
        if current.source_sha256 != expected_source_sha256:
            raise LegacyMigrationRevisionError(
                "드라이런 이후 원본 파일이 변경되어 파일 해시가 일치하지 않습니다. "
                "다시 드라이런하세요."
            )
        if current.master_revision != expected_master_revision:
            raise LegacyMigrationRevisionError(
                "드라이런 이후 납품처/별칭/그룹 기준정보가 변경되었습니다. "
                "다시 드라이런하세요."
            )
        if current.database_revision != expected_database_revision:
            raise LegacyMigrationRevisionError(
                "드라이런 이후 대상 DB 값 또는 월 잠금 상태가 변경되었습니다. "
                "다시 드라이런하세요."
            )
        if current.revision != expected_revision:
            raise LegacyMigrationRevisionError(
                "드라이런 이후 대상 DB 값 또는 월 잠금 상태가 변경되었습니다. "
                "다시 드라이런하세요."
            )
        if not current.can_commit:
            raise LegacyMigrationBlockedError(
                "드라이런 차이 또는 차단 항목이 남아 있어 커밋할 수 없습니다."
            )

        source_note = (
            f"legacy migration:{current.source_path.name}:"
            f"{current.source_sha256}"
        )
        inserted = 0
        unchanged = 0
        connection = self.database.connect()
        try:
            connection.execute("BEGIN IMMEDIATE")
            if _sha256_file(current.source_path) != current.source_sha256:
                raise LegacyMigrationRevisionError(
                    "커밋 직전에 원본 파일이 변경되었습니다. 다시 드라이런하세요."
                )
            if _master_revision(connection) != current.master_revision:
                raise LegacyMigrationRevisionError(
                    "커밋 직전에 기준정보가 변경되었습니다. 다시 드라이런하세요."
                )
            database_issues, database_snapshot, _, _ = _inspect_database_state(
                connection, current.plans, current.actuals
            )
            database_revision = _database_revision(database_snapshot)
            if database_revision != current.database_revision:
                raise LegacyMigrationRevisionError(
                    "커밋 직전에 대상 DB 값 또는 월 잠금 상태가 변경되었습니다. "
                    "다시 드라이런하세요."
                )
            revision = _dry_run_revision(
                source_sha256=current.source_sha256,
                master_revision=current.master_revision,
                report_month=current.report_month,
                plans=current.plans,
                actuals=current.actuals,
                database_revision=database_revision,
                unknown_aliases=current.unknown_aliases,
                issues=tuple(
                    sorted(
                        (
                            *(
                                issue
                                for issue in current.issues
                                if issue.code not in _DB_ISSUE_CODES
                            ),
                            *database_issues,
                        ),
                        key=lambda item: (
                            item.source_locator,
                            item.code,
                            item.message,
                        ),
                    )
                ),
            )
            if revision != current.revision:
                raise LegacyMigrationRevisionError(
                    "커밋 직전에 대상 DB 값 또는 월 잠금 상태가 변경되었습니다. "
                    "다시 드라이런하세요."
                )

            existing_plans, existing_quantities, existing_costs = (
                _load_existing_monthly_values(
                    connection,
                    sorted(
                        {
                            *(record.report_month for record in current.plans),
                            *(record.report_month for record in current.actuals),
                        }
                    ),
                )
            )
            for record in current.plans:
                key = (record.report_month, record.destination_id)
                existing = existing_plans.get(key)
                if existing is not None:
                    unchanged += 1
                    continue
                connection.execute(
                    "INSERT INTO monthly_plans"
                    "(report_month, destination_id, quantity_ea_text, cost_won, "
                    "source_note) VALUES (?, ?, ?, ?, ?)",
                    (
                        record.report_month,
                        record.destination_id,
                        _quantity_text(record.quantity_ea),
                        record.cost_won,
                        source_note,
                    ),
                )
                inserted += 1
            for record in current.actuals:
                key = (record.report_month, record.destination_id)
                if record.quantity_ea is not None:
                    existing_quantity = existing_quantities.get(key)
                    if existing_quantity is None:
                        connection.execute(
                            "INSERT INTO monthly_actual_quantities"
                            "(report_month, destination_id, quantity_ea_text, "
                            "source_note) VALUES (?, ?, ?, ?)",
                            (
                                record.report_month,
                                record.destination_id,
                                _quantity_text(record.quantity_ea),
                                source_note,
                            ),
                        )
                        inserted += 1
                    else:
                        unchanged += 1
                if record.cost_won is not None:
                    existing_cost = existing_costs.get(key)
                    if existing_cost is None:
                        connection.execute(
                            "INSERT INTO monthly_actual_costs"
                            "(report_month, destination_id, cost_won, source_type, "
                            "source_note) VALUES (?, ?, ?, ?, ?)",
                            (
                                record.report_month,
                                record.destination_id,
                                record.cost_won,
                                LEGACY_SOURCE_TYPE,
                                source_note,
                            ),
                        )
                        inserted += 1
                    else:
                        unchanged += 1
            if _sha256_file(current.source_path) != current.source_sha256:
                raise LegacyMigrationRevisionError(
                    "저장 중 원본 파일이 변경되어 전체 작업을 롤백했습니다. "
                    "다시 드라이런하세요."
                )
            connection.commit()
        except LegacyMigrationError:
            connection.rollback()
            raise
        except Exception as error:
            connection.rollback()
            raise LegacyMigrationError(
                f"마이그레이션 저장에 실패하여 전체 작업을 롤백했습니다: {error}"
            ) from error
        finally:
            connection.close()
        return MigrationCommitResult(
            source_sha256=current.source_sha256,
            revision=current.revision,
            inserted=inserted,
            unchanged=unchanged,
        )


_DB_ISSUE_CODES = frozenset(
    {"EXISTING_VALUE_CONFLICT", "MONTH_LOCKED"}
)


def _parse_source(
    path: Path,
    *,
    report_month: str,
    aliases: dict[str, int],
    initial_issues: tuple[MigrationIssue, ...],
) -> _ParsedSource:
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        missing_sheets = [
            name for name in (_HISTORY_SHEET, _PLAN_SHEET) if name not in cached.sheetnames
        ]
        if missing_sheets:
            empty = _empty_parsed(
                *initial_issues,
                *(
                    MigrationIssue(
                        code="MISSING_SHEET",
                        message=f"필수 시트 '{name}'을 찾을 수 없습니다.",
                        source_locator=name,
                    )
                    for name in missing_sheets
                ),
            )
            return empty
        history = _parse_history(
            formulas[_HISTORY_SHEET], cached[_HISTORY_SHEET], aliases
        )
        plans = _parse_plans(
            formulas[_PLAN_SHEET],
            cached[_PLAN_SHEET],
            aliases,
            report_month,
        )
        return _ParsedSource(
            plans=plans.plans,
            actuals=history.actuals,
            # Preserve source traversal order: history A:F first, then the
            # fixed plan table. This keeps repeated dry-runs deterministic and
            # presents the earliest business-data occurrence first.
            unknown_aliases=tuple((*history.unknown_aliases, *plans.unknown_aliases)),
            issues=tuple((*initial_issues, *history.issues, *plans.issues)),
            controls=plans.controls,
        )
    finally:
        formulas.close()
        cached.close()


def _empty_parsed(*issues: MigrationIssue) -> _ParsedSource:
    return _ParsedSource((), (), (), tuple(issues), ())


def _parse_history(formula_sheet, cached_sheet, aliases: dict[str, int]) -> _ParsedSource:
    issues: list[MigrationIssue] = []
    unknown: list[UnknownAlias] = []
    records: list[StagedActualRecord] = []
    for column, expected in enumerate(_HISTORY_HEADERS, start=1):
        actual = cached_sheet.cell(1, column).value
        if actual != expected:
            issues.append(
                MigrationIssue(
                    code="HEADER_MISMATCH",
                    message=(
                        f"헤더는 '{expected}'이어야 합니다. 원본 양식의 A:F 열을 "
                        "확인하세요."
                    ),
                    source_locator=f"{_HISTORY_SHEET}!{cached_sheet.cell(1, column).coordinate}",
                    expected=expected,
                    actual=None if actual is None else str(actual),
                )
            )

    for row in range(282, cached_sheet.max_row + 1):
        if any(
            cached_sheet.cell(row, column).value is not None
            for column in range(1, 7)
        ):
            issues.append(
                MigrationIssue(
                    code="HISTORY_ROW_OUT_OF_RANGE",
                    message=(
                        "누적 데이터의 업무 영역은 A2:F281(280행)입니다. "
                        "범위 밖 값을 삭제하거나 올바른 행으로 이동하세요."
                    ),
                    source_locator=f"{_HISTORY_SHEET}!A{row}:F{row}",
                )
            )
    seen: set[tuple[str, int]] = set()
    for row in range(2, 282):
        values = tuple(cached_sheet.cell(row, column).value for column in range(1, 7))
        if all(value is None for value in values):
            issues.append(
                MigrationIssue(
                    code="UNEXPECTED_BLANK_HISTORY_ROW",
                    message=(
                        "누적 데이터 A:F 영역 중간에 빈 행이 있습니다. "
                        "행을 삭제하거나 값을 복원하세요."
                    ),
                    source_locator=f"{_HISTORY_SHEET}!A{row}:F{row}",
                )
            )
            continue
        year = _parse_year(values[0], f"{_HISTORY_SHEET}!A{row}", issues)
        month = _parse_month_number(values[1], f"{_HISTORY_SHEET}!B{row}", issues)
        alias = _source_alias(values[3], f"{_HISTORY_SHEET}!D{row}", issues)
        quantity_value = _cached_value(
            formula_sheet, cached_sheet, row, 5, issues
        )
        cost_value = _cached_value(formula_sheet, cached_sheet, row, 6, issues)
        quantity = _parse_optional_quantity(
            quantity_value, f"{_HISTORY_SHEET}!E{row}", issues
        )
        cost = _parse_optional_cost(
            cost_value, f"{_HISTORY_SHEET}!F{row}", issues
        )
        if alias is None:
            continue
        destination_id = aliases.get(_normalize(alias))
        if destination_id is None:
            unknown.append(
                UnknownAlias(alias, LEGACY_SOURCE_TYPE, f"{_HISTORY_SHEET}!D{row}")
            )
            continue
        if year is None or month is None:
            continue
        if quantity is None and cost is None:
            continue
        report_month = f"{year:04d}-{month:02d}"
        key = (report_month, destination_id)
        if key in seen:
            issues.append(
                MigrationIssue(
                    code="DUPLICATE_SOURCE_RECORD",
                    message="같은 월과 납품처의 실적 행이 두 번 있습니다.",
                    source_locator=f"{_HISTORY_SHEET}!D{row}",
                )
            )
            continue
        seen.add(key)
        records.append(
            StagedActualRecord(
                report_month=report_month,
                destination_id=destination_id,
                destination_alias=alias,
                quantity_ea=quantity,
                cost_won=cost,
                source_sheet=_HISTORY_SHEET,
                source_row=row,
            )
        )
    records.sort(key=lambda item: (item.report_month, item.destination_id))
    return _ParsedSource((), tuple(records), tuple(unknown), tuple(issues), ())


def _parse_plans(
    formula_sheet,
    cached_sheet,
    aliases: dict[str, int],
    report_month: str,
) -> _ParsedSource:
    issues: list[MigrationIssue] = []
    unknown: list[UnknownAlias] = []
    records: list[StagedPlanRecord] = []
    controls: list[_SourceControl] = []
    report_year, report_month_number = map(int, report_month.split("-"))
    expected_sheet_year = 2000 + int(_PLAN_SHEET[:2])
    if report_year != expected_sheet_year:
        issues.append(
            MigrationIssue(
                code="PLAN_YEAR_MISMATCH",
                message=(
                    f"'{_PLAN_SHEET}'은 {expected_sheet_year}년 계획 시트이므로 "
                    f"{report_month} 보고월과 일치하지 않습니다."
                ),
                source_locator=_PLAN_SHEET,
            )
        )
    if cached_sheet["B4"].value != "납품처":
        issues.append(
            MigrationIssue(
                code="HEADER_MISMATCH",
                message="계획 시트 B4 헤더는 '납품처'이어야 합니다.",
                source_locator=f"{_PLAN_SHEET}!B4",
                expected="납품처",
                actual=_as_text(cached_sheet["B4"].value),
            )
        )
    for month in range(1, 13):
        quantity_column = 3 + (month - 1) * 2
        for row, column, expected in (
            (4, quantity_column, f"{month}월"),
            (5, quantity_column, _PLAN_QUANTITY_HEADER),
            (5, quantity_column + 1, _PLAN_COST_HEADER),
        ):
            actual = cached_sheet.cell(row, column).value
            if actual != expected:
                issues.append(
                    MigrationIssue(
                        code="HEADER_MISMATCH",
                        message=f"계획 시트 헤더는 '{expected}'이어야 합니다.",
                        source_locator=(
                            f"{_PLAN_SHEET}!{cached_sheet.cell(row, column).coordinate}"
                        ),
                        expected=expected,
                        actual=_as_text(actual),
                    )
                )

    total_row = 20
    total_label = _as_text(cached_sheet.cell(total_row, 2).value)
    if total_label != "합계":
        issues.append(
            MigrationIssue(
                code="PLAN_TOTAL_ROW_MISMATCH",
                message=(
                    "계획 데이터는 B6:B19 납품처 행과 B20 합계 행 구조여야 "
                    "합니다. B20의 '합계'를 복원하세요."
                ),
                source_locator=f"{_PLAN_SHEET}!B20",
                expected="합계",
                actual=total_label,
            )
        )

    seen: set[tuple[str, int]] = set()
    for row in range(6, total_row):
        alias = _source_alias(
            cached_sheet.cell(row, 2).value,
            f"{_PLAN_SHEET}!B{row}",
            issues,
        )
        if alias is None:
            continue
        destination_id = aliases.get(_normalize(alias))
        if destination_id is None:
            unknown.append(
                UnknownAlias(alias, LEGACY_SOURCE_TYPE, f"{_PLAN_SHEET}!B{row}")
            )
            continue
        for month in (report_month_number,):
            quantity_column = 3 + (month - 1) * 2
            cost_column = quantity_column + 1
            quantity_value = _cached_value(
                formula_sheet,
                cached_sheet,
                row,
                quantity_column,
                issues,
            )
            cost_value = _cached_value(
                formula_sheet,
                cached_sheet,
                row,
                cost_column,
                issues,
            )
            if quantity_value is None and cost_value is None:
                continue
            if quantity_value is None or cost_value is None:
                issues.append(
                    MigrationIssue(
                        code="PARTIAL_PLAN_VALUE",
                        message=(
                            "계획 수량과 운반비 중 하나만 있습니다. 두 값을 모두 "
                            "입력하거나 모두 비워야 합니다."
                        ),
                        source_locator=(
                            f"{_PLAN_SHEET}!{cached_sheet.cell(row, quantity_column).coordinate}:"
                            f"{cached_sheet.cell(row, cost_column).coordinate}"
                        ),
                    )
                )
                continue
            quantity = _parse_quantity(
                quantity_value,
                f"{_PLAN_SHEET}!{cached_sheet.cell(row, quantity_column).coordinate}",
                issues,
            )
            cost = _parse_cost(
                cost_value,
                f"{_PLAN_SHEET}!{cached_sheet.cell(row, cost_column).coordinate}",
                issues,
            )
            if quantity is None or cost is None:
                continue
            month_key = f"{report_year:04d}-{month:02d}"
            key = (month_key, destination_id)
            if key in seen:
                issues.append(
                    MigrationIssue(
                        code="DUPLICATE_SOURCE_RECORD",
                        message="같은 월과 납품처의 계획 행이 두 번 있습니다.",
                        source_locator=f"{_PLAN_SHEET}!B{row}",
                    )
                )
                continue
            seen.add(key)
            records.append(
                StagedPlanRecord(
                    report_month=month_key,
                    destination_id=destination_id,
                    destination_alias=alias,
                    quantity_ea=quantity,
                    cost_won=cost,
                    source_sheet=_PLAN_SHEET,
                    source_row=row,
                )
            )

    if total_row <= cached_sheet.max_row:
        current_quantity_column = 3 + (report_month_number - 1) * 2
        current_cost_column = current_quantity_column + 1
        quantity_value = _cached_value(
            formula_sheet,
            cached_sheet,
            total_row,
            current_quantity_column,
            issues,
        )
        cost_value = _cached_value(
            formula_sheet,
            cached_sheet,
            total_row,
            current_cost_column,
            issues,
        )
        controls.extend(
            (
                _SourceControl(
                    "CURRENT_PLAN_QUANTITY_TOTAL",
                    _cell_locator(
                        cached_sheet, total_row, current_quantity_column
                    ),
                    _parse_quantity(
                        quantity_value,
                        _cell_locator(
                            cached_sheet, total_row, current_quantity_column
                        ),
                        issues,
                    )
                    if quantity_value is not None
                    else None,
                ),
                _SourceControl(
                    "CURRENT_PLAN_COST_TOTAL",
                    _cell_locator(cached_sheet, total_row, current_cost_column),
                    _parse_cost(
                        cost_value,
                        _cell_locator(
                            cached_sheet, total_row, current_cost_column
                        ),
                        issues,
                    )
                    if cost_value is not None
                    else None,
                ),
            )
        )
    records.sort(key=lambda item: (item.report_month, item.destination_id))
    return _ParsedSource(tuple(records), (), tuple(unknown), tuple(issues), tuple(controls))


def _cached_value(formula_sheet, cached_sheet, row, column, issues):
    formula = formula_sheet.cell(row, column).value
    cached = cached_sheet.cell(row, column).value
    if isinstance(formula, str) and formula.startswith("=") and cached is None:
        coordinate = cached_sheet.cell(row, column).coordinate
        issues.append(
            MigrationIssue(
                code="MISSING_FORMULA_CACHE",
                message=(
                    "수식의 저장값(캐시)이 없습니다. Excel에서 원본을 계산 후 "
                    "저장하고 다시 드라이런하세요."
                ),
                source_locator=f"{cached_sheet.title}!{coordinate}",
            )
        )
    return cached


def _parse_year(value, locator, issues) -> int | None:
    if isinstance(value, bool):
        year = None
    elif isinstance(value, int):
        year = value if 2000 <= value <= 9999 else None
    elif isinstance(value, str):
        match = re.fullmatch(r"\s*([0-9]{2}|[0-9]{4})년\s*", value)
        year = None if match is None else int(match.group(1))
        if year is not None and year < 100:
            year += 2000
    else:
        year = None
    if year is None:
        issues.append(
            MigrationIssue(
                code="INVALID_YEAR",
                message="연도는 '25년' 또는 2025 형식이어야 합니다.",
                source_locator=locator,
                actual=_as_text(value),
            )
        )
    return year


def _parse_month_number(value, locator, issues) -> int | None:
    if isinstance(value, bool):
        month = None
    elif isinstance(value, int):
        month = value
    elif isinstance(value, str):
        match = re.fullmatch(r"\s*([0-9]{1,2})월\s*", value)
        month = None if match is None else int(match.group(1))
    else:
        month = None
    if month is None or not 1 <= month <= 12:
        issues.append(
            MigrationIssue(
                code="INVALID_MONTH",
                message="월은 '1월'부터 '12월' 형식이어야 합니다.",
                source_locator=locator,
                actual=_as_text(value),
            )
        )
        return None
    return month


def _source_alias(value, locator, issues) -> str | None:
    if not isinstance(value, str) or not value.strip():
        issues.append(
            MigrationIssue(
                code="INVALID_DESTINATION",
                message="납품처 이름이 비어 있습니다.",
                source_locator=locator,
                actual=_as_text(value),
            )
        )
        return None
    return unicodedata.normalize("NFKC", value).strip()


def _parse_quantity(value, locator, issues) -> Decimal | None:
    quantity: Decimal | None
    if isinstance(value, bool) or value is None:
        quantity = None
    else:
        try:
            quantity = Decimal(str(value))
        except (InvalidOperation, ValueError):
            quantity = None
    valid = (
        quantity is not None
        and quantity.is_finite()
        and Decimal(0) <= quantity <= MAX_QUANTITY_EA
        and _decimal_places(quantity) <= MAX_QUANTITY_SCALE
    )
    if not valid:
        issues.append(
            MigrationIssue(
                code="INVALID_QUANTITY",
                message=(
                    f"수량은 0 이상, 소수 {MAX_QUANTITY_SCALE}자리 이하의 숫자여야 합니다."
                ),
                source_locator=locator,
                actual=_as_text(value),
            )
        )
        return None
    return quantity


def _parse_optional_quantity(value, locator, issues) -> Decimal | None:
    if value is None:
        return None
    return _parse_quantity(value, locator, issues)


def _parse_cost(value, locator, issues) -> int | None:
    cost: int | None = None
    if not isinstance(value, bool):
        if isinstance(value, int):
            cost = value
        elif (
            isinstance(value, Decimal)
            and value.is_finite()
            and value == value.to_integral_value()
        ):
            cost = int(value)
        elif isinstance(value, float) and math.isfinite(value) and value.is_integer():
            cost = int(value)
    if cost is None or not 0 <= cost <= MAX_SQLITE_INTEGER:
        issues.append(
            MigrationIssue(
                code="INVALID_COST",
                message="운반비는 0 이상의 원 단위 정수여야 합니다.",
                source_locator=locator,
                actual=_as_text(value),
            )
        )
        return None
    return cost


def _parse_optional_cost(value, locator, issues) -> int | None:
    if value is None:
        return None
    return _parse_cost(value, locator, issues)


def _load_destinations(connection: sqlite3.Connection) -> tuple[Destination, ...]:
    rows = connection.execute(
        "SELECT * FROM destinations ORDER BY display_order, id"
    ).fetchall()
    return tuple(
        Destination(
            id=int(row["id"]),
            name=str(row["name"]),
            display_order=int(row["display_order"]),
            active=bool(row["active"]),
            required_for_report=bool(row["required_for_report"]),
            representative_item=row["representative_item"],
            include_quantity_total=bool(row["include_quantity_total"]),
            include_cost_total=bool(row["include_cost_total"]),
            include_sales_total=bool(row["include_sales_total"]),
        )
        for row in rows
    )


@contextmanager
def _read_only_connection(database: Database) -> Iterator[sqlite3.Connection]:
    """Open SQLite in URI read-only mode so dry-run cannot mutate the DB."""
    uri = f"{database.path.resolve().as_uri()}?mode=ro"
    connection = sqlite3.connect(uri, uri=True)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA query_only = ON")
    connection.execute("BEGIN")
    try:
        yield connection
    finally:
        connection.close()


def _load_resolution(connection):
    mapping: dict[str, int] = {}
    issues: list[MigrationIssue] = []
    rows = connection.execute(
        "SELECT raw_name, destination_id FROM destination_aliases "
        "WHERE source_type = ? ORDER BY raw_name, id",
        (LEGACY_SOURCE_TYPE,),
    ).fetchall()
    for row in rows:
        _add_resolution(
            mapping,
            issues,
            str(row["raw_name"]),
            int(row["destination_id"]),
            "destination_aliases",
        )
    return mapping, tuple(issues)


def _add_resolution(mapping, issues, raw_name, destination_id, locator):
    key = _normalize(raw_name)
    existing = mapping.get(key)
    if existing is not None and existing != destination_id:
        issues.append(
            MigrationIssue(
                code="AMBIGUOUS_MASTER_ALIAS",
                message=(
                    f"정규화한 납품처 이름 '{raw_name}'이 둘 이상의 기준정보와 겹칩니다."
                ),
                source_locator=locator,
            )
        )
        return
    mapping[key] = destination_id


def _load_groups(connection):
    group_rows = connection.execute(
        "SELECT id, name FROM report_groups WHERE active = 1 "
        "ORDER BY display_order, id"
    ).fetchall()
    groups = []
    for group in group_rows:
        member_rows = connection.execute(
            "SELECT gm.*, d.name FROM report_group_members AS gm "
            "JOIN destinations AS d ON d.id = gm.destination_id "
            "WHERE gm.group_id = ? ORDER BY gm.display_order, gm.destination_id",
            (group["id"],),
        ).fetchall()
        members = tuple(
            GroupMember(
                group_id=int(row["group_id"]),
                destination_id=int(row["destination_id"]),
                name=str(row["name"]),
                display_order=int(row["display_order"]),
                include_quantity=bool(row["include_quantity"]),
                include_cost=bool(row["include_cost"]),
            )
            for row in member_rows
        )
        groups.append((int(group["id"]), str(group["name"]), members))
    return tuple(groups)


def _reconcile(
    *, report_month,
    plans,
    actuals,
    destinations,
    groups,
    source_controls,
):
    current_plans = {
        record.destination_id: record
        for record in plans
        if record.report_month == report_month
    }
    current_actuals = {
        record.destination_id: record
        for record in actuals
        if record.report_month == report_month
    }
    plan_calculations = {}
    quantity_calculations = {}
    cost_calculations = {}
    for destination in destinations:
        plan = current_plans.get(destination.id)
        actual = current_actuals.get(destination.id)
        plan_calculations[destination.id] = calculate_destination(
            None if plan is None else plan.quantity_ea,
            None if plan is None else plan.cost_won,
            None,
            None,
            destination_id=destination.id,
        )
        quantity = None if actual is None else actual.quantity_ea
        cost = None if actual is None else actual.cost_won
        quantity_calculations[destination.id] = calculate_destination(
            None,
            None,
            quantity,
            None if quantity is None else 0,
            destination_id=destination.id,
        )
        cost_calculations[destination.id] = calculate_destination(
            None,
            None,
            None if cost is None else Decimal(0),
            cost,
            destination_id=destination.id,
        )
    plan_total = calculate_total(plan_calculations, destinations)
    quantity_total = calculate_total(quantity_calculations, destinations)
    cost_total = calculate_total(cost_calculations, destinations)
    total = ReconciledMetrics(
        planned_quantity=plan_total.planned_quantity,
        planned_cost_won=plan_total.planned_cost_won,
        actual_quantity=quantity_total.actual_quantity,
        actual_cost_won=cost_total.actual_cost_won,
    )
    group_results_list = []
    for group_id, name, members in groups:
        plan_group = calculate_group(
            plan_calculations, members, group_id=group_id
        )
        quantity_group = calculate_group(
            quantity_calculations, members, group_id=group_id
        )
        cost_group = calculate_group(
            cost_calculations, members, group_id=group_id
        )
        group_results_list.append(
            GroupReconciliation(
                group_id=group_id,
                name=name,
                result=ReconciledMetrics(
                    planned_quantity=plan_group.planned_quantity,
                    planned_cost_won=plan_group.planned_cost_won,
                    actual_quantity=quantity_group.actual_quantity,
                    actual_cost_won=cost_group.actual_cost_won,
                ),
            )
        )
    group_results = tuple(group_results_list)

    actual_by_month = {}
    for record in actuals:
        actual_by_month.setdefault(record.report_month, {})[record.destination_id] = record
    quantity_values = {}
    cost_values = {}
    for month, month_records in actual_by_month.items():
        month_quantity_calculations = {}
        month_cost_calculations = {}
        for destination in destinations:
            record = month_records.get(destination.id)
            quantity = None if record is None else record.quantity_ea
            cost = None if record is None else record.cost_won
            month_quantity_calculations[destination.id] = calculate_destination(
                None,
                None,
                quantity,
                None if quantity is None else 0,
                destination_id=destination.id,
            )
            month_cost_calculations[destination.id] = calculate_destination(
                None,
                None,
                None if cost is None else Decimal(0),
                cost,
                destination_id=destination.id,
            )
        quantity_values[month] = calculate_total(
            month_quantity_calculations, destinations
        ).actual_quantity
        cost_values[month] = calculate_total(
            month_cost_calculations, destinations
        ).actual_cost_won
    quantity_history = historical_averages(report_month, quantity_values)
    cost_history = historical_averages(
        report_month, cost_values, value_kind=HistoricalValueKind.MONEY
    )

    metric_by_name = {
        "CURRENT_PLAN_QUANTITY_TOTAL": total.planned_quantity,
        "CURRENT_PLAN_COST_TOTAL": total.planned_cost_won,
    }
    controls = []
    issues = []
    current_metrics = (
        ("계획 수량", total.planned_quantity),
        ("계획 운반비", total.planned_cost_won),
        ("실적 수량", total.actual_quantity),
        ("실적 운반비", total.actual_cost_won),
    )
    for label, value in current_metrics:
        if value is None:
            issues.append(
                MigrationIssue(
                    code="INCOMPLETE_CURRENT_TOTAL",
                    message=(
                        f"{report_month} {label} 합계에 필요한 납품처 값이 없습니다. "
                        "누락과 0을 구분해 원본 또는 포함 규칙을 확인하세요."
                    ),
                    source_locator=f"{report_month} 재검증:{label}",
                )
            )
    for group in group_results:
        missing_metrics = [
            label
            for label, value in (
                ("계획 수량", group.result.planned_quantity),
                ("계획 운반비", group.result.planned_cost_won),
                ("실적 수량", group.result.actual_quantity),
                ("실적 운반비", group.result.actual_cost_won),
            )
            if value is None
        ]
        if missing_metrics:
            issues.append(
                MigrationIssue(
                    code="INCOMPLETE_GROUP_TOTAL",
                    message=(
                        f"그룹 '{group.name}'의 {', '.join(missing_metrics)} 계산에 "
                        "필요한 납품처 값이 없습니다."
                    ),
                    source_locator=f"report_groups:{group.group_id}",
                )
            )
    for value_label, averages in (
        ("실적 수량", quantity_history),
        ("실적 운반비", cost_history),
    ):
        for period_label, average in (
            ("3개월", averages.three_month),
            ("6개월", averages.six_month),
            ("12개월", averages.twelve_month),
            ("전년도", averages.comparison_year),
        ):
            if not average.complete:
                issues.append(
                    MigrationIssue(
                        code="INCOMPLETE_HISTORY",
                        message=(
                            f"{value_label} {period_label} 평균에 필요한 월이 없습니다: "
                            f"{', '.join(average.missing_months)}."
                        ),
                        source_locator=f"누적 데이터:{value_label}:{period_label}",
                    )
                )
    for control in source_controls:
        actual = metric_by_name[control.name]
        matches = control.value is not None and actual is not None and control.value == actual
        controls.append(
            ReconciliationControl(
                name=control.name,
                source_locator=control.locator,
                expected=control.value,
                actual=actual,
                matches=matches,
            )
        )
        if not matches:
            issues.append(
                MigrationIssue(
                    code="RECONCILIATION_MISMATCH",
                    message=(
                        "원본 합계와 납품처 포함 규칙으로 다시 계산한 합계가 다릅니다. "
                        "기준정보와 원본을 확인하세요."
                    ),
                    source_locator=control.locator,
                    expected=_as_text(control.value),
                    actual=_as_text(actual),
                )
            )
    return (
        MigrationReconciliation(
            current_total=total,
            groups=group_results,
            actual_quantity_history=quantity_history,
            actual_cost_history=cost_history,
            controls=tuple(controls),
        ),
        tuple(issues),
    )


def _inspect_database_state(connection, plans, actuals):
    issues: list[MigrationIssue] = []
    targets: list[TargetRowSnapshot] = []
    insert_count = 0
    unchanged_count = 0
    months = sorted(
        {
            *(record.report_month for record in plans),
            *(record.report_month for record in actuals),
        }
    )
    locked = {
        str(row["report_month"])
        for row in connection.execute(
            "SELECT report_month FROM month_locks WHERE is_locked = 1"
        )
        if str(row["report_month"]) in months
    }
    existing_plans, existing_quantities, existing_costs = (
        _load_existing_monthly_values(connection, months)
    )
    for record in plans:
        key = (record.report_month, record.destination_id)
        existing = existing_plans.get(key)
        existing_text = (
            None if existing is None else (str(existing[0]), str(existing[1]))
        )
        incoming = (_quantity_text(record.quantity_ea), str(record.cost_won))
        action = _target_action(existing_text, incoming)
        targets.append(
            TargetRowSnapshot(
                table="monthly_plans",
                report_month=record.report_month,
                destination_id=record.destination_id,
                existing_value=existing_text,
                incoming_value=incoming,
                action=action,
            )
        )
        if action == "insert":
            insert_count += 1
        elif action == "unchanged":
            unchanged_count += 1
        else:
            issues.append(
                _conflict_issue(
                    "monthly_plans",
                    record.report_month,
                    record.destination_id,
                    existing_text,
                    incoming,
                )
            )
    for record in actuals:
        if record.quantity_ea is not None:
            key = (record.report_month, record.destination_id)
            existing_quantity = existing_quantities.get(key)
            existing_text = (
                None if existing_quantity is None else (existing_quantity,)
            )
            incoming = (_quantity_text(record.quantity_ea),)
            action = _target_action(existing_text, incoming)
            targets.append(
                TargetRowSnapshot(
                    table="monthly_actual_quantities",
                    report_month=record.report_month,
                    destination_id=record.destination_id,
                    existing_value=existing_text,
                    incoming_value=incoming,
                    action=action,
                )
            )
            if action == "insert":
                insert_count += 1
            elif action == "unchanged":
                unchanged_count += 1
            else:
                issues.append(
                    _conflict_issue(
                        "monthly_actual_quantities",
                        record.report_month,
                        record.destination_id,
                        existing_text,
                        incoming,
                    )
                )
        if record.cost_won is not None:
            key = (record.report_month, record.destination_id)
            existing_cost = existing_costs.get(key)
            existing_text = None if existing_cost is None else (str(existing_cost),)
            incoming = (str(record.cost_won),)
            action = _target_action(existing_text, incoming)
            targets.append(
                TargetRowSnapshot(
                    table="monthly_actual_costs",
                    report_month=record.report_month,
                    destination_id=record.destination_id,
                    existing_value=existing_text,
                    incoming_value=incoming,
                    action=action,
                )
            )
            if action == "insert":
                insert_count += 1
            elif action == "unchanged":
                unchanged_count += 1
            else:
                issues.append(
                    _conflict_issue(
                        "monthly_actual_costs",
                        record.report_month,
                        record.destination_id,
                        existing_text,
                        incoming,
                    )
                )
    lock_snapshots = tuple(
        MonthLockSnapshot(report_month=month, is_locked=month in locked)
        for month in months
    )
    for lock in lock_snapshots:
        if lock.is_locked and any(
            target.report_month == lock.report_month
            and target.action != "unchanged"
            for target in targets
        ):
            issues.append(
                MigrationIssue(
                    code="MONTH_LOCKED",
                    message=(
                        f"{lock.report_month}은 마감 잠금 상태이며 새로 저장하거나 "
                        "변경할 값이 있어 마이그레이션할 수 없습니다."
                    ),
                    source_locator=f"DB month_locks:{lock.report_month}",
                )
            )
    snapshot = DatabaseSnapshot(lock_snapshots, tuple(targets))
    return tuple(issues), snapshot, insert_count, unchanged_count


def _target_action(existing, incoming) -> str:
    if existing is None:
        return "insert"
    return "unchanged" if existing == incoming else "conflict"


def _load_existing_monthly_values(connection, months):
    if not months:
        return {}, {}, {}
    placeholders = ",".join("?" for _ in months)
    plans = {
        (str(row["report_month"]), int(row["destination_id"])): (
            str(row["quantity_ea_text"]),
            int(row["cost_won"]),
        )
        for row in connection.execute(
            "SELECT report_month, destination_id, quantity_ea_text, cost_won "
            f"FROM monthly_plans WHERE report_month IN ({placeholders})",
            months,
        )
    }
    quantities = {
        (str(row["report_month"]), int(row["destination_id"])): str(
            row["quantity_ea_text"]
        )
        for row in connection.execute(
            "SELECT report_month, destination_id, quantity_ea_text "
            f"FROM monthly_actual_quantities WHERE report_month IN ({placeholders})",
            months,
        )
    }
    costs = {
        (str(row["report_month"]), int(row["destination_id"])): int(
            row["cost_won"]
        )
        for row in connection.execute(
            "SELECT report_month, destination_id, cost_won "
            f"FROM monthly_actual_costs WHERE report_month IN ({placeholders})",
            months,
        )
    }
    return plans, quantities, costs


def _conflict_issue(table, month, destination_id, existing, incoming):
    return MigrationIssue(
        code="EXISTING_VALUE_CONFLICT",
        message=(
            "기존 월별 값과 마이그레이션 값에 차이가 있습니다. 기존 값을 "
            "덮어쓰지 않았습니다."
        ),
        source_locator=f"DB {table}:{month}:destination={destination_id}",
        expected=_as_text(existing),
        actual=_as_text(incoming),
    )


def _master_revision(connection):
    payload = {}
    for table, columns, order in (
        (
            "destinations",
            "id,name,display_order,active,required_for_report,representative_item,"
            "include_quantity_total,include_cost_total,include_sales_total",
            "id",
        ),
        (
            "destination_aliases",
            "id,raw_name,source_type,destination_id",
            "id",
        ),
        ("report_groups", "id,name,display_order,active", "id"),
        (
            "report_group_members",
            "group_id,destination_id,display_order,include_quantity,include_cost",
            "group_id,destination_id",
        ),
    ):
        payload[table] = [
            tuple(row)
            for row in connection.execute(
                f"SELECT {columns} FROM {table} ORDER BY {order}"
            )
        ]
    return _json_hash(payload)


def _database_revision(snapshot: DatabaseSnapshot) -> str:
    return _json_hash(
        {
            "month_locks": [
                (item.report_month, item.is_locked)
                for item in snapshot.month_locks
            ],
            "targets": [
                (
                    item.table,
                    item.report_month,
                    item.destination_id,
                    item.existing_value,
                )
                for item in snapshot.targets
            ],
        }
    )


def _dry_run_revision(
    *,
    source_sha256,
    master_revision,
    report_month,
    plans,
    actuals,
    database_revision,
    unknown_aliases,
    issues,
):
    payload = {
        "source_sha256": source_sha256,
        "master_revision": master_revision,
        "report_month": report_month,
        "plans": [
            (
                item.report_month,
                item.destination_id,
                _quantity_text(item.quantity_ea),
                item.cost_won,
                item.source_sheet,
                item.source_row,
            )
            for item in plans
        ],
        "actuals": [
            (
                item.report_month,
                item.destination_id,
                None
                if item.quantity_ea is None
                else _quantity_text(item.quantity_ea),
                item.cost_won,
                item.source_sheet,
                item.source_row,
            )
            for item in actuals
        ],
        "database_revision": database_revision,
        "unknown_aliases": [
            (item.alias, item.source_type, item.source_locator)
            for item in unknown_aliases
        ],
        "issues": [
            (
                item.code,
                item.source_locator,
                item.expected,
                item.actual,
                item.blocking,
            )
            for item in issues
        ],
    }
    return _json_hash(payload)


def _json_hash(value) -> str:
    encoded = json.dumps(
        value,
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _quantity_text(value: Decimal) -> str:
    if value == 0:
        return "0"
    text = format(value, "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def _decimal_places(value: Decimal) -> int:
    if value == 0 or value.as_tuple().exponent >= 0:
        return 0
    trailing_zeroes = 0
    for digit in reversed(value.as_tuple().digits):
        if digit:
            break
        trailing_zeroes += 1
    return max(0, -value.as_tuple().exponent - trailing_zeroes)


def _normalize(value: str) -> str:
    return unicodedata.normalize("NFKC", value).strip()


def _cell_locator(sheet, row: int, column: int) -> str:
    return f"{sheet.title}!{get_column_letter(column)}{row}"


def _as_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, tuple):
        return " | ".join(str(item) for item in value)
    return str(value)


def _require_report_month(value: str) -> None:
    if not isinstance(value, str) or _MONTH.fullmatch(value) is None:
        raise LegacyMigrationError("보고월은 ASCII YYYY-MM 형식이어야 합니다.")
