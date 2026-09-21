"""Export an auditable, read-only review workbook for one monthly report.

The public adapter combines the canonical :class:`PptReport` produced for the
fixed report with validation, import provenance, and master snapshots.  The
exporter writes the already-calculated domain values; it does not reproduce
transport-cost arithmetic in spreadsheet formulas.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from functools import lru_cache
import math
import os
from pathlib import Path
import re
import tempfile
from typing import Literal
import unicodedata
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.worksheet import Worksheet

try:
    from PIL import ImageFont
except ImportError:  # pragma: no cover - conservative metrics remain available
    ImageFont = None

from app.domain.calculations import (
    CalculationKind,
    DestinationCalculation,
    HistoricalValueKind,
    calculate_group,
    calculate_total,
    historical_averages,
)
from app.domain.models import Destination, DestinationAlias, GroupMember, ReportGroup
from app.domain.validation import ValidationContext, ValidationResult, validate_report
from app.reporting.pptx_report import PptReport, validate_ppt_report


SHEET_NAMES = ("월간 종합", "운행실적", "검증 결과", "마스터 기준")
FONT_NAME = "Malgun Gothic"
NUMBER_FORMAT = "#,##0.####"
MONEY_FORMAT = "#,##0"
UNIT_COST_FORMAT = "#,##0.00"
PERCENT_FORMAT = "0.0%"
DATE_FORMAT = "yyyy-mm-dd"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"
ERROR_TOKENS = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
    "#SPILL!",
    "#CALC!",
}
_SHA256 = re.compile(r"^[0-9a-fA-F]{64}$")
_ABSOLUTE_PATH_START = re.compile(
    r"(?i)(?<![\w/\\])(?:[A-Z]:[\\/]|\\\\(?=[^\\/\s])|//(?=[^/\s])|/(?=[^/\s]))"
)
_WEB_URL_TOKEN = re.compile(
    r"(?i)(?<![\w+])(?:https?|ftps?)://[^\s\"'<>]+"
)
_LOCATOR_SUFFIX = re.compile(r"(![A-Za-z0-9_$:.\-]+)$")
_PATH_BOUNDARY = re.compile(r"[\r\n,;\)\]\}>]")
_REDACTED_PATH = "경로 숨김"
_VALIDATION_FONT_SIZE = 9
_VALIDATION_MIN_WIDTH = 59
# Wider than ordinary tables but still reviewable without Excel's 255-unit extreme.
_VALIDATION_MAX_WIDTH = 80
_VALIDATION_MIN_ROW_HEIGHT = 36
_VALIDATION_MAX_ROW_HEIGHT = 120
_VALIDATION_CELL_MARGIN_PIXELS = 12
_SCREEN_DPI = 96

_NAVY = "17365D"
_BLUE = "D9EAF7"
_YELLOW = "FFF2CC"
_CALCULATED = "E7E6E6"
_AMBER = "FCE4D6"
_RED = "F4CCCC"
_LIGHT_BLUE = "DDEBF7"
_WHITE = "FFFFFF"
_DARK_TEXT = "1F1F1F"
_RED_TEXT = "9C0006"
_AMBER_TEXT = "9C5700"
_THIN_GREY = Side(style="thin", color="B7C9D6")
_BOTTOM_BORDER = Border(bottom=_THIN_GREY)


@dataclass(frozen=True, slots=True)
class ImportBatchEvidence:
    """One persisted source batch used by the canonical report."""

    batch_id: int
    provenance_id: str
    report_month: str
    source_type: str
    source_filename: str
    file_sha256: str
    imported_at: datetime

    def __post_init__(self) -> None:
        _positive_int(self.batch_id, "batch_id")
        _text(self.provenance_id, "provenance_id")
        _month(self.report_month, "report_month")
        _text(self.source_type, "source_type")
        _text(self.source_filename, "source_filename")
        if not isinstance(self.file_sha256, str) or not _SHA256.fullmatch(
            self.file_sha256
        ):
            raise ValueError("file_sha256 must be 64 hexadecimal characters")
        object.__setattr__(self, "file_sha256", self.file_sha256.lower())
        if not isinstance(self.imported_at, datetime):
            raise TypeError("imported_at must be a datetime")
        if self.imported_at.tzinfo is not None:
            raise ValueError("imported_at must be a naive local datetime")


@dataclass(frozen=True, slots=True)
class OperationEvidence:
    """One imported or manually entered value supporting the report."""

    record_kind: Literal["destination", "nonregular", "sales"]
    value_role: Literal["plan", "actual"]
    input_kind: Literal["imported", "manual", "derived"]
    batch_id: int | None
    provenance_id: str
    report_month: str
    raw_destination: str | None
    destination_id: int | None
    normalized_destination: str
    quantity_ea: Decimal | None
    cost_won: int | None
    source_locator: str

    def __post_init__(self) -> None:
        if self.record_kind not in {"destination", "nonregular", "sales"}:
            raise ValueError("record_kind must be destination, nonregular, or sales")
        if self.value_role not in {"plan", "actual"}:
            raise ValueError("value_role must be plan or actual")
        if self.input_kind not in {"imported", "manual", "derived"}:
            raise ValueError("input_kind must be imported, manual, or derived")
        _text(self.source_locator, "source_locator")
        if self.input_kind == "imported":
            _positive_int(self.batch_id, "batch_id")
            if self.record_kind != "sales":
                _text(self.raw_destination, "raw_destination")
            if "!" not in self.source_locator:
                raise ValueError("imported evidence requires a workbook source locator")
        else:
            if self.batch_id is not None:
                raise ValueError("manual evidence cannot have a batch_id")
            if self.raw_destination is not None:
                raise ValueError("manual evidence cannot have a raw destination")
            if _safe_display_text(self.source_locator) != self.source_locator:
                raise ValueError("manual evidence source locator cannot expose a path")
        _text(self.provenance_id, "provenance_id")
        _month(self.report_month, "report_month")
        if self.record_kind == "destination":
            _positive_int(self.destination_id, "destination_id")
        elif self.destination_id is not None:
            raise ValueError("summary evidence cannot have a destination_id")
        _text(self.normalized_destination, "normalized_destination")
        if self.quantity_ea is not None:
            if not isinstance(self.quantity_ea, Decimal) or not self.quantity_ea.is_finite():
                raise TypeError("quantity_ea must be a finite Decimal or None")
            if self.quantity_ea < 0:
                raise ValueError("quantity_ea must be nonnegative")
        if self.cost_won is not None:
            _nonnegative_int(self.cost_won, "cost_won")
        if self.record_kind == "sales" and self.quantity_ea is not None:
            raise ValueError("sales evidence cannot have a quantity")


@dataclass(frozen=True, slots=True)
class ReviewWorkbookData:
    """Canonical report plus review-only evidence and master snapshots."""

    report: PptReport
    validation_context: ValidationContext
    destinations: tuple[Destination, ...]
    aliases: tuple[DestinationAlias, ...]
    groups: tuple[ReportGroup, ...]
    group_members: tuple[GroupMember, ...]
    import_batches: tuple[ImportBatchEvidence, ...]
    operations: tuple[OperationEvidence, ...]
    validation_result: ValidationResult | None = None

    def __post_init__(self) -> None:
        if not isinstance(self.report, PptReport):
            raise TypeError("report must be a PptReport")
        if not isinstance(self.validation_context, ValidationContext):
            raise TypeError("validation_context must be a ValidationContext")
        if self.validation_result is not None and not isinstance(
            self.validation_result, ValidationResult
        ):
            raise TypeError("validation_result must be a ValidationResult or None")
        _tuple_of(self.destinations, Destination, "destinations")
        _tuple_of(self.aliases, DestinationAlias, "aliases")
        _tuple_of(self.groups, ReportGroup, "groups")
        _tuple_of(self.group_members, GroupMember, "group_members")
        _tuple_of(self.import_batches, ImportBatchEvidence, "import_batches")
        _tuple_of(self.operations, OperationEvidence, "operations")

    @property
    def validation(self) -> ValidationResult:
        """Return Task7 validation recomputed from the immutable input snapshot."""
        result = validate_report(self.validation_context)
        if self.validation_result is not None and self.validation_result != result:
            raise ValueError("supplied validation result conflicts with recomputed result")
        return result


# Backward-compatible type name; the constructor now requires authoritative
# Task7 inputs instead of accepting a free-standing ValidationResult.
ReviewWorkbookReport = ReviewWorkbookData


def export_review_workbook(
    report: ReviewWorkbookData, output: str | Path
) -> None:
    """Write four review sheets atomically and validate the saved XLSX.

    ``output`` is never replaced until a sibling temporary workbook has been
    reopened and checked.  Any failure removes the temporary file and leaves a
    pre-existing output untouched.
    """

    if not isinstance(report, ReviewWorkbookData):
        raise TypeError("report must be a ReviewWorkbookData")
    output = Path(output)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("output must use the .xlsx extension")
    _validate_report_identity(report)
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = _build_workbook(report)
    handle = tempfile.NamedTemporaryFile(
        mode="wb",
        prefix=f".{output.stem}.",
        suffix=".tmp.xlsx",
        dir=output.parent,
        delete=False,
    )
    temporary = Path(handle.name)
    handle.close()
    try:
        workbook.save(temporary)
        workbook.close()
        _validate_saved_workbook(temporary, report)
        os.replace(temporary, output)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)


def _build_workbook(bundle: ReviewWorkbookReport) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = workbook.create_sheet(SHEET_NAMES[0])
    evidence = workbook.create_sheet(SHEET_NAMES[1])
    validation = workbook.create_sheet(SHEET_NAMES[2])
    masters = workbook.create_sheet(SHEET_NAMES[3])

    _write_summary(summary, bundle)
    _write_evidence(evidence, bundle)
    _write_validation(validation, bundle)
    _write_masters(masters, bundle)
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        _apply_font_to_used_range(sheet)
    summary.sheet_properties.tabColor = _NAVY
    evidence.sheet_properties.tabColor = "5B9BD5"
    validation.sheet_properties.tabColor = "A5A5A5"
    masters.sheet_properties.tabColor = "7F8C8D"
    return workbook


def _write_summary(sheet: Worksheet, bundle: ReviewWorkbookReport) -> None:
    report = bundle.report
    _title(sheet, "2026년 8월 운반비 월간 종합")
    sheet["A3"] = "보고 월"
    sheet["B3"] = date(2026, 8, 1)
    sheet["B3"].number_format = DATE_FORMAT
    sheet["D3"] = "작성일"
    sheet["E3"] = report.created_on
    sheet["E3"].number_format = DATE_FORMAT
    _legend(sheet, 5)

    headers = (
        "납품처/그룹",
        "구분",
        "계획 수량(EA)",
        "실적 수량(EA)",
        "수량 증감(EA)",
        "수량 증감률",
        "계획 운반비(원)",
        "실적 운반비(원)",
        "운반비 증감(원)",
        "운반비 증감률",
        "계획 단위 운반비(원/EA)",
        "실적 단위 운반비(원/EA)",
        "단위 운반비 증감(원/EA)",
        "단위 운반비 증감률",
        "운반비 3개월 평균(원)",
        "운반비 6개월 평균(원)",
        "운반비 12개월 평균(원)",
        "운반비 전년도 평균(원)",
    )
    header_row = 8
    _write_header(sheet, header_row, headers)
    rows = [row for row in report.rows if row is not None]
    rows.extend((report.nonregular, report.total))
    for row_number, row in enumerate(rows, start=header_row + 1):
        calculation = row.calculation
        values = (
            *_summary_values(row.label, calculation),
            *_history_average_values(
                report.report_month,
                row.cost_won_by_month,
                HistoricalValueKind.MONEY,
            ),
        )
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
        _format_summary_row(sheet, row_number, calculation, bundle.operations)
    detail_last = header_row + len(rows)
    sheet.auto_filter.ref = f"A{header_row}:R{detail_last}"
    sheet.freeze_panes = "C9"

    sales_row = detail_last + 3
    _write_header(sheet, sales_row - 1, ("지표", "구분", "계획", "실적", "증감", "증감률"))
    sales = report.sales
    sales_variance = _optional_difference(sales.actual_won, sales.planned_won)
    sales_variance_pct = _optional_ratio(sales_variance, sales.planned_won)
    _write_values(
        sheet,
        sales_row,
        ("매출액", "수기 입력", sales.planned_won, sales.actual_won, sales_variance, _excel_number(sales_variance_pct)),
    )
    total = report.total.calculation
    _write_values(
        sheet,
        sales_row + 1,
        (
            "운반비/매출",
            "계산값",
            _excel_number(_optional_ratio(total.planned_cost_won, sales.planned_won)),
            _excel_number(_optional_ratio(total.actual_cost_won, sales.actual_won)),
            None,
            None,
        ),
    )
    for column in range(3, 6):
        sheet.cell(sales_row, column).number_format = MONEY_FORMAT
    sheet.cell(sales_row, 6).number_format = PERCENT_FORMAT
    for column in (3, 4):
        sheet.cell(sales_row + 1, column).number_format = PERCENT_FORMAT
    _fill_range(sheet, sales_row, 2, 6, _YELLOW)
    _fill_range(sheet, sales_row + 1, 2, 6, _CALCULATED)

    averages_header = sales_row + 4
    _write_header(
        sheet,
        averages_header,
        (
            "이력 평균",
            "3개월 평균",
            "3개월 상태",
            "6개월 평균",
            "6개월 상태",
            "12개월 평균",
            "12개월 상태",
            "전년도 평균",
            "전년도 상태",
        ),
    )
    average_sources = (
        ("수량(EA)", report.total.quantity_by_month, HistoricalValueKind.QUANTITY, NUMBER_FORMAT),
        ("운반비(원)", report.total.cost_won_by_month, HistoricalValueKind.MONEY, MONEY_FORMAT),
        ("매출액(원)", report.sales.actual_won_by_month, HistoricalValueKind.MONEY, MONEY_FORMAT),
    )
    for offset, (label, values, kind, number_format) in enumerate(average_sources, start=1):
        averages = historical_averages(report.report_month, values, value_kind=kind)
        periods = (
            averages.three_month,
            averages.six_month,
            averages.twelve_month,
            averages.comparison_year,
        )
        output = [label]
        for period in periods:
            output.extend((_excel_number(period.value), _average_status(period.complete, period.missing_months)))
        _write_values(sheet, averages_header + offset, tuple(output))
        for column in (2, 4, 6, 8):
            sheet.cell(averages_header + offset, column).number_format = number_format
        _fill_range(sheet, averages_header + offset, 2, 9, _CALCULATED)

    plan_title_row = averages_header + 5
    sheet.cell(plan_title_row, 1, "2026년 9월 운반비 계획 검토")
    sheet.cell(plan_title_row, 1).font = Font(
        name=FONT_NAME, size=13, bold=True, color=_NAVY
    )
    plan_header = plan_title_row + 1
    _write_header(
        sheet,
        plan_header,
        (
            "납품처",
            "구분",
            "계획 수량(EA)",
            "계획 운반비(원)",
            "계획 단위 운반비(원/EA)",
            "운반비 3개월 평균(원)",
            "운반비 6개월 평균(원)",
            "운반비 12개월 평균(원)",
            "운반비 전년도 평균(원)",
        ),
    )
    plan_rows = [row for row in report.next_month.rows if row is not None]
    plan_rows.extend((report.next_month.nonregular, report.next_month.total))
    for row_number, row in enumerate(plan_rows, start=plan_header + 1):
        _write_values(
            sheet,
            row_number,
            (
                row.label,
                _calculation_label(row.calculation),
                _excel_number(row.calculation.planned_quantity),
                row.calculation.planned_cost_won,
                _excel_number(row.calculation.planned_unit_cost),
                *_history_average_values(
                    report.next_month.month,
                    row.cost_won_by_month,
                    HistoricalValueKind.MONEY,
                ),
            ),
        )
        sheet.cell(row_number, 3).number_format = NUMBER_FORMAT
        sheet.cell(row_number, 4).number_format = MONEY_FORMAT
        sheet.cell(row_number, 5).number_format = UNIT_COST_FORMAT
        for column in range(6, 10):
            sheet.cell(row_number, column).number_format = MONEY_FORMAT
        if row.calculation.kind is CalculationKind.DESTINATION:
            _fill_range(sheet, row_number, 3, 5, _YELLOW)
            _fill_range(sheet, row_number, 6, 9, _CALCULATED)
        else:
            _fill_range(sheet, row_number, 1, 9, _CALCULATED)
            for column in range(1, 10):
                sheet.cell(row_number, column).font = Font(
                    name=FONT_NAME, bold=True, color=_DARK_TEXT
                )

    plan_summary_row = plan_header + len(plan_rows)
    total_history_costs = _history_average_values(
        report.next_month.month,
        report.next_month.total.cost_won_by_month,
        HistoricalValueKind.MONEY,
    )
    total_history_quantities = _history_average_values(
        report.next_month.month,
        report.next_month.total.quantity_by_month,
        HistoricalValueKind.QUANTITY,
    )
    sales_history_values = _history_average_values(
        report.next_month.month,
        report.next_month.sales.actual_won_by_month,
        HistoricalValueKind.MONEY,
    )
    for offset, values, number_format in (
        (
            1,
            (
                "총 대당 운반비",
                "계산값",
                None,
                None,
                _excel_number(report.next_month.total.calculation.planned_unit_cost),
                *(
                    _excel_number(_optional_ratio(cost, quantity))
                    for cost, quantity in zip(
                        total_history_costs, total_history_quantities
                    )
                ),
            ),
            UNIT_COST_FORMAT,
        ),
        (
            2,
            (
                "매출액",
                "수기 입력",
                None,
                report.next_month.sales.planned_won,
                None,
                *sales_history_values,
            ),
            MONEY_FORMAT,
        ),
        (
            3,
            (
                "매출액 대비 운반비",
                "계산값",
                None,
                _excel_number(
                    _optional_ratio(
                        report.next_month.total.calculation.planned_cost_won,
                        report.next_month.sales.planned_won,
                    )
                ),
                None,
                *(
                    _excel_number(_optional_ratio(cost, sales))
                    for cost, sales in zip(total_history_costs, sales_history_values)
                ),
            ),
            PERCENT_FORMAT,
        ),
    ):
        row_number = plan_summary_row + offset
        _write_values(sheet, row_number, values)
        for column in range(4, 10):
            sheet.cell(row_number, column).number_format = number_format
        _fill_range(sheet, row_number, 1, 9, _CALCULATED)

    widths = {
        "A": 22,
        "B": 13,
        "C": 16,
        "D": 16,
        "E": 16,
        "F": 13,
        "G": 18,
        "H": 18,
        "I": 18,
        "J": 13,
        "K": 20,
        "L": 20,
        "M": 22,
        "N": 15,
        "O": 20,
        "P": 20,
        "Q": 20,
        "R": 20,
    }
    _set_widths(sheet, widths)
    sheet.row_dimensions[header_row].height = 36


def _write_evidence(sheet: Worksheet, bundle: ReviewWorkbookReport) -> None:
    _title(sheet, "운행실적 입력 및 가져오기 증거")
    batch_headers = (
        "배치 ID",
        "원천 유형",
        "보고 월",
        "원본 파일명",
        "원본 SHA-256",
        "가져온 시각",
        "프로비넌스 ID",
    )
    _write_header(sheet, 3, batch_headers)
    for row_number, batch in enumerate(bundle.import_batches, start=4):
        _write_values(
            sheet,
            row_number,
            (
                batch.batch_id,
                batch.source_type,
                _month_date(batch.report_month),
                _safe_filename(batch.source_filename),
                batch.file_sha256,
                batch.imported_at,
                batch.provenance_id,
            ),
        )
        sheet.cell(row_number, 3).number_format = DATE_FORMAT
        sheet.cell(row_number, 6).number_format = DATETIME_FORMAT
        for column in (4, 5, 7):
            sheet.cell(row_number, column).alignment = Alignment(
                horizontal="left", vertical="center", wrap_text=True
            )
        sheet.row_dimensions[row_number].height = 36
        _fill_range(sheet, row_number, 1, 7, _BLUE)

    headers = (
        "입력 구분",
        "증거 구분",
        "값 구분",
        "배치 ID",
        "프로비넌스 ID",
        "보고 월",
        "원본 납품처",
        "표준 납품처",
        "수량(EA)",
        "운반비(원)",
        "원본 위치",
        "원천 유형",
    )
    header_row = max(9, 5 + len(bundle.import_batches))
    _write_header(sheet, header_row, headers)
    batch_by_id = {item.batch_id: item for item in bundle.import_batches}
    for row_number, item in enumerate(bundle.operations, start=header_row + 1):
        values = _evidence_values(item, batch_by_id)
        _write_values(sheet, row_number, values)
        sheet.cell(row_number, 6).number_format = DATE_FORMAT
        sheet.cell(row_number, 9).number_format = NUMBER_FORMAT
        sheet.cell(row_number, 10).number_format = MONEY_FORMAT
        for column in (1, 2, 3, 4, 6):
            sheet.cell(row_number, column).alignment = Alignment(
                horizontal="center", vertical="center"
            )
        fill = {"imported": _BLUE, "manual": _YELLOW, "derived": _CALCULATED}[item.input_kind]
        _fill_range(sheet, row_number, 1, 12, fill)
    last_row = header_row + len(bundle.operations)
    sheet.auto_filter.ref = f"A{header_row}:L{last_row}"
    sheet.freeze_panes = f"A{header_row + 1}"
    _set_widths(
        sheet,
        {"A": 13, "B": 22, "C": 13, "D": 22, "E": 42, "F": 20, "G": 28, "H": 20, "I": 14, "J": 16, "K": 38, "L": 16},
    )


def _write_validation(sheet: Worksheet, bundle: ReviewWorkbookReport) -> None:
    _title(sheet, "검증 결과")
    blocking_count = sum(issue.severity == "error" for issue in bundle.validation.issues)
    sheet["A3"] = "생성 가능"
    sheet["B3"] = "가능" if bundle.validation.can_generate else "불가"
    sheet["A4"] = "차단 건수"
    sheet["B4"] = blocking_count
    sheet["A5"] = "상태"
    sheet["B5"] = "생성 가능" if bundle.validation.can_generate else "차단 항목 해결 필요"
    for cell in (sheet["A3"], sheet["A4"], sheet["A5"]):
        cell.font = Font(name=FONT_NAME, bold=True, color=_DARK_TEXT)

    headers = ("코드", "심각도", "메시지", "납품처", "표시 순서", "원본 위치")
    _write_header(sheet, 7, headers)
    destination_by_id = {item.id: item for item in bundle.destinations}
    if bundle.validation.issues:
        safe_messages = tuple(
            _safe_display_text(issue.message) for issue in bundle.validation.issues
        )
        safe_locators = tuple(
            ""
            if issue.source_locator is None
            else _safe_locator(issue.source_locator)
            for issue in bundle.validation.issues
        )
        message_width = _validation_column_width(safe_messages)
        locator_width = _validation_column_width(safe_locators)
        lines_per_row = _validation_line_capacity(
            _VALIDATION_MAX_ROW_HEIGHT, _VALIDATION_FONT_SIZE
        )
        row_number = 8
        for issue, safe_message, safe_locator in zip(
            bundle.validation.issues, safe_messages, safe_locators
        ):
            destination = destination_by_id.get(issue.destination_id)
            message_lines = _wrap_text_lines(
                safe_message, message_width, _VALIDATION_FONT_SIZE
            )
            locator_lines = _wrap_text_lines(
                safe_locator, locator_width, _VALIDATION_FONT_SIZE
            )
            part_count = max(
                1,
                math.ceil(max(len(message_lines), len(locator_lines)) / lines_per_row),
            )
            color = _RED if issue.severity == "error" else _AMBER
            text = _RED_TEXT if issue.severity == "error" else _AMBER_TEXT
            for part_index in range(part_count):
                line_start = part_index * lines_per_row
                line_end = line_start + lines_per_row
                message_part = "\n".join(message_lines[line_start:line_end])
                locator_part = "\n".join(locator_lines[line_start:line_end])
                continued = part_index > 0
                values = (
                    (
                        f"{issue.code} (계속 {part_index + 1}/{part_count})"
                        if continued
                        else issue.code
                    ),
                    "계속"
                    if continued
                    else "오류"
                    if issue.severity == "error"
                    else "경고",
                    message_part,
                    None
                    if continued or destination is None
                    else destination.name,
                    None
                    if continued or destination is None
                    else destination.display_order,
                    locator_part,
                )
                _write_values(sheet, row_number, values)
                _fill_range(sheet, row_number, 1, 6, color)
                for column in range(1, 7):
                    sheet.cell(row_number, column).font = Font(
                        name=FONT_NAME,
                        size=_VALIDATION_FONT_SIZE,
                        color=text,
                        bold=column in (1, 2),
                    )
                for column in (3, 6):
                    sheet.cell(row_number, column).alignment = Alignment(
                        horizontal="left", vertical="top", wrap_text=True
                    )
                for column in (2, 5):
                    sheet.cell(row_number, column).alignment = Alignment(
                        horizontal="center", vertical="center"
                    )
                wrapped_lines = max(
                    _wrapped_line_count(
                        message_part, message_width, _VALIDATION_FONT_SIZE
                    ),
                    _wrapped_line_count(
                        locator_part, locator_width, _VALIDATION_FONT_SIZE
                    ),
                )
                sheet.row_dimensions[row_number].height = _validation_row_height(
                    wrapped_lines, _VALIDATION_FONT_SIZE
                )
                row_number += 1
        last_row = row_number - 1
    else:
        _write_values(
            sheet,
            8,
            ("VALID", "정상", "차단 또는 경고 항목이 없습니다.", None, None, None),
        )
        last_row = 8
        message_width = _VALIDATION_MIN_WIDTH
        locator_width = _VALIDATION_MIN_WIDTH
    sheet.auto_filter.ref = f"A7:F{last_row}"
    sheet.freeze_panes = "A8"
    _set_widths(
        sheet,
        {
            "A": 32,
            "B": 12,
            "C": message_width,
            "D": 20,
            "E": 12,
            "F": locator_width,
        },
    )


def _write_masters(sheet: Worksheet, bundle: ReviewWorkbookReport) -> None:
    _title(sheet, "마스터 기준 스냅샷")
    sheet["A4"] = "읽기 전용 검토 스냅샷입니다. 이 파일을 수정해도 기준정보는 변경되지 않습니다."
    sheet["A4"].font = Font(name=FONT_NAME, italic=True, color="595959")
    destination_headers = (
        "표시 순서",
        "납품처명",
        "별칭 (원천 유형)",
        "활성",
        "보고 필수",
        "대표 품목",
        "수량 합계 포함",
        "운반비 합계 포함",
        "매출 합계 포함",
    )
    _write_header(sheet, 5, destination_headers)
    aliases_by_destination: dict[int, list[DestinationAlias]] = {}
    for alias in bundle.aliases:
        aliases_by_destination.setdefault(alias.destination_id, []).append(alias)
    ordered_destinations = sorted(bundle.destinations, key=lambda item: (item.display_order, item.id))
    for row_number, destination in enumerate(ordered_destinations, start=6):
        aliases = sorted(
            aliases_by_destination.get(destination.id, []),
            key=lambda item: (item.source_type, item.raw_name, item.id),
        )
        alias_text = ", ".join(
            f"{alias.raw_name} ({alias.source_type})" for alias in aliases
        )
        values = (
            destination.display_order,
            destination.name,
            alias_text or None,
            _yes_no(destination.active),
            _yes_no(destination.required_for_report),
            destination.representative_item,
            _yes_no(destination.include_quantity_total),
            _yes_no(destination.include_cost_total),
            _yes_no(destination.include_sales_total),
        )
        _write_values(sheet, row_number, values)
        _fill_range(sheet, row_number, 1, 9, _LIGHT_BLUE)
        for column in (1, 4, 5, 7, 8, 9):
            sheet.cell(row_number, column).alignment = Alignment(
                horizontal="center", vertical="center"
            )
    destination_last = 5 + len(ordered_destinations)
    sheet.auto_filter.ref = f"A5:I{destination_last}"
    sheet.freeze_panes = "A6"

    group_header = destination_last + 3
    group_headers = (
        "그룹명",
        "그룹 표시 순서",
        "그룹 활성",
        "멤버 납품처",
        "멤버 표시 순서",
        "수량 포함",
        "운반비 포함",
    )
    _write_header(sheet, group_header, group_headers)
    destination_by_id = {item.id: item for item in bundle.destinations}
    members_by_group: dict[int, list[GroupMember]] = {}
    for member in bundle.group_members:
        members_by_group.setdefault(member.group_id, []).append(member)
    group_rows: list[tuple[ReportGroup, GroupMember | None]] = []
    for group in sorted(bundle.groups, key=lambda item: (item.display_order, item.id)):
        members = sorted(
            members_by_group.get(group.id, ()),
            key=lambda item: (item.display_order, item.destination_id),
        )
        group_rows.extend((group, member) for member in members or (None,))
    for row_number, (group, member) in enumerate(group_rows, start=group_header + 1):
        _write_values(
            sheet,
            row_number,
            (
                group.name,
                group.display_order,
                _yes_no(group.active),
                None if member is None else destination_by_id[member.destination_id].name,
                None if member is None else member.display_order,
                None if member is None else _yes_no(member.include_quantity),
                None if member is None else _yes_no(member.include_cost),
            ),
        )
        _fill_range(sheet, row_number, 1, 7, _CALCULATED)
        for column in (2, 3, 5, 6, 7):
            sheet.cell(row_number, column).alignment = Alignment(
                horizontal="center", vertical="center"
            )
    sheet.protection.sheet = True
    sheet.protection.enable()
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=9):
        for cell in row:
            cell.protection = Protection(locked=True)
    _set_widths(sheet, {"A": 14, "B": 22, "C": 38, "D": 12, "E": 14, "F": 22, "G": 16, "H": 18, "I": 16})


def _validate_report_identity(bundle: ReviewWorkbookReport) -> None:
    report = bundle.report
    if report.nonregular.calculation.actual_cost_won is None:
        raise ValueError("missing report input has no blocking validation issue")
    validate_ppt_report(report)
    # Force authoritative Task7 recomputation before any workbook/temp file exists.
    bundle.validation
    if report.report_month != "2026-08":
        raise ValueError("review workbook supports only report month 2026-08")
    if not bundle.import_batches:
        raise ValueError("at least one current import batch is required")

    _unique((item.id for item in bundle.destinations), "destination identity")
    _unique((item.name for item in bundle.destinations), "destination name identity")
    _unique(
        (item.display_order for item in bundle.destinations),
        "destination display order",
    )
    _unique((item.id for item in bundle.aliases), "alias identity")
    _unique(
        ((item.source_type, item.raw_name) for item in bundle.aliases),
        "alias source identity",
    )
    _unique((item.id for item in bundle.groups), "group identity")
    _unique((item.name for item in bundle.groups), "group name identity")
    _unique((item.display_order for item in bundle.groups), "group display order")
    _unique((item.batch_id for item in bundle.import_batches), "batch identity")
    _unique((item.provenance_id for item in bundle.import_batches), "batch provenance identity")
    _unique((item.source_type for item in bundle.import_batches), "batch source type")

    destination_by_id = {item.id: item for item in bundle.destinations}
    group_by_id = {item.id: item for item in bundle.groups}
    alias_by_source = {
        (item.source_type, item.raw_name): item.destination_id
        for item in bundle.aliases
    }
    for alias in bundle.aliases:
        if alias.destination_id not in destination_by_id:
            raise ValueError("alias has a missing destination identity")
    seen_members: set[tuple[int, int]] = set()
    seen_member_orders: set[tuple[int, int]] = set()
    for member in bundle.group_members:
        identity = (member.group_id, member.destination_id)
        if identity in seen_members:
            raise ValueError("group members contain a duplicate identity")
        seen_members.add(identity)
        if member.group_id not in group_by_id or member.destination_id not in destination_by_id:
            raise ValueError("group member has a missing identity")
        if member.name != destination_by_id[member.destination_id].name:
            raise ValueError("group member name does not match canonical destination")
        order_identity = (member.group_id, member.display_order)
        if order_identity in seen_member_orders:
            raise ValueError("group member display order contains a duplicate")
        seen_member_orders.add(order_identity)

    displayed = [row for row in report.rows if row is not None]
    keys = [row.key for row in displayed] + [report.nonregular.key, report.total.key]
    for key in keys:
        _text(key, "report row identity")
    _unique(keys, "report row identity")
    direct_calculations: dict[int, DestinationCalculation] = {}
    group_ids: list[int] = []
    for row in displayed:
        calculation = row.calculation
        if calculation.kind is CalculationKind.DESTINATION:
            destination_id = calculation.provenance.destination_id
            if destination_id is None or destination_id not in destination_by_id:
                raise ValueError("report row has a missing destination identity")
            if row.label != destination_by_id[destination_id].name:
                raise ValueError("report row label does not match destination identity")
            if destination_id in direct_calculations:
                raise ValueError("displayed destination identity contains a duplicate identity")
            direct_calculations[destination_id] = calculation
        elif calculation.kind is CalculationKind.DERIVED_GROUP:
            group_id = calculation.provenance.group_id
            if group_id not in group_by_id:
                raise ValueError("report row has a missing group identity")
            if row.label != group_by_id[group_id].name:
                raise ValueError("report row label does not match group identity")
            group_ids.append(group_id)
        else:
            raise ValueError("only the total row may use the grand-total identity")
    _unique(group_ids, "displayed group identity")
    expected_destinations = {
        item.id for item in bundle.destinations if item.active and item.required_for_report
    }
    if set(direct_calculations) != expected_destinations:
        raise ValueError("report has a missing destination identity")
    expected_groups = {item.id for item in bundle.groups if item.active}
    if set(group_ids) != expected_groups:
        raise ValueError("report has a missing group identity")
    if report.total.calculation.kind is not CalculationKind.GRAND_TOTAL:
        raise ValueError("total row has a missing grand-total identity")

    members_by_group: dict[int, list[GroupMember]] = {}
    for member in bundle.group_members:
        members_by_group.setdefault(member.group_id, []).append(member)
    for row in displayed:
        if row.calculation.kind is not CalculationKind.DERIVED_GROUP:
            continue
        group_id = row.calculation.provenance.group_id
        try:
            expected = calculate_group(
                direct_calculations,
                sorted(
                    members_by_group.get(group_id, []),
                    key=lambda item: (item.display_order, item.destination_id),
                ),
                group_id=group_id,
            )
        except (TypeError, ValueError) as error:
            raise ValueError("group calculation does not match master provenance") from error
        if row.calculation != expected:
            raise ValueError("group calculation does not match master provenance")
    total_rules = tuple(
        sorted(
            (
                item
                for item in bundle.destinations
                if item.id in direct_calculations
            ),
            key=lambda item: (item.display_order, item.id),
        )
    )
    try:
        expected_total = calculate_total(direct_calculations, total_rules, nonregular=report.nonregular.calculation)
    except (TypeError, ValueError) as error:
        raise ValueError("total calculation does not match master provenance") from error
    if report.total.calculation != expected_total:
        raise ValueError("total calculation does not match master provenance")

    next_rows = tuple(report.next_month.rows)
    expected_next_identity = tuple((item.id, item.name) for item in total_rules)
    actual_next_identity = tuple(
        (row.calculation.destination_id, row.label) for row in next_rows
    )
    if actual_next_identity != expected_next_identity:
        raise ValueError("next-month rows do not match master identity and order")
    next_calculations = {
        row.calculation.destination_id: row.calculation for row in next_rows
    }
    try:
        expected_next_total = calculate_total(next_calculations, total_rules, nonregular=report.next_month.nonregular.calculation)
    except (TypeError, ValueError) as error:
        raise ValueError(
            "next-month total calculation does not match master provenance"
        ) from error
    if report.next_month.total.calculation != expected_next_total:
        raise ValueError(
            "next-month total calculation does not match master provenance"
        )

    _validate_validation_snapshot(bundle, direct_calculations)

    batch_by_id = {item.batch_id: item for item in bundle.import_batches}
    for batch in bundle.import_batches:
        if batch.report_month != report.report_month:
            raise ValueError("batch report month does not match canonical report")
    _unique(
        (
            (
                item.input_kind,
                item.record_kind,
                item.value_role,
                item.batch_id,
                item.provenance_id,
                item.destination_id,
                item.source_locator,
            )
            for item in bundle.operations
        ),
        "operation evidence identity",
    )
    operation_by_input: dict[
        tuple[str, str, str, int | None], list[OperationEvidence]
    ] = {}
    for item in bundle.operations:
        identity = (
            item.report_month,
            item.record_kind,
            item.value_role,
            item.destination_id,
        )
        operation_by_input.setdefault(identity, []).append(item)
        if item.record_kind == "destination":
            destination = destination_by_id.get(item.destination_id)
            if destination is None or item.normalized_destination != destination.name:
                raise ValueError("operation has a missing destination identity")
            if item.destination_id not in direct_calculations:
                raise ValueError("operation has a missing report destination identity")
        elif item.record_kind == "nonregular":
            if item.normalized_destination != report.nonregular.label:
                raise ValueError("nonregular evidence label does not match canonical report")
        elif item.normalized_destination != "매출액":
            raise ValueError("sales evidence label does not match canonical report")
        if item.input_kind == "imported":
            batch = batch_by_id.get(item.batch_id)
            if batch is None or batch.provenance_id != item.provenance_id:
                raise ValueError("operation has a missing batch identity")
            if (
                item.record_kind == "destination"
                and alias_by_source.get((batch.source_type, item.raw_destination))
                != item.destination_id
            ):
                raise ValueError("operation alias identity does not match master aliases")
    expected_inputs: dict[
        tuple[str, str, str, int | None],
        tuple[Decimal | None, int | None],
    ] = {}
    for destination_id, calculation in direct_calculations.items():
        expected_inputs[(report.report_month, "destination", "plan", destination_id)] = (
            calculation.planned_quantity,
            calculation.planned_cost_won,
        )
        expected_inputs[(report.report_month, "destination", "actual", destination_id)] = (
            calculation.actual_quantity,
            calculation.actual_cost_won,
        )
    for role, quantity, cost in (
        (
            "plan",
            report.nonregular.calculation.planned_quantity,
            report.nonregular.calculation.planned_cost_won,
        ),
        (
            "actual",
            report.nonregular.calculation.actual_quantity,
            report.nonregular.calculation.actual_cost_won,
        ),
    ):
        expected_inputs[(report.report_month, "nonregular", role, None)] = (
            quantity,
            cost,
        )
    expected_inputs[(report.report_month, "sales", "plan", None)] = (
        None,
        report.sales.planned_won,
    )
    expected_inputs[(report.report_month, "sales", "actual", None)] = (
        None,
        report.sales.actual_won,
    )
    for row in report.next_month.rows:
        destination_id = row.calculation.destination_id
        expected_inputs[(report.next_month.month, "destination", "plan", destination_id)] = (
            row.calculation.planned_quantity,
            row.calculation.planned_cost_won,
        )
    expected_inputs[(report.next_month.month, "nonregular", "plan", None)] = (
        report.next_month.nonregular.calculation.planned_quantity,
        report.next_month.nonregular.calculation.planned_cost_won,
    )
    expected_inputs[(report.next_month.month, "sales", "plan", None)] = (
        None,
        report.next_month.sales.planned_won,
    )
    if set(operation_by_input) - set(expected_inputs):
        raise ValueError("operation evidence contains an input not displayed by the report")
    blocking_issues = tuple(
        issue for issue in bundle.validation.issues if issue.severity == "error"
    )
    for identity, (expected_quantity, expected_cost) in expected_inputs.items():
        records = operation_by_input.get(identity, [])
        destination_id = identity[3]
        evidence_label = (
            "nonregular evidence"
            if identity[1] == "nonregular"
            else "operation evidence"
        )
        missing = (
            expected_cost is None
            if identity[1] == "sales"
            else expected_quantity is None or expected_cost is None
        )
        if missing and not _missing_input_is_blocked(identity, blocking_issues, report):
            raise ValueError("missing report input has no blocking validation issue")
        if expected_quantity is None and expected_cost is None:
            if len(records) > 1 or any(
                item.quantity_ea is not None or item.cost_won is not None
                for item in records
            ):
                raise ValueError(
                    f"{evidence_label} does not reconcile to missing input"
                )
            continue
        if len(records) != 1:
            if len(records) > 1:
                raise ValueError("every present input requires exactly one evidence record")
            raise ValueError(
                f"{evidence_label} does not reconcile to canonical inputs"
            )
        record = records[0]
        if (
            record.quantity_ea != expected_quantity
            or record.cost_won != expected_cost
        ):
            raise ValueError(
                f"{evidence_label} does not reconcile to canonical inputs"
            )


def _missing_input_is_blocked(
    identity: tuple[str, str, str, int | None],
    blocking_issues: tuple[object, ...],
    report: PptReport,
) -> bool:
    month, record_kind, value_role, destination_id = identity
    if (
        month == report.report_month
        and record_kind == "destination"
        and value_role == "actual"
    ):
        return any(
            issue.code == "MISSING_ACTUAL_QUANTITY"
            and issue.destination_id == destination_id
            for issue in blocking_issues
        )
    if (
        month == report.next_month.month
        and record_kind == "destination"
        and value_role == "plan"
    ):
        return any(
            issue.code in {"MISSING_NEXT_MONTH_PLAN", "REPORT_MONTH_MISMATCH"}
            and issue.destination_id == destination_id
            for issue in blocking_issues
        )
    if (
        month == report.report_month
        and record_kind == "sales"
        and value_role == "actual"
    ):
        return any(issue.code == "MISSING_SALES" for issue in blocking_issues)
    return False


def _validate_validation_snapshot(
    bundle: ReviewWorkbookReport,
    direct_calculations: dict[int, DestinationCalculation],
) -> None:
    """Reconcile Task7 inputs with the canonical report rather than trusting flags."""
    report = bundle.report
    context = bundle.validation_context
    if context.report_month != report.report_month:
        raise ValueError("validation snapshot report month conflicts with canonical report")
    if (
        context.sales is None
        or context.sales.amount_won != report.sales.actual_won
        or context.sales.confirmed_at is None
    ):
        raise ValueError("validation snapshot sales conflicts with canonical report")
    expected_history = historical_averages(
        report.report_month, report.total.quantity_by_month
    ).comparison_year
    if context.prior_year_history != expected_history:
        raise ValueError("validation snapshot history conflicts with canonical report")

    destination_by_id = {item.id: item for item in bundle.destinations}
    validation_by_id = {item.destination_id: item for item in context.destinations}
    if set(validation_by_id) != set(direct_calculations):
        raise ValueError("validation snapshot destinations conflict with canonical report")
    next_by_id = {
        row.calculation.provenance.destination_id: row.calculation
        for row in report.next_month.rows
        if row is not None
        and row.calculation.provenance.destination_id is not None
    }
    for destination_id, calculation in direct_calculations.items():
        master = destination_by_id[destination_id]
        item = validation_by_id[destination_id]
        next_plan = item.next_month_plan
        if (
            item.name != master.name
            or item.display_order != master.display_order
            or item.required_for_report != master.required_for_report
            or item.actual_quantity != calculation.actual_quantity
            or next_plan is None
            or next_plan.report_month != report.next_month.month
            or destination_id not in next_by_id
            or next_plan.quantity != next_by_id[destination_id].planned_quantity
        ):
            raise ValueError(
                "validation snapshot destination values conflict with canonical report"
            )

    if len(context.current_import_batches) != len(bundle.import_batches):
        raise ValueError("validation snapshot import provenance conflicts with evidence")
    evidence_by_id = {item.batch_id: item for item in bundle.import_batches}
    for item in context.current_import_batches:
        evidence = evidence_by_id.get(item.batch_id)
        if (
            evidence is None
            or not item.is_current
            or item.report_month != evidence.report_month
            or item.source_type != evidence.source_type
            or item.file_sha256 != evidence.file_sha256
        ):
            raise ValueError("validation snapshot import provenance conflicts with evidence")


def _validate_saved_workbook(path: Path, bundle: ReviewWorkbookReport) -> None:
    workbook = load_workbook(path, data_only=False, keep_links=True)
    try:
        if tuple(workbook.sheetnames) != SHEET_NAMES:
            raise ValueError("saved workbook sheet order changed")
        if workbook.vba_archive is not None or workbook._external_links:
            raise ValueError("saved workbook contains macros or external links")
        summary = workbook[SHEET_NAMES[0]]
        report_month_value = summary["B3"].value
        if isinstance(report_month_value, datetime):
            report_month_value = report_month_value.date()
        if report_month_value != date(2026, 8, 1) or not summary["B3"].is_date:
            raise ValueError("saved report month is not a typed date")
        summary_rows = [row for row in bundle.report.rows if row is not None]
        summary_rows.extend((bundle.report.nonregular, bundle.report.total))
        for row_number, report_row in enumerate(summary_rows, start=9):
            for column, expected in enumerate(
                (
                    *_summary_values(report_row.label, report_row.calculation),
                    *_history_average_values(
                        bundle.report.report_month,
                        report_row.cost_won_by_month,
                        HistoricalValueKind.MONEY,
                    ),
                ),
                start=1,
            ):
                _assert_saved_value(summary.cell(row_number, column), expected)
            for column in (6, 10, 14):
                if summary.cell(row_number, column).number_format != PERCENT_FORMAT:
                    raise ValueError("saved percentage format changed")

        sales_row = 8 + len(summary_rows) + 3
        sales = bundle.report.sales
        sales_variance = _optional_difference(sales.actual_won, sales.planned_won)
        sales_values = (
            "매출액",
            "수기 입력",
            sales.planned_won,
            sales.actual_won,
            sales_variance,
            _excel_number(_optional_ratio(sales_variance, sales.planned_won)),
        )
        ratio_values = (
            "운반비/매출",
            "계산값",
            _excel_number(
                _optional_ratio(
                    bundle.report.total.calculation.planned_cost_won,
                    sales.planned_won,
                )
            ),
            _excel_number(
                _optional_ratio(
                    bundle.report.total.calculation.actual_cost_won,
                    sales.actual_won,
                )
            ),
            None,
            None,
        )
        for offset, values in enumerate((sales_values, ratio_values)):
            for column, expected in enumerate(values, start=1):
                _assert_saved_value(
                    summary.cell(sales_row + offset, column), expected
                )

        averages_header = sales_row + 4
        average_sources = (
            (
                "수량(EA)",
                bundle.report.total.quantity_by_month,
                HistoricalValueKind.QUANTITY,
            ),
            (
                "운반비(원)",
                bundle.report.total.cost_won_by_month,
                HistoricalValueKind.MONEY,
            ),
            (
                "매출액(원)",
                bundle.report.sales.actual_won_by_month,
                HistoricalValueKind.MONEY,
            ),
        )
        for offset, (label, values, kind) in enumerate(average_sources, start=1):
            averages = historical_averages(
                bundle.report.report_month, values, value_kind=kind
            )
            expected_values: list[object] = [label]
            for average in (
                averages.three_month,
                averages.six_month,
                averages.twelve_month,
                averages.comparison_year,
            ):
                expected_values.extend(
                    (
                        _excel_number(average.value),
                        _average_status(average.complete, average.missing_months),
                    )
                )
            for column, expected in enumerate(expected_values, start=1):
                _assert_saved_value(
                    summary.cell(averages_header + offset, column), expected
                )

        plan_title_row = averages_header + 5
        _assert_saved_value(
            summary.cell(plan_title_row, 1), "2026년 9월 운반비 계획 검토"
        )
        plan_rows = [
            row for row in bundle.report.next_month.rows if row is not None
        ]
        plan_rows.extend(
            (bundle.report.next_month.nonregular, bundle.report.next_month.total)
        )
        for row_number, report_row in enumerate(
            plan_rows, start=plan_title_row + 2
        ):
            expected_values = (
                report_row.label,
                _calculation_label(report_row.calculation),
                _excel_number(report_row.calculation.planned_quantity),
                report_row.calculation.planned_cost_won,
                _excel_number(report_row.calculation.planned_unit_cost),
                *_history_average_values(
                    bundle.report.next_month.month,
                    report_row.cost_won_by_month,
                    HistoricalValueKind.MONEY,
                ),
            )
            for column, expected in enumerate(expected_values, start=1):
                _assert_saved_value(summary.cell(row_number, column), expected)
        plan_last_row = plan_title_row + 1 + len(plan_rows)
        costs = _history_average_values(
            bundle.report.next_month.month,
            bundle.report.next_month.total.cost_won_by_month,
            HistoricalValueKind.MONEY,
        )
        quantities = _history_average_values(
            bundle.report.next_month.month,
            bundle.report.next_month.total.quantity_by_month,
            HistoricalValueKind.QUANTITY,
        )
        sales_history = _history_average_values(
            bundle.report.next_month.month,
            bundle.report.next_month.sales.actual_won_by_month,
            HistoricalValueKind.MONEY,
        )
        plan_summary_values = (
            (
                "총 대당 운반비",
                "계산값",
                None,
                None,
                _excel_number(
                    bundle.report.next_month.total.calculation.planned_unit_cost
                ),
                *(
                    _excel_number(_optional_ratio(cost, quantity))
                    for cost, quantity in zip(costs, quantities)
                ),
            ),
            (
                "매출액",
                "수기 입력",
                None,
                bundle.report.next_month.sales.planned_won,
                None,
                *sales_history,
            ),
            (
                "매출액 대비 운반비",
                "계산값",
                None,
                _excel_number(
                    _optional_ratio(
                        bundle.report.next_month.total.calculation.planned_cost_won,
                        bundle.report.next_month.sales.planned_won,
                    )
                ),
                None,
                *(
                    _excel_number(_optional_ratio(cost, sales))
                    for cost, sales in zip(costs, sales_history)
                ),
            ),
        )
        for offset, values in enumerate(plan_summary_values, start=1):
            for column, expected in enumerate(values, start=1):
                _assert_saved_value(
                    summary.cell(plan_last_row + offset, column), expected
                )

        evidence = workbook[SHEET_NAMES[1]]
        for row_number, batch in enumerate(bundle.import_batches, start=4):
            values = (
                batch.batch_id,
                batch.source_type,
                _month_date(batch.report_month),
                _safe_filename(batch.source_filename),
                batch.file_sha256,
                batch.imported_at,
                batch.provenance_id,
            )
            for column, expected in enumerate(values, start=1):
                _assert_saved_value(evidence.cell(row_number, column), expected)
        header_row = max(9, 5 + len(bundle.import_batches))
        batch_by_id = {item.batch_id: item for item in bundle.import_batches}
        for row_number, item in enumerate(
            bundle.operations, start=header_row + 1
        ):
            values = _evidence_values(item, batch_by_id)
            for column, expected in enumerate(values, start=1):
                _assert_saved_value(evidence.cell(row_number, column), expected)

        checks = workbook[SHEET_NAMES[2]]
        _assert_saved_value(
            checks["B3"], "가능" if bundle.validation.can_generate else "불가"
        )
        _assert_saved_value(
            checks["B4"],
            sum(issue.severity == "error" for issue in bundle.validation.issues),
        )

        masters = workbook[SHEET_NAMES[3]]
        if not masters.protection.sheet:
            raise ValueError("saved master snapshot is not read-only")
        for sheet in workbook.worksheets:
            if not sheet.freeze_panes or not sheet.auto_filter.ref:
                raise ValueError(f"{sheet.title} is missing freeze panes or filters")
            if not any(cell.style_id for row in sheet.iter_rows() for cell in row):
                raise ValueError(f"{sheet.title} is missing review styles")
            for column, dimension in sheet.column_dimensions.items():
                max_width = (
                    _VALIDATION_MAX_WIDTH
                    if sheet.title == SHEET_NAMES[2] and column in {"C", "F"}
                    else 59
                )
                if dimension.width is None or not 0 < dimension.width <= max_width:
                    raise ValueError(f"{sheet.title} has an invalid column width")
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "e":
                        raise ValueError(
                            f"{sheet.title}!{cell.coordinate} contains an Excel error"
                        )
                    if isinstance(cell.value, str):
                        if _contains_absolute_path(cell.value):
                            raise ValueError(
                                f"{sheet.title}!{cell.coordinate} exposes an absolute path"
                            )
                        if cell.value in ERROR_TOKENS:
                            raise ValueError(
                                f"{sheet.title}!{cell.coordinate} contains a formula error"
                            )
                        if cell.value.startswith("="):
                            raise ValueError(
                                f"{sheet.title}!{cell.coordinate} contains a formula"
                            )
    finally:
        workbook.close()
    with ZipFile(path) as archive:
        members = {name.lower() for name in archive.namelist()}
        if any("vbaproject.bin" in name or name.startswith("xl/externallinks/") for name in members):
            raise ValueError("saved workbook package contains macros or external links")


def _summary_values(
    label: str, calculation: DestinationCalculation
) -> tuple[object, ...]:
    return (
        label,
        _calculation_label(calculation),
        _excel_number(calculation.planned_quantity),
        _excel_number(calculation.actual_quantity),
        _excel_number(calculation.quantity_variance),
        _excel_number(calculation.quantity_variance_pct),
        calculation.planned_cost_won,
        calculation.actual_cost_won,
        calculation.cost_variance_won,
        _excel_number(calculation.cost_variance_pct),
        _excel_number(calculation.planned_unit_cost),
        _excel_number(calculation.actual_unit_cost),
        _excel_number(calculation.actual_unit_cost_variance),
        _excel_number(calculation.actual_unit_cost_variance_pct),
    )


def _evidence_values(
    item: OperationEvidence, batch_by_id: dict[int, ImportBatchEvidence]
) -> tuple[object, ...]:
    batch = batch_by_id.get(item.batch_id) if item.batch_id is not None else None
    return (
        {"imported": "가져오기", "manual": "수기 입력", "derived": "혼합 원천 집계"}[item.input_kind],
        {
            "destination": "납품처",
            "nonregular": "비정규",
            "sales": "매출액",
        }[item.record_kind],
        "계획" if item.value_role == "plan" else "실적",
        item.batch_id,
        item.provenance_id,
        _month_date(item.report_month),
        item.raw_destination,
        item.normalized_destination,
        _excel_number(item.quantity_ea),
        item.cost_won,
        _safe_locator(item.source_locator),
        None if batch is None else batch.source_type,
    )


def _format_summary_row(
    sheet: Worksheet,
    row: int,
    calculation: DestinationCalculation,
    operations: tuple[OperationEvidence, ...],
) -> None:
    for column in (3, 4, 5):
        sheet.cell(row, column).number_format = NUMBER_FORMAT
    for column in (6, 10, 14):
        sheet.cell(row, column).number_format = PERCENT_FORMAT
    for column in (7, 8, 9):
        sheet.cell(row, column).number_format = MONEY_FORMAT
    for column in (11, 12, 13):
        sheet.cell(row, column).number_format = UNIT_COST_FORMAT
    for column in (15, 16, 17, 18):
        sheet.cell(row, column).number_format = MONEY_FORMAT
    if calculation.kind is CalculationKind.DESTINATION:
        for column in (3, 7):
            sheet.cell(row, column).fill = PatternFill("solid", fgColor=_YELLOW)
        destination_id = calculation.provenance.destination_id
        sheet.cell(row, 4).fill = PatternFill(
            "solid",
            fgColor=_operation_fill(operations, destination_id, "quantity_ea"),
        )
        sheet.cell(row, 8).fill = PatternFill(
            "solid",
            fgColor=_operation_fill(operations, destination_id, "cost_won"),
        )
        for column in (5, 6, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18):
            sheet.cell(row, column).fill = PatternFill("solid", fgColor=_CALCULATED)
    else:
        _fill_range(sheet, row, 1, 18, _CALCULATED)
        for column in range(1, 19):
            sheet.cell(row, column).font = Font(name=FONT_NAME, bold=True, color=_DARK_TEXT)


def _operation_fill(
    operations: tuple[OperationEvidence, ...],
    destination_id: int | None,
    metric: Literal["quantity_ea", "cost_won"],
) -> str:
    kinds = {
        item.input_kind
        for item in operations
        if item.destination_id == destination_id and getattr(item, metric) is not None
    }
    if kinds == {"manual"}:
        return _YELLOW
    if kinds == {"imported"}:
        return _BLUE
    return _CALCULATED


def _legend(sheet: Worksheet, row: int) -> None:
    entries = (
        ("범례", None),
        ("수기 입력", _YELLOW),
        ("가져오기", _BLUE),
        ("계산값", _CALCULATED),
        ("경고", _AMBER),
        ("오류", _RED),
    )
    for column, (label, fill) in enumerate(entries, start=1):
        cell = sheet.cell(row, column, label)
        cell.font = Font(name=FONT_NAME, bold=column == 1, color=_DARK_TEXT)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)


def _title(sheet: Worksheet, text: str) -> None:
    sheet["A2"] = text
    sheet["A2"].font = Font(name=FONT_NAME, size=15, bold=True, color=_NAVY)
    sheet["A2"].border = Border(bottom=Side(style="medium", color=_NAVY))
    sheet.row_dimensions[2].height = 26


def _write_header(sheet: Worksheet, row: int, headers: tuple[str, ...]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row, column, header)
        cell.fill = PatternFill("solid", fgColor=_NAVY)
        cell.font = Font(name=FONT_NAME, bold=True, color=_WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BOTTOM_BORDER
    sheet.row_dimensions[row].height = 30


def _write_values(sheet: Worksheet, row: int, values: tuple[object, ...]) -> None:
    for column, value in enumerate(values, start=1):
        if isinstance(value, str):
            value = _safe_display_text(value)
        cell = sheet.cell(row, column, value)
        cell.alignment = Alignment(
            horizontal=_default_horizontal(value), vertical="center"
        )


def _fill_range(sheet: Worksheet, row: int, first_column: int, last_column: int, color: str) -> None:
    fill = PatternFill("solid", fgColor=color)
    for column in range(first_column, last_column + 1):
        sheet.cell(row, column).fill = fill


def _apply_font_to_used_range(sheet: Worksheet) -> None:
    for row in sheet.iter_rows(
        min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column
    ):
        for cell in row:
            if cell.value is None:
                continue
            current = cell.font
            cell.font = Font(
                name=FONT_NAME,
                size=current.sz or 10,
                bold=current.bold,
                italic=current.italic,
                color=current.color,
            )
            if cell.alignment == Alignment():
                cell.alignment = Alignment(
                    horizontal=_default_horizontal(cell.value), vertical="center"
                )
            elif cell.alignment.vertical is None:
                cell.alignment = Alignment(
                    horizontal=(
                        cell.alignment.horizontal
                        or _default_horizontal(cell.value)
                    ),
                    vertical="center",
                    wrap_text=cell.alignment.wrap_text,
                )
            if cell.row not in (2, 5, 7, 8, 9):
                sheet.row_dimensions[cell.row].height = sheet.row_dimensions[cell.row].height or 22


def _set_widths(sheet: Worksheet, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _safe_filename(value: str) -> str:
    return value.replace("\\", "/").rstrip("/").rsplit("/", 1)[-1]


def _safe_locator(value: str) -> str:
    return _safe_display_text(value)


def _safe_display_text(value: str) -> str:
    """Remove absolute directory components from embedded user-visible paths."""
    output: list[str] = []
    cursor = 0
    protected_spans = tuple(match.span() for match in _WEB_URL_TOKEN.finditer(value))
    while match := _next_absolute_path_match(value, cursor, protected_spans):
        start = match.start()
        end = _absolute_path_end(value, start)
        path_text = value[start:end]
        basename = _safe_path_basename(path_text)
        output.extend((value[cursor:start], basename))
        cursor = end
    output.append(value[cursor:])
    return "".join(output)


def _contains_absolute_path(value: str) -> bool:
    """Detect an absolute path independently from its display replacement."""
    protected_spans = tuple(match.span() for match in _WEB_URL_TOKEN.finditer(value))
    return _next_absolute_path_match(value, 0, protected_spans) is not None


def _next_absolute_path_match(
    value: str, cursor: int, protected_spans: tuple[tuple[int, int], ...]
):
    while match := _ABSOLUTE_PATH_START.search(value, cursor):
        protected_end = next(
            (
                end
                for start, end in protected_spans
                if start <= match.start() < end
            ),
            None,
        )
        if protected_end is not None:
            cursor = protected_end
            continue
        if match.group(0) == "/" and not _is_credible_posix_path(
            value[match.start() : _absolute_path_end(value, match.start())]
        ):
            cursor = match.end()
            continue
        return match
    return None


def _is_credible_posix_path(candidate: str) -> bool:
    clean = candidate.strip()
    locator_match = _LOCATOR_SUFFIX.search(clean)
    path = clean if locator_match is None else clean[: locator_match.start()]
    components = [part.strip() for part in path.split("/")[1:] if part.strip()]
    if len(components) >= 2:
        return True
    if len(components) != 1:
        return False
    return locator_match is not None or bool(
        re.search(r"\.[^./\\\s]+$", components[0])
    )


def _absolute_path_end(value: str, start: int) -> int:
    """Find a strong embedded-path boundary without treating spaces as separators."""
    if start > 0 and value[start - 1] in {'"', "'"}:
        quote = value[start - 1]
        closing = value.find(quote, start)
        if closing >= 0:
            return closing
    boundary = _PATH_BOUNDARY.search(value, start)
    return len(value) if boundary is None else boundary.start()


def _safe_path_basename(path_text: str) -> str:
    clean = path_text.strip()
    locator_match = _LOCATOR_SUFFIX.search(clean)
    locator = "" if locator_match is None else locator_match.group(1)
    path = clean if locator_match is None else clean[: locator_match.start()]
    windows_path = bool(re.match(r"(?i)^[A-Z]:[\\/]", path)) or path.startswith(
        ("\\\\", "//")
    )
    parts = [part for part in re.split(r"[\\/]", path) if part]
    if windows_path and path.startswith(("\\\\", "//")) and len(parts) <= 2:
        basename = ""
    elif windows_path and len(parts) <= 1:
        basename = ""
    else:
        basename = "" if not parts else parts[-1].strip()
    if basename in {"", ".", ".."}:
        basename = _REDACTED_PATH
    return f"{basename}{locator}"


def _yes_no(value: bool) -> str:
    return "예" if value else "아니오"


def _default_horizontal(value: object) -> str:
    if isinstance(value, bool):
        return "center"
    if isinstance(value, (int, float, Decimal, date, datetime)):
        return "right"
    return "left"


def _calculation_label(calculation: DestinationCalculation) -> str:
    if calculation.kind is CalculationKind.DESTINATION:
        return "납품처"
    if calculation.kind is CalculationKind.DERIVED_GROUP:
        return "그룹 계산"
    return "총계 계산"


def _average_status(complete: bool, missing_months: tuple[str, ...]) -> str:
    if complete:
        return "완전"
    return f"자료 부족 ({', '.join(missing_months)})"


def _wrapped_line_count(
    value: str, column_width: float, font_size: float = 10
) -> int:
    return len(_wrap_text_lines(value, column_width, font_size))


def _wrap_text_lines(
    value: str, column_width: float, font_size: float
) -> tuple[str, ...]:
    available_pixels = max(
        1,
        int(
            (_excel_column_width_pixels(column_width) - _VALIDATION_CELL_MARGIN_PIXELS)
            * 0.85
        ),
    )
    wrapped: list[str] = []
    for logical_line in value.split("\n"):
        if not logical_line:
            wrapped.append("")
            continue
        current: list[str] = []
        current_pixels = 0.0
        for character in logical_line:
            character_pixels = _glyph_pixel_width(character, font_size)
            if current and current_pixels + character_pixels > available_pixels:
                wrapped.append("".join(current))
                current = []
                current_pixels = 0.0
            current.append(character)
            current_pixels += character_pixels
        wrapped.append("".join(current))
    return tuple(wrapped or ("",))


def _validation_column_width(values: tuple[str, ...]) -> int:
    max_lines = _validation_line_capacity(
        _VALIDATION_MAX_ROW_HEIGHT, _VALIDATION_FONT_SIZE
    )
    for width in range(_VALIDATION_MIN_WIDTH, _VALIDATION_MAX_WIDTH + 1):
        if all(
            _wrapped_line_count(value, width, _VALIDATION_FONT_SIZE) <= max_lines
            for value in values
        ):
            return width
    return _VALIDATION_MAX_WIDTH


def _validation_row_height(line_count: int, font_size: float) -> float:
    return min(
        _VALIDATION_MAX_ROW_HEIGHT,
        max(
            _VALIDATION_MIN_ROW_HEIGHT,
            6 + (_validation_line_height_points(font_size) * max(1, line_count)),
        ),
    )


def _validation_line_capacity(row_height: float, font_size: float) -> int:
    return max(
        1,
        math.floor(
            (row_height - 6) / _validation_line_height_points(font_size)
        ),
    )


def _excel_column_width_pixels(width: float) -> int:
    return max(1, math.floor((width * 7) + 5))


@lru_cache(maxsize=16)
def _validation_font(font_size: float):
    if ImageFont is None:
        return None
    pixel_size = max(1, round(font_size * _SCREEN_DPI / 72))
    windows_directory = os.environ.get("WINDIR")
    candidates = [Path("C:/Windows/Fonts/malgun.ttf")]
    if windows_directory:
        candidates.insert(0, Path(windows_directory) / "Fonts" / "malgun.ttf")
    for candidate in candidates:
        if not candidate.is_file():
            continue
        try:
            return ImageFont.truetype(str(candidate), pixel_size)
        except OSError:
            continue
    return None


@lru_cache(maxsize=4096)
def _glyph_pixel_width(character: str, font_size: float) -> float:
    font = _validation_font(font_size)
    if font is not None:
        return max(1.0, float(font.getlength(character)))
    pixel_size = font_size * _SCREEN_DPI / 72
    if character == "\t":
        return pixel_size * 2.8
    if character.isspace():
        return pixel_size * 0.55
    if unicodedata.east_asian_width(character) in {"W", "F", "A"}:
        return pixel_size
    if character in "MW@%#&":
        return pixel_size * 0.95
    return pixel_size * 0.72


def _validation_line_height_points(font_size: float) -> float:
    font = _validation_font(font_size)
    if font is None:
        return font_size * 1.5
    ascent, descent = font.getmetrics()
    metric_height = (ascent + descent) * 72 / _SCREEN_DPI
    return max(font_size * 1.45, metric_height + 1.2)


def _history_average_values(
    month: str,
    values: dict[str, object] | object,
    kind: HistoricalValueKind,
) -> tuple[float | int | None, ...]:
    averages = historical_averages(month, values, value_kind=kind)
    return tuple(
        _excel_number(period.value)
        for period in (
            averages.three_month,
            averages.six_month,
            averages.twelve_month,
            averages.comparison_year,
        )
    )


def _excel_number(value: Decimal | int | None) -> float | int | None:
    if value is None:
        return None
    if isinstance(value, int):
        return value
    return float(value)


def _optional_difference(actual: int | None, planned: int | None) -> int | None:
    if actual is None or planned is None:
        return None
    return actual - planned


def _optional_ratio(
    numerator: Decimal | int | None, denominator: Decimal | int | None
) -> Decimal | None:
    if numerator is None or denominator in (None, 0):
        return None
    return Decimal(numerator) / Decimal(denominator)


def _month_date(value: str) -> date:
    _month(value, "report_month")
    return date.fromisoformat(f"{value}-01")


def _assert_saved_value(cell: Cell, expected: object) -> None:
    if isinstance(expected, str):
        expected = _safe_display_text(expected)
    if isinstance(expected, datetime):
        if cell.value != expected or not cell.is_date:
            raise ValueError(f"{cell.coordinate} datetime changed after save")
        return
    if isinstance(expected, date):
        actual = cell.value.date() if isinstance(cell.value, datetime) else cell.value
        if actual != expected or not cell.is_date:
            raise ValueError(f"{cell.coordinate} date changed after save")
        return
    if isinstance(expected, (Decimal, int, float)) and not isinstance(expected, bool):
        _assert_numeric_equal(cell, expected)
        return
    if cell.value != expected:
        raise ValueError(f"{cell.coordinate} changed after save")


def _assert_numeric_equal(
    cell: Cell, expected: Decimal | int | float | None
) -> None:
    if expected is None:
        if cell.value is not None:
            raise ValueError(f"{cell.coordinate} must remain blank")
        return
    if cell.data_type != "n" or cell.value is None:
        raise ValueError(f"{cell.coordinate} must contain a typed number")
    if isinstance(expected, float):
        matches = math.isclose(
            float(cell.value), expected, rel_tol=1e-12, abs_tol=1e-12
        )
    else:
        matches = Decimal(str(cell.value)) == Decimal(str(expected))
    if not matches:
        raise ValueError(f"{cell.coordinate} does not reconcile to the canonical report")


def _tuple_of(value: object, item_type: type, field: str) -> None:
    if not isinstance(value, tuple) or any(not isinstance(item, item_type) for item in value):
        raise TypeError(f"{field} must be a tuple of {item_type.__name__}")


def _unique(values, label: str) -> None:
    values = tuple(values)
    if len(values) != len(set(values)):
        raise ValueError(f"{label} contains a duplicate identity")


def _month(value: object, field: str) -> None:
    if not isinstance(value, str) or not re.fullmatch(r"[0-9]{4}-(0[1-9]|1[0-2])", value):
        raise ValueError(f"{field} must use YYYY-MM")


def _text(value: object, field: str) -> None:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field} must be nonblank text")


def _positive_int(value: object, field: str) -> None:
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        raise ValueError(f"{field} must be a positive integer")


def _nonnegative_int(value: object, field: str) -> None:
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise ValueError(f"{field} must be a nonnegative integer")
