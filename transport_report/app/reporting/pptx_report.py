"""Update the fixed six-slide August report, without rebuilding its layout.

Rows are ordered display slots: 12 destinations, one derived group and one
miscellaneous row. None clears a vacant slot. Plan rows are the 12 September
destinations. The caller supplies Task6 results (including flag-aware totals),
monthly histories and explicitly selected review details. This adapter does
not infer destination identity from template text or sum derived groups twice.
"""
from __future__ import annotations

from collections.abc import Mapping
from copy import deepcopy
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
import math
import os
from pathlib import Path
import re
import tempfile
from zipfile import ZipFile

from PIL import Image, ImageFont, PngImagePlugin
from matplotlib import font_manager
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.domain.calculations import (
    DestinationCalculation, GrandTotalResult, HistoricalValueKind, MAX_SQLITE_INTEGER,
    ReviewCandidate, historical_averages, select_unit_cost_reviews,
)
from app.reporting.charts import (
    ChartReport, render_quantity_chart, render_cost_chart, render_combined_chart,
)


SOURCE_TEMPLATE_SHA256 = '7bb6269b6d3834e9eb4f7b43aa32d17aa627861d2352a29b5a7db9d7166a9cce'
SLIDE_SIZE = (9906000, 6858000)
TABLE_SIZES = {'report.monthly_table': (21, 18), 'report.review_table': (7, 8), 'report.plan_table': (19, 8)}
# Exact original shape IDs AND names, protected by the whole-file SHA above.
# Runtime targets are always the semantic names, never these IDs or text.
_SOURCE_ROLES = {
    1: {'report.title': (6, 'Text Box 5'), 'report.date': (2, 'Text Box 9')},
    2: {'report.title': (5123, '제목 1'), 'report.date': (16, 'Text Box 2'), 'report.monthly_table': (2, '표 1'), 'report.review_table': (7, '표 6')},
    3: {'report.title': (5123, '제목 1'), 'report.date': (9, 'Text Box 2'), 'report.chart': (8, '그림 7'), 'report.quantity_average': (7, 'TextBox 6')},
    4: {'report.title': (5123, '제목 1'), 'report.date': (11, 'Text Box 2'), 'report.chart': (6, '그림 5'), 'report.cost_average': (10, 'TextBox 9')},
    5: {'report.title': (5123, '제목 1'), 'report.date': (14, 'Text Box 2'), 'report.chart': (8, '그림 7'), 'report.quantity_average': (16, 'TextBox 15'), 'report.cost_average': (7, 'TextBox 6')},
    6: {'report.title': (5123, '제목 1'), 'report.date': (10, 'Text Box 2'), 'report.plan_table': (2, '표 1')},
}


@dataclass(frozen=True, slots=True)
class ReportRow:
    key: str
    label: str
    calculation: DestinationCalculation
    quantity_by_month: Mapping[str, Decimal | None]
    cost_won_by_month: Mapping[str, int | None]


@dataclass(frozen=True, slots=True)
class SalesReport:
    planned_won: int | None
    actual_won: int | None
    actual_won_by_month: Mapping[str, int | None]


@dataclass(frozen=True, slots=True)
class PlanReport:
    month: str
    rows: tuple[ReportRow | None, ...]
    nonregular: ReportRow
    total: ReportRow
    sales: SalesReport


@dataclass(frozen=True, slots=True)
class ReviewDetail:
    row_key: str
    planned_trips: int | None
    actual_trips: int | None
    standard_load: Decimal | None
    reason: str


@dataclass(frozen=True, slots=True)
class PptReport:
    report_month: str
    created_on: date
    rows: tuple[ReportRow | None, ...]
    nonregular: ReportRow
    total: ReportRow
    sales: SalesReport
    next_month: PlanReport
    charts: ChartReport
    reviews: tuple[ReviewDetail, ...] = ()


class TemplateStructureChanged(ValueError):
    code = 'TEMPLATE_STRUCTURE_CHANGED'

    def __init__(self, detail: str):
        super().__init__(f'{self.code}: {detail}')


def generate_pptx(template: str | Path, report: PptReport, output: str | Path) -> None:
    """Write a complete report atomically; source and prior output survive failure.

    An unmodified packaged source is normalized in memory. Prepared templates
    must already contain every semantic role and the fixed table topology.
    Charts are rendered by Task10 into a private temporary directory.
    """
    template, output = Path(template), Path(output)
    _different_paths(template, output)
    presentation = load_template(template)
    targets = validate_template(presentation)
    _validate_report(report)
    _validate_review_fit(presentation, targets[2, 'report.review_table'], report)
    _write(targets[1, 'report.title'].text_frame, '26년 8월 화성공장 운반비 보고')
    _write(targets[2, 'report.title'].text_frame, '▣ 26년 8월 운반비 종합')
    _write(targets[6, 'report.title'].text_frame, '▣ 26년 9월 운반비 계획')
    stamp = f'{report.created_on.year}.{report.created_on.month}.{report.created_on.day}'
    for number in range(1, 7):
        frame = targets[number, 'report.date'].text_frame
        _write_paragraph(frame.paragraphs[0], stamp)
    _monthly_table(targets[2, 'report.monthly_table'].table, report)
    _review_table(targets[2, 'report.review_table'].table, report)
    _plan_table(targets[6, 'report.plan_table'].table, report.next_month)
    quantity = _averages(report.report_month, report.charts.actual_quantity_by_month)[3]
    cost = _averages(report.report_month, report.charts.actual_cost_won_by_month, money=True)[3]
    for number in (3, 5):
        _write(targets[number, 'report.quantity_average'].text_frame, _number(quantity))
    for number in (4, 5):
        _write(targets[number, 'report.cost_average'].text_frame, _number(cost, scale=1000))
    with tempfile.TemporaryDirectory(prefix='transport-ppt-') as directory:
        for number, renderer in ((3, render_quantity_chart), (4, render_cost_chart), (5, render_combined_chart)):
            chart = Path(directory) / f'chart-{number}.png'
            renderer(report.charts, chart)
            shape = targets[number, 'report.chart']
            _fit_chart_to_crop(chart, shape)
            # Replace only this picture's image relationship. Keeping the
            # picture XML is more exact than delete/add: crop, effects, name,
            # shape ID, geometry and z-order all remain byte-identical.
            _, relationship = shape.part.get_or_add_image_part(str(chart))
            shape._pic.blipFill.blip.rEmbed = relationship
        _save_atomic(presentation, output, template)


def _fit_chart_to_crop(path, shape):
    """Keep the Task10 canvas size but fit all pixels into the legacy crop.

    Original pictures crop up to 10% off the bottom. Padding that invisible
    area prevents cropping the newly generated data table. Picture geometry
    and crop values themselves are never changed.
    """
    with Image.open(path) as original:
        width, height = original.size
        left = math.ceil(max(0, shape.crop_left) * width)
        top = math.ceil(max(0, shape.crop_top) * height)
        right = math.floor(min(1, 1 - shape.crop_right) * width)
        bottom = math.floor(min(1, 1 - shape.crop_bottom) * height)
        if right <= left or bottom <= top:
            raise TemplateStructureChanged('chart has no visible crop area')
        if (left, top, right, bottom) == (0, 0, width, height):
            return
        image = Image.new('RGB', original.size, 'white')
        image.paste(original.convert('RGB').resize((right - left, bottom - top), Image.Resampling.LANCZOS), (left, top))
        metadata = PngImagePlugin.PngInfo()
        for key, value in original.info.items():
            if isinstance(value, str):
                metadata.add_text(key, value)
    image.save(path, format='PNG', pnginfo=metadata)


def prepare_template(source: str | Path, output: str | Path) -> None:
    """Explicitly name roles on a protected copy, never on the original file."""
    source, output = Path(source), Path(output)
    _different_paths(source, output)
    presentation = load_template(source)
    _save_atomic(presentation, output, source)


def load_template(path: str | Path):
    data = Path(path).read_bytes()
    try:
        presentation = Presentation(BytesIO(data))
    except Exception as error:
        raise TemplateStructureChanged('cannot open PowerPoint package') from error
    if sha256(data).hexdigest() == SOURCE_TEMPLATE_SHA256:
        for number, roles in _SOURCE_ROLES.items():
            for role, (shape_id, old_name) in roles.items():
                matches = [s for s in presentation.slides[number - 1].shapes if s.shape_id == shape_id and s.name == old_name]
                if len(matches) != 1:
                    raise TemplateStructureChanged(f'slide {number}: source role {role} is ambiguous')
                matches[0].name = role
    validate_template(presentation)
    return presentation


def validate_template(presentation):
    """Validate every required role and merged-cell topology before any edits."""
    if len(presentation.slides) != 6 or (presentation.slide_width, presentation.slide_height) != SLIDE_SIZE:
        raise TemplateStructureChanged('expected six slides at the original slide size')
    targets = {}
    for number, roles in _SOURCE_ROLES.items():
        slide = presentation.slides[number - 1]
        for role in roles:
            matches = [s for s in slide.shapes if s.name == role]
            if len(matches) != 1:
                raise TemplateStructureChanged(f'slide {number}: expected one {role}, found {len(matches)}')
            shape = matches[0]
            if role in TABLE_SIZES:
                if not shape.has_table:
                    raise TemplateStructureChanged(f'slide {number}: {role} is not a table')
                table = shape.table
                if (len(table.rows), len(table.columns)) != TABLE_SIZES[role]:
                    raise TemplateStructureChanged(f'slide {number}: {role} row/column count changed')
                if any(len(row.cells) != TABLE_SIZES[role][1] for row in table.rows):
                    raise TemplateStructureChanged(f'slide {number}: {role} contains an incomplete row')
                if _merge_map(table) != _expected_merges(role):
                    raise TemplateStructureChanged(f'slide {number}: {role} merged cells changed')
            elif role == 'report.chart':
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    raise TemplateStructureChanged(f'slide {number}: chart is not a picture')
            elif not shape.has_text_frame:
                raise TemplateStructureChanged(f'slide {number}: {role} is not text')
            targets[number, role] = shape
    return targets


def _merge_map(table):
    return {(r, c): (cell.span_height, cell.span_width) for r, row in enumerate(table.rows) for c, cell in enumerate(row.cells) if cell.is_merge_origin}


def _expected_merges(role):
    if role == 'report.review_table':
        return {(0, 0): (2, 1), (0, 1): (1, 7)}
    if role == 'report.plan_table':
        return {(0, 0): (2, 1), (0, 1): (1, 3), (0, 4): (1, 4), (14, 1): (1, 2), (14, 3): (2, 1), **{(r, 1): (1, 3) for r in (16, 17, 18)}}
    result = {(0, 0): (2, 2), (0, 2): (1, 3), (0, 5): (1, 3), (0, 8): (1, 6), (0, 14): (1, 4), (11, 0): (4, 1), (16, 5): (1, 3), (17, 12): (1, 2)}
    result.update({(1, c): (1, 2) for c in (8, 10, 12)})
    result.update({(r, 0): (1, 2) for r in (*range(2, 11), *range(16, 21))})
    result.update({(r, c): (1, w) for r in (15, 16) for c, w in ((2, 3), (8, 2), (12, 2))})
    result.update({(r, c): (1, 3) for r in (18, 19, 20) for c in (2, 5, 8, 11)})
    return result


def _different_paths(source, output):
    if source.resolve() == output.resolve() or (output.exists() and os.path.samefile(source, output)):
        raise ValueError('output must be separate from the source template')


def _save_atomic(presentation, output, template):
    # Import here so the standalone verifier can also import template helpers.
    from tools.verify_ppt_layout import verify_ppt_layout
    descriptor, name = tempfile.mkstemp(prefix=f'.{output.name}.', suffix='.pptx', dir=output.parent)
    os.close(descriptor)
    temporary = Path(name)
    try:
        serialized = BytesIO()
        presentation.save(serialized)
        # python-pptx rewrites legacy XML declarations/relationship order.
        # These parts are deliberately unchanged, so preserve their exact bytes.
        with ZipFile(template) as original, ZipFile(serialized) as generated, ZipFile(temporary, 'w') as package:
            for part in generated.infolist():
                preserve = part.filename.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/theme/'))
                value = original.read(part.filename) if preserve else generated.read(part.filename)
                package.writestr(part, value)
        problems = verify_ppt_layout(template, temporary)
        if problems:
            raise TemplateStructureChanged('; '.join(problems))
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def _validate_report(report):
    if not isinstance(report, PptReport) or report.report_month != '2026-08':
        raise ValueError('PPT generation supports only the August 2026 PptReport')
    if not isinstance(report.created_on, date):
        raise ValueError('created_on must be a date')
    if not isinstance(report.next_month, PlanReport) or report.next_month.month != '2026-09':
        raise ValueError('September 2026 next-month plan is required')
    for rows, expected in ((report.rows, 14), (report.next_month.rows, 12)):
        if len(rows) != expected:
            raise ValueError(f'exactly {expected} row slots are required; use None for vacant slots')
        keys = [row.key for row in rows if row is not None]
        if len(keys) != len(set(keys)):
            raise ValueError('duplicate report row key')
        identities = [row.calculation.provenance for row in rows if row is not None]
        if len(identities) != len(set(identities)):
            raise ValueError('duplicate calculation identity')
        for row in rows:
            if row is not None and (not isinstance(row.calculation, DestinationCalculation) or not row.label.strip()):
                raise ValueError('rows require canonical calculations and a label')
    for period in (report, report.next_month):
        if not isinstance(period.total.calculation, GrandTotalResult):
            raise ValueError('total requires a Task6 GrandTotalResult')
    for row in (*report.next_month.rows, report.next_month.nonregular, report.next_month.total):
        if row is not None and (row.calculation.planned_quantity is None or row.calculation.planned_cost_won is None):
            raise ValueError('next-month plan values are missing')
    if report.next_month.sales.planned_won is None:
        raise ValueError('next-month sales plan is missing')
    if type(report.sales.actual_won) is not int or not 0 <= report.sales.actual_won <= MAX_SQLITE_INTEGER:
        raise ValueError('current actual sales must be present nonnegative integer won within the supported range')
    if report.charts.report_month != report.report_month:
        raise ValueError('chart month conflicts with report')
    for mapping, field in ((report.charts.planned_quantity_by_month, 'planned_quantity'), (report.charts.actual_quantity_by_month, 'actual_quantity'), (report.charts.planned_cost_won_by_month, 'planned_cost_won'), (report.charts.actual_cost_won_by_month, 'actual_cost_won')):
        if mapping.get(report.report_month) != getattr(report.total.calculation, field):
            raise ValueError(f'chart current total conflicts with {field}')
    _validate_actual_history(report)
    _selected_reviews(report)


def _validate_actual_history(report):
    """All consumers must agree on each identity/measure/month, including None.

    Current canonical Task6 actuals anchor August. History maps may be sparse,
    but an overlapping observation cannot be silently replaced or preferred.
    Totals have one report-wide identity even if next-month membership changes.
    """
    observations = {}
    def accept(identity, measure, mapping, source):
        for month, value in mapping.items():
            key = identity, measure, month
            if key in observations and observations[key][0] != value:
                raise ValueError(f'actual history conflict: {identity}, {measure}, {month}: {observations[key][1]} / {source}')
            observations[key] = value, source

    for period in (report, report.next_month):
        for row in (*period.rows, period.nonregular, period.total):
            if row is None:
                continue
            if row is period.total:
                identity = ('total',)
            elif row is period.nonregular:
                identity = ('nonregular',)
            elif getattr(row.calculation, 'destination_id', None) is not None:
                identity = ('destination', row.calculation.destination_id)
            elif getattr(row.calculation, 'group_id', None) is not None:
                identity = ('group', row.calculation.group_id)
            else:
                identity = ('row', row.key)
            for measure, mapping, actual in (
                ('quantity', row.quantity_by_month, row.calculation.actual_quantity),
                ('cost_won', row.cost_won_by_month, row.calculation.actual_cost_won),
            ):
                accept(identity, measure, mapping, f'{type(period).__name__}.{row.key}')
                if period is report:
                    accept(identity, measure, {report.report_month: actual}, 'Task6 current actual')
    accept(('total',), 'quantity', report.charts.actual_quantity_by_month, 'Task10 chart/callout')
    accept(('total',), 'cost_won', report.charts.actual_cost_won_by_month, 'Task10 chart/callout')


def _averages(month, values, *, money=False):
    history = historical_averages(month, values, value_kind=HistoricalValueKind.MONEY if money else HistoricalValueKind.QUANTITY)
    return tuple(getattr(history, name).value for name in ('three_month', 'six_month', 'twelve_month', 'comparison_year'))


def _number(value, *, scale=1, places=0, percent=False):
    if value is None:
        return ''
    value = Decimal(value) / Decimal(scale)
    if percent:
        value *= 100
    rounded = value.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    if rounded == 0:
        rounded = abs(rounded)
    return f'{rounded:,.{places}f}' + ('%' if percent else '')


def _ratio(numerator, denominator):
    return None if numerator is None or denominator in (None, 0) else Decimal(numerator) / Decimal(denominator)


def _difference(actual, planned):
    return None if actual is None or planned is None else actual - planned


class _TableUpdate:
    """Stage a complete body so every cell is replaced exactly once.

    Clearing populated cells first would lose the original lengths that keep
    mixed-format run boundaries stable during replacement.
    """
    def __init__(self, table):
        self.table = table
        self.values = {(r, c): '' for r, row in enumerate(table.rows) if r >= 2 for c, cell in enumerate(row.cells) if not cell.is_spanned}

    def apply(self):
        for (row, column), text in self.values.items():
            _write(self.table.cell(row, column).text_frame, text)


def _put(update, row, column, text):
    cell = update.table.cell(row, column)
    if cell.is_spanned:
        if text:
            raise TemplateStructureChanged(f'cannot write merged-away cell {row},{column}')
        return
    update.values[row, column] = text


def _history_cells(table, row, start, values, *, scale=1000, percent=False, places=0):
    for column, value in enumerate(values, start):
        _put(table, row, column, _number(value, scale=scale, percent=percent, places=places))


def _monthly_table(table, report):
    table = _TableUpdate(table)
    for index, row in enumerate(report.rows, 2):
        if row is None:
            continue
        _put(table, index, 0 if index < 11 else 1, row.label)
        calc = row.calculation
        values = (calc.planned_quantity, calc.planned_cost_won, calc.planned_unit_cost, calc.actual_quantity, calc.actual_cost_won, calc.actual_unit_cost, calc.quantity_variance, calc.quantity_variance_pct, calc.cost_variance_won, calc.cost_variance_pct, calc.actual_unit_cost_variance, calc.actual_unit_cost_variance_pct)
        for column, value in enumerate(values, 2):
            # The miscellaneous slot has no planned/quantity variance/unit cells.
            if index == 15 and column in (2, 3, 4, 8, 9, 12, 13):
                continue
            _put(table, index, column, _number(value, scale=1000 if column in (3, 6, 10) else 1, places=1 if column in (9, 11, 13) else 0, percent=column in (9, 11, 13)))
        _history_cells(table, index, 14, _averages(report.report_month, row.cost_won_by_month, money=True))
    extra = report.nonregular.calculation
    _put(table, 16, 0, report.nonregular.label)
    for column, value in ((2, extra.planned_cost_won), (5, extra.actual_cost_won), (10, extra.cost_variance_won)):
        _put(table, 16, column, _number(value, scale=1000))
    _put(table, 16, 11, _number(extra.cost_variance_pct, percent=True, places=1))
    _history_cells(table, 16, 14, _averages(report.report_month, report.nonregular.cost_won_by_month, money=True))
    total = report.total.calculation
    _put(table, 17, 0, '합계')
    for column, field in ((2, 'planned_quantity'), (3, 'planned_cost_won'), (5, 'actual_quantity'), (6, 'actual_cost_won'), (8, 'quantity_variance'), (9, 'quantity_variance_pct'), (10, 'cost_variance_won'), (11, 'cost_variance_pct')):
        _put(table, 17, column, _number(getattr(total, field), scale=1000 if column in (3, 6, 10) else 1, places=1 if column in (9, 11) else 0, percent=column in (9, 11)))
    cost_history = _averages(report.report_month, report.total.cost_won_by_month, money=True)
    quantity_history = _averages(report.report_month, report.total.quantity_by_month)
    sales_history = _averages(report.report_month, report.sales.actual_won_by_month, money=True)
    _history_cells(table, 17, 14, cost_history)
    unit_history = tuple(_ratio(c, q) for c, q in zip(cost_history, quantity_history))
    _summary_row(table, 18, '총 대당 운반비', total.planned_unit_cost, total.actual_unit_cost, unit_history, variance=total.actual_unit_cost_variance, variance_pct=total.actual_unit_cost_variance_pct)
    _summary_row(table, 19, '매출액', report.sales.planned_won, report.sales.actual_won, sales_history, scale=1000)
    ratios = tuple(_ratio(c, s) for c, s in zip(cost_history, sales_history))
    _summary_row(table, 20, '매출액 대비 운반비', _ratio(total.planned_cost_won, report.sales.planned_won), _ratio(total.actual_cost_won, report.sales.actual_won), ratios, percent=True)
    table.apply()


def _summary_row(table, row, label, plan, actual, history, *, scale=1, percent=False, variance=None, variance_pct=None):
    _put(table, row, 0, label)
    delta = _difference(actual, plan) if variance is None else variance
    relative = _ratio(delta, plan) if variance_pct is None else variance_pct
    for column, value in ((2, plan), (5, actual), (8, delta)):
        _put(table, row, column, _number(value, scale=scale, percent=percent, places=2 if percent else 0))
    _put(table, row, 11, _number(relative, percent=True, places=1))
    _history_cells(table, row, 14, history, scale=scale, percent=percent, places=1 if percent else 0)


def _plan_table(table, plan):
    table = _TableUpdate(table)
    for index, row in enumerate(plan.rows, 2):
        if row is None:
            continue
        _put(table, index, 0, row.label)
        for column, value in ((1, row.calculation.planned_quantity), (2, row.calculation.planned_cost_won), (3, row.calculation.planned_unit_cost)):
            _put(table, index, column, _number(value, scale=1000 if column == 2 else 1))
        _history_cells(table, index, 4, _averages(plan.month, row.cost_won_by_month, money=True))
    _put(table, 14, 0, plan.nonregular.label)
    _put(table, 14, 1, _number(plan.nonregular.calculation.planned_cost_won, scale=1000))
    _history_cells(table, 14, 4, _averages(plan.month, plan.nonregular.cost_won_by_month, money=True))
    total = plan.total.calculation
    _put(table, 15, 0, '합계')
    _put(table, 15, 1, _number(total.planned_quantity))
    _put(table, 15, 2, _number(total.planned_cost_won, scale=1000))
    costs = _averages(plan.month, plan.total.cost_won_by_month, money=True)
    quantities = _averages(plan.month, plan.total.quantity_by_month)
    sales = _averages(plan.month, plan.sales.actual_won_by_month, money=True)
    _history_cells(table, 15, 4, costs)
    for row, label, value, history, scale, percent in (
        (16, '총 대당 운반비', total.planned_unit_cost, tuple(_ratio(c, q) for c, q in zip(costs, quantities)), 1, False),
        (17, '매출액', plan.sales.planned_won, sales, 1000, False),
        (18, '매출액 대비 운반비', _ratio(total.planned_cost_won, plan.sales.planned_won), tuple(_ratio(c, s) for c, s in zip(costs, sales)), 1, True),
    ):
        _put(table, row, 0, label)
        _put(table, row, 1, _number(value, scale=scale, percent=percent, places=1 if percent else 0))
        _history_cells(table, row, 4, history, scale=scale, percent=percent, places=1 if percent else 0)
    table.apply()


def _selected_reviews(report):
    rows = {row.key: row for row in report.rows if row is not None}
    candidates = [ReviewCandidate(row.calculation.destination_id, row.label, index, row.calculation) for index, row in enumerate(rows.values()) if getattr(row.calculation, 'destination_id', None) is not None]
    eligible_ids = {item.destination_id for item in select_unit_cost_reviews(candidates).automatic_items}
    keys = [detail.row_key for detail in report.reviews]
    if len(keys) > 5 or len(set(keys)) != len(keys):
        raise ValueError('review table permits at most five distinct selected rows')
    selected = []
    for detail in report.reviews:
        row = rows.get(detail.row_key)
        if row is None:
            raise ValueError('review row is missing from monthly report')
        calculation = row.calculation
        destination_id = getattr(calculation, 'destination_id', None)
        eligible = destination_id in eligible_ids if destination_id is not None else (calculation.actual_unit_cost_variance_pct is not None and abs(calculation.actual_unit_cost_variance_pct) >= Decimal('.15'))
        if not eligible:
            raise ValueError('selected review does not meet the inclusive +/-15% threshold')
        for trips in (detail.planned_trips, detail.actual_trips):
            if trips is not None and (type(trips) is not int or trips < 0):
                raise ValueError('review trips must be nonnegative integers or missing')
        if detail.standard_load is not None and (not isinstance(detail.standard_load, Decimal) or not detail.standard_load.is_finite() or detail.standard_load < 0):
            raise ValueError('review standard load must be a nonnegative Decimal')
        selected.append((row, detail))
    return selected


def _review_table(table, report):
    table = _TableUpdate(table)
    for index, (row, detail) in enumerate(_selected_reviews(report), 2):
        for column, value in enumerate(_review_values(row, detail)):
            _put(table, index, column, value)
    table.apply()


def _review_values(row, detail):
    quantity = row.calculation.actual_quantity
    load = _ratio(quantity, detail.actual_trips)
    return (row.label, _number(detail.planned_trips), _number(detail.actual_trips), _number(quantity), _number(load), _number(detail.standard_load), _number(_ratio(load, detail.standard_load), percent=True, places=1), detail.reason)


def _validate_review_fit(presentation, shape, report):
    """Conservative fixed-row preflight; never resize or alter text formatting.

    Font metrics are measured at 4 pixels/point. Unknown/theme fonts use a
    conservative em bound instead of assuming a narrower substitute. Existing
    empty paragraphs also occupy height, just as they do in native PowerPoint.
    """
    table = shape.table
    defaults = presentation._element.xpath('./p:defaultTextStyle/a:lvl1pPr')
    defaults += presentation.slides[1].slide_layout.slide_master._element.xpath('./p:txStyles/p:otherStyle/a:lvl1pPr')
    for row_index, (row, detail) in enumerate(_selected_reviews(report), 2):
        for column, text in enumerate(_review_values(row, detail)):
            if not text:
                continue
            cell = table.cell(row_index, column)
            width = (table.columns[column].width - cell.margin_left - cell.margin_right) / 12700
            available = (table.rows[row_index].height - cell.margin_top - cell.margin_bottom) / 12700
            paragraphs = cell.text_frame.paragraphs
            lines = text.split('\n')
            height = 0.0
            fits = width > 0 and available > 0 and '\t' not in text and '\v' not in text
            for index in range(max(len(lines), len(paragraphs))):
                paragraph = paragraphs[index] if index < len(paragraphs) else paragraphs[0]
                line = lines[index] if index < len(lines) else ''
                props = list(defaults) + cell._tc.xpath('./a:txBody/a:lstStyle/a:lvl1pPr')
                if paragraph._p.pPr is not None:
                    props.append(paragraph._p.pPr)
                attrs, fonts = {}, {}
                for prop in props:
                    for child in prop.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr'):
                        attrs.update(child.attrib)
                        fonts.update({el.tag.rsplit('}', 1)[-1]: el.get('typeface') for el in child})
                direct = paragraph._p.xpath('./a:endParaRPr | ./a:r/a:rPr')
                sizes = [float(attrs.get('sz', 1800)) / 100]
                for prop in direct:
                    attrs.update(prop.attrib)
                    sizes.append(float(attrs.get('sz', 1800)) / 100)
                    fonts.update({el.tag.rsplit('}', 1)[-1]: el.get('typeface') for el in prop})
                # Direct run sizes override inherited defaults, not vice versa.
                size = max(sizes[1:] or sizes)
                family = (fonts.get('ea') or fonts.get('latin')) if any(ord(char) >= 0x2E80 for char in line) else fonts.get('latin')
                font = None
                if family and not family.startswith('+'):
                    try:
                        path = font_manager.findfont(font_manager.FontProperties(family=family, weight='bold' if attrs.get('b') == '1' else 'normal'), fallback_to_default=False)
                        font = ImageFont.truetype(path, max(1, math.ceil(size * 4)))
                    except (ValueError, OSError):
                        pass
                def advance(token):
                    if font:
                        # Leave headroom for native shaping, and do not trust
                        # missing-glyph boxes to measure Korean text narrowly.
                        bound = sum(size if ord(char) >= 0x2E80 else 0 for char in token)
                        return max(bound, font.getlength(token) / 4) * 1.15
                    return sum(.5 if char.isspace() else 1.2 for char in token) * size
                indent = max((float(prop.get('marL', 0)) / 12700 for prop in props), default=0)
                usable = width - max(0, indent)
                count, used = 1, 0.0
                for token in re.findall(r'[A-Za-z0-9]+|.', line):
                    extent = advance(token)
                    if usable <= 0 or extent > usable:
                        fits = False
                        break
                    if used + extent > usable:
                        count += 1
                        used = 0
                    used += extent
                if count > 1 and cell.text_frame.word_wrap is False:
                    fits = False
                def spacing(name, fallback):
                    result = fallback
                    for prop in props:
                        children = prop.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}' + name)
                        if children and len(children[0]):
                            spec = children[0][0]
                            result = float(spec.get('val')) / (100 if spec.tag.endswith('spcPts') else 100000) * (1 if spec.tag.endswith('spcPts') else size * 1.2)
                    return result
                height += max(size * 1.2, spacing('lnSpc', size * 1.2)) * count + spacing('spcBef', 0) + spacing('spcAft', 0)
            if not fits or height > available + .01:
                raise ValueError(f'REPORT_TEXT_DOES_NOT_FIT: 2번 슬라이드 검토표 {row_index + 1}행 {column + 1}열의 문구가 고정 셀 크기를 초과합니다. 문구를 줄여 주세요.')


def _write(frame, text):
    # Reuse destination paragraphs/runs, including their direct formatting.
    lines = text.split('\n')
    for index, paragraph in enumerate(frame.paragraphs):
        if index < len(lines):
            _write_paragraph(paragraph, lines[index])
        else:
            # Existing blank paragraphs still participate in PowerPoint's
            # minimum table-row height, so retain them and their formatting.
            _write_paragraph(paragraph, '')
    for line in lines[len(frame.paragraphs):]:
        paragraph = frame.add_paragraph()
        source = frame.paragraphs[0]
        if source._p.pPr is not None:
            paragraph._p.insert(0, deepcopy(source._p.pPr))
        _write_paragraph(paragraph, line, source=source)


def _write_paragraph(paragraph, text, *, source=None):
    for line_break in paragraph._p.xpath('./a:br'):
        paragraph._p.remove(line_break)
    runs = list(paragraph.runs)
    if not runs:
        if not text:
            return
        run = paragraph.add_run()
        # Copy direct formatting only when it exists. Otherwise leave the
        # new run unformatted so its destination's master/table style applies.
        properties = paragraph._p.find('{http://schemas.openxmlformats.org/drawingml/2006/main}endParaRPr')
        if properties is None and paragraph._p.pPr is not None:
            properties = paragraph._p.pPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr')
        if properties is None and source is not None and source.runs:
            properties = source.runs[0]._r.rPr
        if properties is not None:
            copied = deepcopy(properties)
            copied.tag = '{http://schemas.openxmlformats.org/drawingml/2006/main}rPr'
            run._r.insert(0, copied)
        runs = [run]
    if not text and runs:
        # PowerPoint 2007 ignores the font on an empty run and falls back to
        # 18pt when endParaRPr is absent, expanding otherwise unchanged rows.
        if not paragraph._p.xpath('./a:endParaRPr') and runs[0]._r.rPr is not None:
            end = deepcopy(runs[0]._r.rPr)
            end.tag = '{http://schemas.openxmlformats.org/drawingml/2006/main}endParaRPr'
            paragraph._p.append(end)
    if runs and all(not run.text for run in runs):
        runs[0].text = text
        return
    remaining = text
    for index, run in enumerate(runs):
        size = len(run.text) if index < len(runs) - 1 else len(remaining)
        run.text, remaining = remaining[:size], remaining[size:]
