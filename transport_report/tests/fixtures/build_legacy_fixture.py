from __future__ import annotations

from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook


HISTORY_SHEET = "누적 데이터"
PLAN_SHEET = "26년 월계획"
DESTINATIONS = (
    "Alpha",
    "당진",
    "South A",
    "South B",
    *(f"Extra {number:02d}" for number in range(1, 11)),
)


def build_legacy_fixture(
    path: Path,
    *,
    unknown_alias: bool = False,
    formula_without_cache: bool = False,
    history_formula_without_cache: bool = False,
    total_mismatch: bool = False,
    malformed: tuple[str, object] | None = None,
    omit_history: tuple[str, int, int] | None = None,
    missing_actual_quantity: tuple[str, int, int] | None = None,
    zero_current: bool = False,
) -> Path:
    workbook = Workbook()
    history = workbook.active
    history.title = HISTORY_SHEET
    history.append(("연도", "월", "연도+월", "납품처", "실적수량", "실적운반비"))

    row_number = 2
    source_destinations = (
        ("Mystery", *DESTINATIONS[1:]) if unknown_alias else DESTINATIONS
    )
    history_cache: dict[str, str | None] = {}
    for year, last_month in ((2025, 12), (2026, 8)):
        for month in range(1, last_month + 1):
            for offset, destination in enumerate(source_destinations, start=1):
                if omit_history == (destination, year, month):
                    continue
                quantity: object = year % 100 * 100 + month * 10 + offset
                cost: object = int(quantity) * 100
                if zero_current and destination == "Alpha" and (year, month) == (2026, 8):
                    quantity = 0
                    cost = 0
                if missing_actual_quantity == (destination, year, month):
                    quantity = None
                history.cell(row_number, 1, f"{year % 100}년")
                history.cell(row_number, 2, f"{month}월")
                history.cell(row_number, 3, f"=A{row_number}&B{row_number}")
                history.cell(row_number, 4, destination)
                if row_number == 2 and history_formula_without_cache:
                    history.cell(row_number, 5, f"={quantity}")
                    history_cache[f"E{row_number}"] = None
                else:
                    history.cell(row_number, 5, quantity)
                history.cell(row_number, 6, cost)
                row_number += 1

    history["H1"] = "차트 보조"
    history["H2"] = "Mystery helper"
    history["I2"] = 999_999

    plan = workbook.create_sheet(PLAN_SHEET)
    plan["AC3"] = "단위 : 대 / 원"
    plan["B4"] = "납품처"
    for month in range(1, 13):
        quantity_column = 3 + (month - 1) * 2
        plan.cell(4, quantity_column, f"{month}월")
        plan.cell(5, quantity_column, "수량")
        plan.cell(5, quantity_column + 1, "운반비")

    formula_cache: dict[str, str | None] = {}
    for row, destination in enumerate(source_destinations, start=6):
        plan.cell(row, 2, destination)
        for month in range(1, 9):
            quantity_column = 3 + (month - 1) * 2
            quantity = month * 10 + row
            cost = quantity * 1000
            if month == 8 and destination == "Alpha":
                plan.cell(row, quantity_column, "=10+76")
                formula_cache[plan.cell(row, quantity_column).coordinate] = (
                    None if formula_without_cache else str(quantity)
                )
            else:
                plan.cell(row, quantity_column, quantity)
            plan.cell(row, quantity_column + 1, cost)

    total_row = 6 + len(source_destinations)
    plan.cell(total_row, 2, "합계")
    august_quantity_column = 17
    included_quantity_rows = [
        row
        for row, destination in enumerate(source_destinations, start=6)
        if destination != "당진"
    ]
    plan.cell(
        total_row,
        august_quantity_column,
        "=" + "+".join(f"Q{row}" for row in included_quantity_rows),
    )
    plan.cell(total_row, august_quantity_column + 1, f"=SUM(R6:R{total_row - 1})")
    formula_cache[f"Q{total_row}"] = str(
        sum(80 + row for row in included_quantity_rows) + int(total_mismatch)
    )
    formula_cache[f"R{total_row}"] = str(
        sum((80 + row) * 1000 for row in range(6, total_row))
    )

    # Anything below the total row is a chart/helper range and must not migrate.
    plan.cell(total_row + 5, 2, "Mystery helper")
    plan.cell(total_row + 5, 17, 123_456)
    plan.cell(total_row + 5, 18, 654_321)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    if history_cache:
        _patch_formula_cached_values(path, sheet_number=1, values=history_cache)
    _patch_formula_cached_values(path, sheet_number=2, values=formula_cache)

    if malformed is not None:
        field, value = malformed
        workbook = load_workbook(path)
        cell = {"year": "A2", "month": "B2", "quantity": "E2", "cost": "F2"}[field]
        workbook[HISTORY_SHEET][cell] = value
        workbook.save(path)
        workbook.close()
    return path


def change_source_value(path: Path) -> None:
    workbook = load_workbook(path)
    workbook[HISTORY_SHEET]["E2"] = int(workbook[HISTORY_SHEET]["E2"].value) + 1
    workbook.save(path)
    workbook.close()


def _patch_formula_cached_values(
    path: Path, *, sheet_number: int, values: dict[str, str | None]
) -> None:
    worksheet_name = f"xl/worksheets/sheet{sheet_number}.xml"
    with ZipFile(path, "r") as source:
        members = [(info, source.read(info.filename)) for info in source.infolist()]

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ElementTree.register_namespace("", namespace)
    rewritten: list[tuple[object, bytes]] = []
    for info, data in members:
        if info.filename == worksheet_name:
            root = ElementTree.fromstring(data)
            cells = {
                cell.attrib["r"]: cell
                for cell in root.findall(f".//{{{namespace}}}c")
            }
            for coordinate, cached_value in values.items():
                value_element = cells[coordinate].find(f"{{{namespace}}}v")
                if value_element is None:
                    value_element = ElementTree.SubElement(
                        cells[coordinate], f"{{{namespace}}}v"
                    )
                value_element.text = cached_value
            data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
        rewritten.append((info, data))

    temporary = path.with_name(path.name + ".rewritten")
    with ZipFile(temporary, "w", compression=ZIP_DEFLATED) as destination:
        for info, data in rewritten:
            destination.writestr(info, data)
    temporary.replace(path)
