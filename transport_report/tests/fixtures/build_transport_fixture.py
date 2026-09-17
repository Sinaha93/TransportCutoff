from __future__ import annotations

from datetime import date
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook


REGULAR_SHEET = "화성운반비내역(8월)"
SUBCONTRACT_SHEETS = (
    "용차1-OK로지웰",
    "용차2-대원로지스틱",
    "용차3-정동물류",
)


def build_transport_fixture(
    path: Path, *, subtotal_mismatch: bool = False
) -> Path:
    workbook = Workbook()
    regular = workbook.active
    regular.title = REGULAR_SHEET
    _write_regular_header(regular, 2)

    regular.cell(3, 1, "5톤 / 기사A")
    regular.cell(3, 2, "Known Plant")
    regular.cell(3, 14, 1.25)  # day 12 (C is day 1)
    regular.cell(3, 34, 1.25)  # AH: total count
    regular.cell(3, 36, 2_000)  # AJ: unit cost
    regular.cell(3, 37, 2_501 if subtotal_mismatch else 2_500)  # AK: subtotal

    _write_regular_header(regular, 4)
    regular.row_dimensions[4].hidden = True

    regular.cell(5, 2, "Unknown Plant")
    regular.cell(5, 3, 1)
    regular.cell(5, 34, 1)
    regular.cell(5, 36, 3_000)
    regular.cell(5, 37, 3_000)

    for offset, sheet_name in enumerate(SUBCONTRACT_SHEETS, start=5):
        sheet = workbook.create_sheet(sheet_name)
        sheet.cell(1, 1, "일자")
        sheet.cell(1, 2, "운반지역")
        sheet.cell(1, 8, "톤수")
        sheet.cell(1, 9, "금액")
        sheet.cell(2, 1, date(2026, 8, offset))
        sheet.cell(2, 2, "Known Plant")
        sheet.cell(2, 8, f"{offset}톤")
        sheet.cell(2, 9, offset * 10_000)
        sheet.cell(3, 2, "합계")
        sheet.cell(3, 9, offset * 10_000)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def build_structural_clone_fixture(
    path: Path, *, month: int = 8, include_inactive_regular_rows: bool = False
) -> Path:
    workbook = Workbook()
    regular = workbook.active
    regular.title = f"화성운반비내역({month}월)"
    _write_structural_clone_regular_header(regular, 2)

    regular.cell(3, 1, "5톤 / 기사A")
    regular.cell(3, 2, "Known Plant")
    regular.cell(3, 14, 1.25)
    regular.cell(3, 34, "=SUM(C3:AG3)")
    regular.cell(3, 36, 2_000)
    regular.cell(3, 37, "=AH3*AJ3")

    _write_structural_clone_regular_header(regular, 4)
    regular.row_dimensions[4].hidden = True

    regular.cell(5, 2, "Known Plant")
    regular.cell(5, 3, 1)
    regular.cell(5, 34, "=SUM(C5:AG5)")
    regular.cell(5, 36, 3_000)
    regular.cell(5, 37, "=AH5*AJ5")
    regular.row_dimensions[5].hidden = True

    total_label_columns = (1, 7, 1)
    cached_values: dict[int, dict[str, str | None]] = {
        1: {"AH3": "1.25", "AK3": "2500", "AH5": "1", "AK5": "3000"}
    }
    if include_inactive_regular_rows:
        regular.cell(6, 1, "8-ton / Driver B")
        regular.cell(6, 2, "Dormant Blank Destination")
        regular.cell(6, 34, 0)
        regular.cell(6, 37, 0)

        regular.cell(7, 2, "Dormant Zero Destination")
        for column in range(3, 34):
            regular.cell(7, column, 0)
        regular.cell(7, 34, 0)
        regular.cell(7, 37, 0)

        regular.cell(8, 2, "Known Plant")
        regular.cell(8, 3, 1)
        regular.cell(8, 34, "=SUM(C8:AG8)")
        regular.cell(8, 36, 4_000)
        regular.cell(8, 37, "=AH8*AJ8")
        cached_values[1].update({"AH8": "1", "AK8": "4000"})
    for sheet_number, (sheet_name, total_column) in enumerate(
        zip(SUBCONTRACT_SHEETS, total_label_columns, strict=True), start=2
    ):
        sheet = workbook.create_sheet(sheet_name)
        sheet.cell(2, 1, "No")
        sheet.cell(2, 2, "오더일자")
        sheet.cell(2, 3, "의뢰담당")
        sheet.cell(2, 4, "출발지")
        sheet.cell(2, 5, "출발동")
        sheet.cell(2, 6, "도착지")
        sheet.cell(2, 7, "도착동")
        sheet.cell(2, 8, "차종(t)" if sheet_number == 2 else "차종(톤)")
        sheet.cell(2, 9, "기본요금")
        sheet.cell(2, 10, "처리기사")
        sheet.cell(3, 1, 1)
        sheet.cell(3, 2, date(2026, month, 5 + sheet_number))
        sheet.cell(3, 6, "Known Plant")
        sheet.cell(3, 8, 5)
        amount = (3 + sheet_number) * 10_000
        sheet.cell(3, 9, f"={amount}" if sheet_number == 2 else amount)
        sheet.cell(4, total_column, "합계")
        sheet.cell(4, 9, "=SUM(I3:I3)")
        cached_values[sheet_number] = {"I4": str(amount)}
        if sheet_number == 2:
            cached_values[sheet_number]["I3"] = str(amount)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    for sheet_number, values in cached_values.items():
        patch_formula_cached_values(path, sheet_number, values)
    return path


def patch_formula_cached_values(
    path: Path, sheet_number: int, values: dict[str, str | None]
) -> None:
    worksheet_name = f"xl/worksheets/sheet{sheet_number}.xml"
    with ZipFile(path, "r") as source:
        members = [(info, source.read(info.filename)) for info in source.infolist()]

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ElementTree.register_namespace("", namespace)
    rewritten: list[tuple[object, bytes]] = []
    for info, data in members:
        if info.filename == worksheet_name:
            root = ElementTree.fromstring(data)
            cells = {
                cell.attrib["r"]: cell
                for cell in root.findall(f".//{{{namespace}}}c")
            }
            for coordinate, cached_value in values.items():
                value_element = cells[coordinate].find(f"{{{namespace}}}v")
                if value_element is None:
                    value_element = ElementTree.SubElement(
                        cells[coordinate], f"{{{namespace}}}v"
                    )
                value_element.text = cached_value
            data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
        rewritten.append((info, data))

    temporary = path.with_name(path.name + ".rewritten")
    with ZipFile(temporary, "w", compression=ZIP_DEFLATED) as destination:
        for info, data in rewritten:
            destination.writestr(info, data)
    temporary.replace(path)


def _write_regular_header(sheet, row: int) -> None:
    sheet.cell(row, 1, "차량/기사")
    sheet.cell(row, 2, "운반지역")
    for day in range(1, 32):
        sheet.cell(row, day + 2, day)
    sheet.cell(row, 34, "총회수")
    sheet.cell(row, 36, "단가")
    sheet.cell(row, 37, "소계")


def _write_structural_clone_regular_header(sheet, row: int) -> None:
    _write_regular_header(sheet, row)
    sheet.cell(row, 1, "구분")
    sheet.cell(row, 2, "날짜")
    sheet.cell(row, 34, "계")
    sheet.cell(row, 38, "합계")
