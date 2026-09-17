from __future__ import annotations

import hashlib
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from decimal import Decimal
from importlib import import_module
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.db import Database
from app.repositories.masters import MasterRepository
from app.repositories.monthly_inputs import MonthLockedError, MonthlyInputRepository
from tests.fixtures.build_transport_fixture import (
    REGULAR_SHEET,
    SUBCONTRACT_SHEETS,
    build_structural_clone_fixture,
    build_transport_fixture,
    patch_formula_cached_values,
)


@pytest.fixture
def api():
    try:
        parser_module = import_module("app.importers.hwaseong_workbook")
        repository_module = import_module("app.repositories.transport_entries")
        service_module = import_module("app.services.import_service")
    except ModuleNotFoundError as error:
        pytest.fail(f"transport import API is not implemented: {error}")
    return parser_module, repository_module, service_module


@pytest.fixture
def fixture_path(tmp_path):
    return build_transport_fixture(tmp_path / "transport.xlsx")


@pytest.fixture
def database(tmp_path):
    result = Database(tmp_path / "app.db")
    result.migrate()
    masters = MasterRepository(result)
    destination = masters.create_destination("Known Destination", 1)
    masters.add_alias(destination.id, "Known Plant", "transport")
    return result


def _parsed_entry(parser_module, **overrides):
    values = {
        "report_month": "2026-08",
        "destination_alias": "Known Plant",
        "source_sheet": REGULAR_SHEET,
        "source_row": 3,
        "source_date": date(2026, 8, 1),
        "day": 1,
        "transport_type": "regular",
        "trip_count": Decimal("1"),
        "unit_rate_won": 1,
        "cost_won": 1,
    }
    values.update(overrides)
    return parser_module.ParsedTransportEntry(**values)


def test_parser_normalizes_day_columns_and_fractional_trips(api, fixture_path):
    parser_module, _, _ = api

    rows = parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")

    matching = [
        row
        for row in rows
        if row.source_sheet == REGULAR_SHEET and row.day == 12
    ]
    assert len(matching) == 1
    assert matching[0].trip_count == Decimal("1.25")
    assert matching[0].cost_won == 2_500
    assert matching[0].vehicle_driver_group == "5톤 / 기사A"


def test_parser_skips_repeated_headers_and_displayed_totals(api, fixture_path):
    parser_module, _, _ = api

    rows = parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")

    assert len(rows) == 5
    assert {row.source_sheet for row in rows} == {REGULAR_SHEET, *SUBCONTRACT_SHEETS}
    assert not {"운반지역", "합계"} & {row.destination_alias for row in rows}


def test_parser_accepts_anonymized_real_workbook_structure_and_formula_caches(
    api, tmp_path
):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(tmp_path / "structural-clone.xlsx")

    rows = parser_module.HwaseongWorkbookParser().parse(path, "2026-08")

    assert len(rows) == 5
    assert any(
        row.source_sheet == REGULAR_SHEET
        and row.source_row == 3
        and row.trip_count == Decimal("1.25")
        and row.cost_won == 2_500
        for row in rows
    )
    assert any(
        row.source_sheet == REGULAR_SHEET and row.source_row == 5 for row in rows
    )
    assert not {"합계", "총계", "계"} & {row.destination_alias for row in rows}


def test_parser_skips_anonymized_regular_footer_without_destination(api, tmp_path):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(
        tmp_path / "regular-footer.xlsx", include_regular_footer=True
    )

    rows = parser_module.HwaseongWorkbookParser().parse(path, "2026-08")

    assert not any(
        row.source_sheet == REGULAR_SHEET and row.source_row == 113 for row in rows
    )


def test_parser_rejects_ak_only_regular_footer_lookalike(api, tmp_path):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(
        tmp_path / "regular-footer-lookalike.xlsx",
        include_regular_footer_lookalike=True,
    )

    with pytest.raises(parser_module.WorkbookStructureError, match="destination"):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-08")


def test_parser_skips_both_inactive_regular_row_shapes_and_updates_group(
    api, tmp_path
):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(
        tmp_path / "inactive-rows.xlsx", include_inactive_regular_rows=True
    )

    rows = parser_module.HwaseongWorkbookParser().parse(path, "2026-08")

    aliases = {row.destination_alias for row in rows}
    assert "Dormant Blank Destination" not in aliases
    assert "Dormant Zero Destination" not in aliases
    trailing_row = next(
        row
        for row in rows
        if row.source_sheet == REGULAR_SHEET and row.source_row == 8
    )
    assert trailing_row.vehicle_driver_group == "8-ton / Driver B"


def test_parser_rejects_missing_aj_cache_on_inactive_looking_row(api, tmp_path):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(
        tmp_path / "inactive-missing-aj-cache.xlsx",
        include_inactive_regular_rows=True,
        missing_inactive_regular_unit_formula_cache=True,
    )

    with pytest.raises(
        parser_module.WorkbookStructureError,
        match=r"row 6 cell AJ6 cached value is missing.*recalculate.*Excel",
    ):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-08")


def test_parser_rejects_missing_ah_cache_on_inactive_looking_row(api, tmp_path):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(
        tmp_path / "inactive-missing-ah-cache.xlsx",
        include_inactive_regular_rows=True,
        missing_inactive_regular_count_formula_cache=True,
    )

    with pytest.raises(
        parser_module.WorkbookStructureError,
        match=r"row 6 cell AH6 cached value is missing.*recalculate.*Excel",
    ):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-08")


def test_parser_rejects_missing_ak_cache_on_inactive_looking_row(api, tmp_path):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(
        tmp_path / "inactive-missing-ak-cache.xlsx",
        include_inactive_regular_rows=True,
        missing_inactive_regular_subtotal_formula_cache=True,
    )

    with pytest.raises(
        parser_module.WorkbookStructureError,
        match=r"row 6 cell AK6 cached value is missing.*recalculate.*Excel",
    ):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-08")


def test_parser_rejects_numeric_string_regular_unit_rate(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[REGULAR_SHEET].cell(3, 36, "2000")
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="unit cost"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


@pytest.mark.parametrize(
    ("column", "value"),
    [
        (3, 1),
        (3, "malformed"),
        (34, 1),
        (34, "malformed"),
        (37, 1),
        (37, "malformed"),
        (36, "malformed"),
    ],
)
def test_parser_rejects_inactive_looking_row_with_invalid_activity_or_unit(
    api, fixture_path, column, value
):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    sheet = workbook[REGULAR_SHEET]
    for day_column in range(3, 34):
        sheet.cell(5, day_column).value = None
    sheet.cell(5, 34, 0)
    sheet.cell(5, 36).value = None
    sheet.cell(5, 37, 0)
    sheet.cell(5, column, value)
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_hidden_regular_detail_rows_are_parsed(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[REGULAR_SHEET].row_dimensions[5].hidden = True
    workbook.save(fixture_path)
    workbook.close()

    rows = parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")

    assert any(row.source_row == 5 for row in rows)


def test_hidden_subcontract_detail_rows_are_parsed(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[SUBCONTRACT_SHEETS[0]].row_dimensions[2].hidden = True
    workbook.save(fixture_path)
    workbook.close()

    rows = parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")

    assert any(
        row.source_sheet == SUBCONTRACT_SHEETS[0] and row.source_row == 2
        for row in rows
    )


def test_parser_rejects_missing_required_sheet(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    del workbook[SUBCONTRACT_SHEETS[-1]]
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="required sheet"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_rejects_missing_required_header(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[REGULAR_SHEET].cell(2, 36, "not-unit-cost")
    workbook[REGULAR_SHEET].cell(4, 36, "not-unit-cost")
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="header"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_requires_ah_total_count_header(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[REGULAR_SHEET].cell(2, 34, "not-total-count")
    workbook[REGULAR_SHEET].cell(4, 34, "not-total-count")
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="header"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_rejects_regular_total_count_mismatch(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[REGULAR_SHEET].cell(3, 34, 2)
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.TotalCountMismatchError, match="row 3"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


@pytest.mark.parametrize(
    ("coordinate", "message"),
    [("AH3", "AH3"), ("AK3", "AK3")],
)
def test_parser_rejects_missing_regular_formula_cache(
    api, tmp_path, coordinate, message
):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(tmp_path / "missing-cache.xlsx")
    patch_formula_cached_values(path, 1, {coordinate: None})

    with pytest.raises(
        parser_module.WorkbookStructureError,
        match=rf"row 3 cell {message} cached value is missing.*recalculate.*Excel",
    ):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-08")


@pytest.mark.parametrize(
    ("fixture_option", "coordinate"),
    [
        ("missing_regular_day_formula_cache", "N3"),
        ("missing_regular_unit_formula_cache", "AJ3"),
    ],
)
def test_parser_reports_missing_daily_and_unit_formula_caches(
    api, tmp_path, fixture_option, coordinate
):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(
        tmp_path / f"missing-{coordinate}-cache.xlsx",
        **{fixture_option: True},
    )

    with pytest.raises(
        parser_module.WorkbookStructureError,
        match=rf"row 3 cell {coordinate} cached value is missing.*recalculate.*Excel",
    ):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-08")


def test_parser_rejects_missing_subcontract_formula_cache(api, tmp_path):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(tmp_path / "missing-sub-cache.xlsx")
    patch_formula_cached_values(path, 2, {"I3": None})

    with pytest.raises(
        parser_module.WorkbookStructureError,
        match=r"row 3 cell I3 cached value is missing.*recalculate.*Excel",
    ):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-08")


def test_parser_rejects_missing_subcontract_total_formula_cache(api, tmp_path):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(tmp_path / "missing-total-cache.xlsx")
    patch_formula_cached_values(path, 2, {"I4": None})

    with pytest.raises(
        parser_module.WorkbookStructureError,
        match=r"row 4 cell I4 cached value is missing.*recalculate.*Excel",
    ):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-08")


def test_parser_rejects_stale_subcontract_total_formula_cache(api, tmp_path):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(tmp_path / "stale-total-cache.xlsx")
    patch_formula_cached_values(path, 2, {"I4": "50001"})

    with pytest.raises(parser_module.SubcontractTotalMismatchError, match="row 4"):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-08")


@pytest.mark.parametrize("sheet_name", SUBCONTRACT_SHEETS)
def test_parser_requires_one_displayed_subcontract_total(api, fixture_path, sheet_name):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    sheet = workbook[sheet_name]
    for cell in sheet[3]:
        cell.value = None
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(
        parser_module.WorkbookStructureError, match="displayed cached total"
    ):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


@pytest.mark.parametrize("sheet_name", SUBCONTRACT_SHEETS)
def test_parser_rejects_non_numeric_subcontract_total(api, fixture_path, sheet_name):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[sheet_name].cell(3, 9, "not-a-total")
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="cached total"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


@pytest.mark.parametrize("sheet_name", SUBCONTRACT_SHEETS)
def test_parser_rejects_mismatched_subcontract_total(api, fixture_path, sheet_name):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    sheet = workbook[sheet_name]
    sheet.cell(3, 9, sheet.cell(3, 9).value + 1)
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.SubcontractTotalMismatchError, match="row 3"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_reconciles_anonymized_total_rows_in_columns_a_and_g(api, tmp_path):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(tmp_path / "total-placements.xlsx")

    rows = parser_module.HwaseongWorkbookParser().parse(path, "2026-08")

    assert len([row for row in rows if row.transport_type == "nonregular"]) == 3


def test_parser_rejects_total_label_text_in_ancillary_detail_column(
    api, fixture_path
):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    sheet = workbook[SUBCONTRACT_SHEETS[0]]
    sheet.insert_rows(3)
    sheet.cell(3, 3, "계")
    sheet.cell(3, 9, 50_000)
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="date"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_rejects_subcontract_detail_after_displayed_total(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    sheet = workbook[SUBCONTRACT_SHEETS[0]]
    sheet.cell(4, 1, 2)
    sheet.cell(4, 2, date(2026, 8, 6))
    sheet.cell(4, 6, "Known Plant")
    sheet.cell(4, 8, "5톤")
    sheet.cell(4, 9, 60_000)
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(
        parser_module.WorkbookStructureError, match=r"row 4.*detail after.*total"
    ):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_rejects_second_subcontract_displayed_total(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    sheet = workbook[SUBCONTRACT_SHEETS[0]]
    sheet.cell(4, 1, "총계")
    sheet.cell(4, 9, 50_000)
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(
        parser_module.WorkbookStructureError, match=r"row 4.*second displayed total"
    ):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_accepts_actual_post_total_blank_template_and_formula_footer_shapes(
    api, tmp_path
):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(
        tmp_path / "post-total-footers.xlsx",
        include_subcontract_post_total_rows=True,
    )

    rows = parser_module.HwaseongWorkbookParser().parse(path, "2026-08")

    assert len(rows) == 5


@pytest.mark.parametrize("sheet_name", [REGULAR_SHEET, *SUBCONTRACT_SHEETS[:1]])
def test_parser_rejects_worksheet_row_dimension_abuse(
    api, fixture_path, sheet_name
):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[sheet_name].cell(2_001, 1).number_format = "0"
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(
        parser_module.WorkbookStructureError, match=r"row limit.*2000"
    ):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


@pytest.mark.parametrize("sheet_name", [REGULAR_SHEET, *SUBCONTRACT_SHEETS[:1]])
def test_parser_rejects_worksheet_column_dimension_abuse(
    api, fixture_path, sheet_name
):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[sheet_name].cell(1, 129).number_format = "0"
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(
        parser_module.WorkbookStructureError, match=r"column limit.*128"
    ):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_caps_total_entries_across_workbook(
    api, fixture_path, monkeypatch
):
    parser_module, _, _ = api
    monkeypatch.setattr(parser_module, "MAX_PARSED_ENTRIES", 4, raising=False)

    with pytest.raises(
        parser_module.WorkbookStructureError, match=r"entry limit.*4"
    ):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_rejects_regular_sheet_month_mismatch(api, tmp_path):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(tmp_path / "august.xlsx")

    with pytest.raises(parser_module.WorkbookStructureError, match="sheet month"):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-09")


def test_parser_accepts_previous_calendar_month_as_august_billing(api, tmp_path):
    parser_module, _, _ = api
    path = build_transport_fixture(tmp_path / "previous-month.xlsx")
    workbook = load_workbook(path)
    workbook[SUBCONTRACT_SHEETS[0]].cell(2, 1, date(2026, 7, 31))
    workbook.save(path)
    workbook.close()

    rows = parser_module.HwaseongWorkbookParser().parse(path, "2026-08")

    billed = next(
        row
        for row in rows
        if row.source_sheet == SUBCONTRACT_SHEETS[0] and row.source_row == 2
    )
    assert billed.report_month == "2026-08"
    assert billed.day == 31
    assert billed.source_date == date(2026, 7, 31)


def test_parser_accepts_previous_calendar_month_across_year_boundary(api, tmp_path):
    parser_module, _, _ = api
    path = build_transport_fixture(tmp_path / "january.xlsx", month=1)
    workbook = load_workbook(path)
    workbook[SUBCONTRACT_SHEETS[0]].cell(2, 1, date(2025, 12, 31))
    workbook.save(path)
    workbook.close()

    rows = parser_module.HwaseongWorkbookParser().parse(path, "2026-01")

    billed = next(
        row
        for row in rows
        if row.source_sheet == SUBCONTRACT_SHEETS[0] and row.source_row == 2
    )
    assert billed.report_month == "2026-01"
    assert billed.source_date == date(2025, 12, 31)


@pytest.mark.parametrize("outside_date", [date(2026, 6, 30), date(2026, 9, 1)])
def test_parser_rejects_subcontract_dates_outside_billing_window(
    api, tmp_path, outside_date
):
    parser_module, _, _ = api
    path = build_transport_fixture(tmp_path / "outside-window.xlsx")
    workbook = load_workbook(path)
    workbook[SUBCONTRACT_SHEETS[0]].cell(2, 1, outside_date)
    workbook.save(path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="billing window"):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-08")


def test_parser_preserves_full_regular_source_date(api, fixture_path):
    parser_module, _, _ = api

    rows = parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")

    regular = next(
        row
        for row in rows
        if row.source_sheet == REGULAR_SHEET and row.source_row == 3
    )
    assert regular.source_date == date(2026, 8, 12)


def test_parser_rejects_invalid_generated_calendar_date(api, tmp_path):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(tmp_path / "february.xlsx", month=2)
    workbook = load_workbook(path)
    sheet = workbook["화성운반비내역(2월)"]
    sheet.cell(3, 33, 1)  # day 31
    workbook.save(path)
    workbook.close()
    patch_formula_cached_values(path, 1, {"AH3": "2.25", "AK3": "4500"})

    with pytest.raises(parser_module.WorkbookStructureError, match="calendar date"):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-02")


def test_parser_rejects_data_row_missing_regular_destination(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[REGULAR_SHEET].cell(5, 2).value = None
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="destination"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_rejects_data_row_with_invalid_regular_unit_cost(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[REGULAR_SHEET].cell(5, 36, "invalid")
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="unit cost"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_rejects_nontext_regular_destination(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[REGULAR_SHEET].cell(5, 2, 123)
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="destination"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_rejects_data_row_missing_subcontract_destination(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[SUBCONTRACT_SHEETS[0]].cell(2, 2).value = None
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="destination"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_rejects_subcontract_detail_missing_vehicle(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[SUBCONTRACT_SHEETS[0]].cell(2, 8).value = None
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="vehicle"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


@pytest.mark.parametrize("candidate_column", [3, 4, 10])
def test_parser_rejects_partial_subcontract_candidate_across_full_footprint(
    api, fixture_path, candidate_column
):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    sheet = workbook[SUBCONTRACT_SHEETS[0]]
    sheet.insert_rows(3)
    sheet.cell(3, candidate_column, "partial detail")
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="date"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_checks_driver_column_when_its_header_is_blank(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    sheet = workbook[SUBCONTRACT_SHEETS[0]]
    sheet.cell(1, 10).value = None
    sheet.insert_rows(3)
    sheet.cell(3, 10, "partial driver")
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="date"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_skips_pre_numbered_blank_subcontract_template_rows(
    api, tmp_path
):
    parser_module, _, _ = api
    path = build_structural_clone_fixture(
        tmp_path / "pre-numbered-templates.xlsx",
        include_subcontract_template_rows=True,
    )

    rows = parser_module.HwaseongWorkbookParser().parse(path, "2026-08")

    assert len(rows) == 5


@pytest.mark.parametrize(
    ("column", "value", "message"),
    [(1, None, "date"), (9, None, "amount"), (9, "invalid", "amount")],
)
def test_parser_rejects_missing_or_invalid_subcontract_fields(
    api, fixture_path, column, value, message
):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[SUBCONTRACT_SHEETS[0]].cell(2, column).value = value
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match=message):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_rejects_nontext_subcontract_destination(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    workbook[SUBCONTRACT_SHEETS[0]].cell(2, 2, 123)
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="destination"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_rejects_subtotal_mismatch(api, tmp_path):
    parser_module, _, _ = api
    path = build_transport_fixture(
        tmp_path / "subtotal-mismatch.xlsx", subtotal_mismatch=True
    )

    with pytest.raises(parser_module.SubtotalMismatchError, match="row 3"):
        parser_module.HwaseongWorkbookParser().parse(path, "2026-08")


@pytest.mark.parametrize(
    "invalid_count",
    ["10000.0001", "0.00001", "1e100000", "12345678901234567"],
)
def test_parser_rejects_abusive_trip_counts_before_cost_calculation(
    api, fixture_path, invalid_count
):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    sheet = workbook[REGULAR_SHEET]
    sheet.cell(3, 14, invalid_count)
    sheet.cell(3, 34, invalid_count)
    workbook.save(fixture_path)
    workbook.close()

    with pytest.raises(parser_module.WorkbookStructureError, match="trip count"):
        parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")


def test_parser_accepts_trip_count_upper_boundary(api, fixture_path):
    parser_module, _, _ = api
    workbook = load_workbook(fixture_path)
    sheet = workbook[REGULAR_SHEET]
    sheet.cell(3, 14, 10_000)
    sheet.cell(3, 34, 10_000)
    sheet.cell(3, 36, 1)
    sheet.cell(3, 37, 10_000)
    workbook.save(fixture_path)
    workbook.close()

    rows = parser_module.HwaseongWorkbookParser().parse(fixture_path, "2026-08")

    boundary = next(row for row in rows if row.source_row == 3)
    assert boundary.trip_count == Decimal("10000")


def test_repository_rejects_exponent_abuse_before_persistence(api, database):
    parser_module, repository_module, _ = api
    row = parser_module.ParsedTransportEntry(
        report_month="2026-08",
        destination_alias="Known Plant",
        source_sheet=REGULAR_SHEET,
        source_row=3,
        source_date=date(2026, 8, 1),
        day=1,
        transport_type="regular",
        trip_count=Decimal("1e100000"),
        unit_rate_won=1,
        cost_won=1,
    )

    with pytest.raises(parser_module.WorkbookStructureError, match="trip count"):
        repository_module.TransportEntryRepository(database).import_entries(
            report_month="2026-08",
            source_filename="synthetic.xlsx",
            file_sha256="e" * 64,
            rows=[row],
        )

    with database.connection() as connection:
        assert (
            connection.execute("SELECT COUNT(*) FROM import_batches").fetchone()[0]
            == 0
        )
        assert (
            connection.execute("SELECT COUNT(*) FROM transport_entries").fetchone()[0]
            == 0
        )


@pytest.mark.parametrize(
    ("overrides", "message"),
    [
        ({"report_month": "2026-07"}, "report_month must match"),
        ({"day": 2}, "day must match source_date"),
        ({"source_date": None}, "source_date is required"),
        ({"transport_type": "charter"}, "transport_type"),
        (
            {"source_date": date(2026, 7, 31), "day": 31},
            "regular source_date",
        ),
        (
            {
                "source_date": date(2026, 6, 30),
                "day": 30,
                "transport_type": "nonregular",
                "unit_rate_won": None,
            },
            "nonregular source_date",
        ),
        (
            {
                "source_date": date(2026, 9, 1),
                "transport_type": "nonregular",
                "unit_rate_won": None,
            },
            "nonregular source_date",
        ),
    ],
    ids=(
        "wrong-row-month",
        "wrong-day",
        "missing-source-date",
        "invalid-type",
        "regular-prior-month",
        "nonregular-too-old",
        "nonregular-future",
    ),
)
def test_repository_rejects_inconsistent_entries_before_commit(
    api, database, overrides, message
):
    parser_module, repository_module, _ = api
    row = _parsed_entry(parser_module, **overrides)

    with pytest.raises(parser_module.WorkbookStructureError, match=message):
        repository_module.TransportEntryRepository(database).import_entries(
            report_month="2026-08",
            source_filename="synthetic.xlsx",
            file_sha256="d" * 64,
            rows=[row],
        )

    with database.connection() as connection:
        assert connection.execute("SELECT COUNT(*) FROM import_batches").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM transport_entries").fetchone()[0] == 0


def test_repository_accepts_nonregular_previous_month_across_year_boundary(
    api, database
):
    parser_module, repository_module, _ = api
    row = _parsed_entry(
        parser_module,
        report_month="2026-01",
        source_date=date(2025, 12, 31),
        day=31,
        transport_type="nonregular",
        unit_rate_won=None,
    )

    result = repository_module.TransportEntryRepository(database).import_entries(
        report_month="2026-01",
        source_filename="synthetic.xlsx",
        file_sha256="c" * 64,
        rows=[row],
    )

    assert result.status == "imported"
    assert result.inserted_count == 1


def test_duplicate_import_restores_newline_alias_blockers_from_entries(api, database):
    parser_module, repository_module, _ = api
    rows = [
        parser_module.ParsedTransportEntry(
            report_month="2026-08",
            destination_alias=alias,
            source_sheet=REGULAR_SHEET,
            source_row=source_row,
            source_date=date(2026, 8, 1),
            day=1,
            transport_type="regular",
            trip_count=Decimal("1"),
            unit_rate_won=1,
            cost_won=1,
        )
        for source_row, alias in enumerate(
            ("Unknown\nPlant", "Unknown\nPlant", "Other Plant"), start=3
        )
    ]
    repository = repository_module.TransportEntryRepository(database)

    first = repository.import_entries(
        report_month="2026-08",
        source_filename="synthetic.xlsx",
        file_sha256="f" * 64,
        rows=rows,
    )
    duplicate = repository.import_entries(
        report_month="2026-08",
        source_filename="synthetic.xlsx",
        file_sha256="f" * 64,
        rows=rows,
    )

    expected_errors = (
        "Unknown destination alias: Unknown\nPlant",
        "Unknown destination alias: Other Plant",
    )
    assert first.status == "imported_with_errors"
    assert first.blocking_errors == expected_errors
    assert duplicate.status == "duplicate"
    assert duplicate.inserted_count == 0
    assert duplicate.blocking_errors == expected_errors
    with database.connection() as connection:
        assert (
            connection.execute("SELECT COUNT(*) FROM import_batches").fetchone()[0]
            == 1
        )
        assert (
            connection.execute("SELECT COUNT(*) FROM transport_entries").fetchone()[0]
            == 3
        )


def test_import_is_idempotent_and_stores_file_hash(api, database, fixture_path):
    _, _, service_module = api
    workbook = load_workbook(fixture_path)
    workbook[REGULAR_SHEET].cell(5, 2, "Known Plant")
    workbook.save(fixture_path)
    workbook.close()
    service = service_module.TransportImportService(database)

    first = service.import_transport(fixture_path, "2026-08")
    second = service.import_transport(fixture_path, "2026-08")

    assert first.status == "imported"
    assert first.inserted_count == 5
    assert second.status == "duplicate"
    assert second.inserted_count == 0
    assert second.blocking_errors == ()
    with database.connection() as connection:
        batches = connection.execute("SELECT * FROM import_batches").fetchall()
        entries = connection.execute("SELECT * FROM transport_entries").fetchall()
    assert len(batches) == 1
    assert len(entries) == 5
    assert batches[0]["file_sha256"] == hashlib.sha256(
        fixture_path.read_bytes()
    ).hexdigest()


def test_import_persists_complete_source_dates(api, database, tmp_path):
    _, _, service_module = api
    path = build_transport_fixture(tmp_path / "source-dates.xlsx")
    workbook = load_workbook(path)
    workbook[SUBCONTRACT_SHEETS[0]].cell(2, 1, date(2026, 7, 31))
    workbook.save(path)
    workbook.close()

    service_module.TransportImportService(database).import_transport(path, "2026-08")

    with database.connection() as connection:
        source_dates = {
            row["source_date"]
            for row in connection.execute(
                "SELECT source_date FROM transport_entries"
            )
        }
    assert "2026-07-31" in source_dates
    assert "2026-08-12" in source_dates


def test_import_snapshots_once_then_hashes_and_parses_the_same_bytes(
    api, database, fixture_path
):
    parser_module, repository_module, service_module = api
    original_bytes = fixture_path.read_bytes()
    original_hash = hashlib.sha256(original_bytes).hexdigest()

    class MutatingParser:
        snapshot_path: Path | None = None
        snapshot_bytes: bytes | None = None

        def parse(self, path, report_month):
            self.snapshot_path = Path(path)
            self.snapshot_bytes = self.snapshot_path.read_bytes()
            fixture_path.write_bytes(b"source changed after snapshot")
            return []

    class RecordingRepository:
        source_filename: str | None = None
        file_sha256: str | None = None

        def import_entries(self, **kwargs):
            self.source_filename = kwargs["source_filename"]
            self.file_sha256 = kwargs["file_sha256"]
            return repository_module.ImportCommitResult(1, "imported", 0, ())

    parser = MutatingParser()
    repository = RecordingRepository()
    result = service_module.TransportImportService(
        database, parser=parser, repository=repository
    ).import_transport(fixture_path, "2026-08")

    assert parser.snapshot_bytes == original_bytes
    assert parser.snapshot_path is not None
    assert parser.snapshot_path != fixture_path
    assert not parser.snapshot_path.exists()
    assert result.file_sha256 == original_hash
    assert repository.file_sha256 == original_hash
    assert repository.source_filename == "transport.xlsx"


@pytest.mark.parametrize(
    ("limit_name", "limit", "message"),
    [
        ("MAX_INPUT_BYTES", 1, "input byte limit"),
        ("MAX_ZIP_MEMBERS", 1, "member count"),
        ("MAX_ZIP_MEMBER_BYTES", 1, "member size"),
        ("MAX_ZIP_UNCOMPRESSED_BYTES", 1, "uncompressed size"),
    ],
)
def test_import_rejects_oversized_or_zip_bomb_like_inputs(
    api, database, fixture_path, monkeypatch, limit_name, limit, message
):
    _, _, service_module = api
    monkeypatch.setattr(service_module, limit_name, limit)

    with pytest.raises(service_module.ImportFileLimitError, match=message):
        service_module.TransportImportService(database).import_transport(
            fixture_path, "2026-08"
        )

    with database.connection() as connection:
        assert (
            connection.execute("SELECT COUNT(*) FROM import_batches").fetchone()[0]
            == 0
        )


def test_snapshot_temp_file_is_removed_when_parser_fails(api, database, fixture_path):
    parser_module, _, service_module = api

    class FailingParser:
        snapshot_path: Path | None = None

        def parse(self, path, report_month):
            self.snapshot_path = Path(path)
            raise parser_module.WorkbookStructureError("forced parse failure")

    parser = FailingParser()
    service = service_module.TransportImportService(database, parser=parser)

    with pytest.raises(parser_module.WorkbookStructureError, match="forced"):
        service.import_transport(fixture_path, "2026-08")

    assert parser.snapshot_path is not None
    assert not parser.snapshot_path.exists()


def test_concurrent_identical_imports_commit_one_batch(api, database, fixture_path):
    _, _, service_module = api
    service = service_module.TransportImportService(database)

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = [
            future.result()
            for future in (
                executor.submit(service.import_transport, fixture_path, "2026-08"),
                executor.submit(service.import_transport, fixture_path, "2026-08"),
            )
        ]

    assert sorted(result.status for result in results) == [
        "duplicate",
        "imported_with_errors",
    ]
    assert sorted(result.inserted_count for result in results) == [0, 5]
    with database.connection() as connection:
        assert (
            connection.execute("SELECT COUNT(*) FROM import_batches").fetchone()[0]
            == 1
        )
        assert (
            connection.execute("SELECT COUNT(*) FROM transport_entries").fetchone()[0]
            == 5
        )


def test_unknown_alias_is_imported_as_blocking_unresolved_entry(
    api, database, fixture_path
):
    _, _, service_module = api
    service = service_module.TransportImportService(database)

    result = service.import_transport(fixture_path, "2026-08")
    destination = MasterRepository(database).list_destinations()[0]
    MasterRepository(database).add_alias(destination.id, "Unknown Plant", "transport")
    duplicate = service.import_transport(fixture_path, "2026-08")

    assert result.status == "imported_with_errors"
    assert result.blocking_errors == ("Unknown destination alias: Unknown Plant",)
    assert duplicate.status == "duplicate"
    assert duplicate.inserted_count == 0
    assert duplicate.blocking_errors == result.blocking_errors
    with database.connection() as connection:
        unresolved = connection.execute(
            "SELECT destination_id, unresolved_alias, source_alias "
            "FROM transport_entries "
            "WHERE destination_id IS NULL"
        ).fetchall()
        resolved_aliases = connection.execute(
            "SELECT DISTINCT source_alias FROM transport_entries "
            "WHERE destination_id IS NOT NULL"
        ).fetchall()
        batch = connection.execute("SELECT * FROM import_batches").fetchone()
    assert [tuple(row) for row in unresolved] == [
        (None, "Unknown Plant", "Unknown Plant")
    ]
    assert [row["source_alias"] for row in resolved_aliases] == ["Known Plant"]
    assert "Unknown Plant" in batch["error_summary"]


def test_structural_failure_is_atomic(api, database, tmp_path):
    parser_module, _, service_module = api
    path = build_transport_fixture(
        tmp_path / "subtotal-mismatch.xlsx", subtotal_mismatch=True
    )

    with pytest.raises(parser_module.SubtotalMismatchError):
        service_module.TransportImportService(database).import_transport(
            path, "2026-08"
        )

    with database.connection() as connection:
        assert connection.execute("SELECT COUNT(*) FROM import_batches").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM transport_entries").fetchone()[0] == 0


def test_locked_month_rejects_import_without_partial_batch(
    api, database, fixture_path
):
    _, _, service_module = api
    MonthlyInputRepository(database).finalize_month("2026-08", "revision-1")

    with pytest.raises(MonthLockedError, match="2026-08"):
        service_module.TransportImportService(database).import_transport(
            fixture_path, "2026-08"
        )

    with database.connection() as connection:
        assert connection.execute("SELECT COUNT(*) FROM import_batches").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM transport_entries").fetchone()[0] == 0
