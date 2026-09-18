from __future__ import annotations

from dataclasses import FrozenInstanceError
from decimal import Decimal
from pathlib import Path

import pytest

from app.db import Database
from app.repositories.masters import MasterRepository
from tests.fixtures.build_legacy_fixture import (
    DESTINATIONS,
    build_legacy_fixture,
    change_source_value,
)


@pytest.fixture
def migration_api():
    from app.importers.legacy_migration import (
        LEGACY_SOURCE_TYPE,
        LegacyMigrationBlockedError,
        LegacyMigrationConfirmationError,
        LegacyMigrationError,
        LegacyMigrationRevisionError,
        LegacyMigrationService,
    )

    return {
        "source_type": LEGACY_SOURCE_TYPE,
        "blocked": LegacyMigrationBlockedError,
        "confirmation": LegacyMigrationConfirmationError,
        "error": LegacyMigrationError,
        "revision": LegacyMigrationRevisionError,
        "service": LegacyMigrationService,
    }


@pytest.fixture
def database(tmp_path: Path) -> Database:
    result = Database(tmp_path / "transport.sqlite3")
    result.migrate()
    return result


@pytest.fixture
def masters(database: Database, migration_api):
    repository = MasterRepository(database)
    destinations = {}
    for order, name in enumerate(DESTINATIONS, start=1):
        destination = repository.create_destination(
            name,
            order,
            include_quantity_total=name != "당진",
            include_cost_total=True,
        )
        repository.add_alias(destination.id, name, migration_api["source_type"])
        destinations[name] = destination
    group = repository.create_group("South", 1)
    repository.replace_group_members(
        group.id,
        [destinations["South A"].id, destinations["South B"].id],
    )
    return repository, destinations, group


@pytest.fixture
def service(database: Database, masters, migration_api):
    return migration_api["service"](database)


def _commit(service, dry_run, source: Path):
    return service.commit(
        source,
        report_month="2026-08",
        confirmed=True,
        expected_source_sha256=dry_run.source_sha256,
        expected_master_revision=dry_run.master_revision,
        expected_database_revision=dry_run.database_revision,
        expected_revision=dry_run.revision,
    )


def _table_counts(database: Database) -> tuple[int, int, int]:
    with database.connection() as connection:
        return tuple(
            connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
            for table in (
                "monthly_plans",
                "monthly_actual_quantities",
                "monthly_actual_costs",
            )
        )


def test_parses_exact_regions_cached_formulas_and_ignores_chart_helpers(
    tmp_path, service
):
    source = build_legacy_fixture(tmp_path / "legacy.xlsx")

    result = service.dry_run(source, report_month="2026-08")

    assert result.can_commit
    assert result.summary.plan_records == 14
    assert result.summary.actual_records == 280
    assert result.summary.current_plan_records == 14
    assert result.summary.current_actual_records == 14
    assert result.summary.months == tuple(
        [f"2025-{month:02d}" for month in range(1, 13)]
        + [f"2026-{month:02d}" for month in range(1, 9)]
    )
    assert all(record.destination_alias != "Mystery helper" for record in result.plans)
    assert all(record.destination_alias != "Mystery helper" for record in result.actuals)
    alpha_august = next(
        record
        for record in result.plans
        if record.destination_alias == "Alpha" and record.report_month == "2026-08"
    )
    assert alpha_august.quantity_ea == Decimal("86")
    with pytest.raises(FrozenInstanceError):
        alpha_august.quantity_ea = Decimal("1")


def test_uncached_formula_is_a_blocker_with_exact_locator(tmp_path, service):
    source = build_legacy_fixture(
        tmp_path / "uncached.xlsx", formula_without_cache=True
    )

    result = service.dry_run(source, report_month="2026-08")

    issue = next(item for item in result.issues if item.code == "MISSING_FORMULA_CACHE")
    assert not result.can_commit
    assert issue.source_locator == "26년 월계획!Q6"
    assert "수식의 저장값" in issue.message

    history_source = build_legacy_fixture(
        tmp_path / "uncached-history.xlsx", history_formula_without_cache=True
    )
    history_result = service.dry_run(history_source, report_month="2026-08")
    history_issue = next(
        item
        for item in history_result.issues
        if item.code == "MISSING_FORMULA_CACHE"
    )
    assert history_issue.source_locator == "누적 데이터!E2"


def test_unknown_alias_is_reported_without_guessing(tmp_path, service):
    source = build_legacy_fixture(tmp_path / "unknown.xlsx", unknown_alias=True)

    result = service.dry_run(source, report_month="2026-08")

    assert not result.can_commit
    assert result.unknown_aliases[0].alias == "Mystery"
    assert result.unknown_aliases[0].source_type == "legacy_workbook"
    assert result.unknown_aliases[0].source_locator == "누적 데이터!D2"
    assert any(item.source_locator == "26년 월계획!B6" for item in result.unknown_aliases)


def test_canonical_destination_name_still_requires_legacy_source_alias(
    tmp_path, service, masters
):
    alpha_alias = next(
        item
        for item in masters[0].list_aliases(masters[1]["Alpha"].id)
        if item.source_type == "legacy_workbook"
    )
    masters[0].remove_alias(alpha_alias.id)
    source = build_legacy_fixture(tmp_path / "canonical-without-alias.xlsx")

    result = service.dry_run(source, report_month="2026-08")

    assert not result.can_commit
    assert any(
        item.alias == "Alpha" and item.source_locator == "누적 데이터!D2"
        for item in result.unknown_aliases
    )


def test_dry_run_opens_one_read_transaction(tmp_path, service, monkeypatch):
    import app.importers.legacy_migration as migration

    source = build_legacy_fixture(tmp_path / "snapshot.xlsx")
    original = migration._load_destinations
    observed = []

    def checking_loader(connection):
        observed.append(connection.in_transaction)
        return original(connection)

    monkeypatch.setattr(migration, "_load_destinations", checking_loader)

    service.dry_run(source, report_month="2026-08")

    assert observed == [True]


def test_dry_run_is_strictly_read_only_for_sqlite(tmp_path, database, service):
    source = build_legacy_fixture(tmp_path / "legacy.xlsx")
    source_bytes = source.read_bytes()
    before_bytes = database.path.read_bytes()
    before_counts = _table_counts(database)

    service.dry_run(source, report_month="2026-08")

    assert _table_counts(database) == before_counts == (0, 0, 0)
    assert database.path.read_bytes() == before_bytes
    assert source.read_bytes() == source_bytes


def test_commit_requires_confirmation_and_exact_dry_run_tokens(
    tmp_path, service, migration_api
):
    source = build_legacy_fixture(tmp_path / "legacy.xlsx")
    dry_run = service.dry_run(source, report_month="2026-08")

    with pytest.raises(migration_api["confirmation"], match="확인"):
        service.commit(
            source,
            report_month="2026-08",
            confirmed=False,
            expected_source_sha256=dry_run.source_sha256,
            expected_master_revision=dry_run.master_revision,
            expected_database_revision=dry_run.database_revision,
            expected_revision=dry_run.revision,
        )
    with pytest.raises(migration_api["revision"], match="해시"):
        service.commit(
            source,
            report_month="2026-08",
            confirmed=True,
            expected_source_sha256="0" * 64,
            expected_master_revision=dry_run.master_revision,
            expected_database_revision=dry_run.database_revision,
            expected_revision=dry_run.revision,
        )


def test_changed_source_or_master_after_dry_run_cannot_commit(
    tmp_path, service, masters, migration_api
):
    source = build_legacy_fixture(tmp_path / "legacy.xlsx")
    dry_run = service.dry_run(source, report_month="2026-08")
    change_source_value(source)

    with pytest.raises(migration_api["revision"], match="변경"):
        _commit(service, dry_run, source)

    source = build_legacy_fixture(tmp_path / "legacy.xlsx")
    dry_run = service.dry_run(source, report_month="2026-08")
    masters[0].update_destination(masters[1]["Alpha"].id, display_order=99)
    with pytest.raises(migration_api["revision"], match="기준정보"):
        _commit(service, dry_run, source)


def test_confirmed_commit_is_atomic_and_repeated_import_is_idempotent(
    tmp_path, database, service
):
    source = build_legacy_fixture(tmp_path / "legacy.xlsx")
    first_dry_run = service.dry_run(source, report_month="2026-08")

    first = _commit(service, first_dry_run, source)
    first_counts = _table_counts(database)
    second_dry_run = service.dry_run(source, report_month="2026-08")
    second = _commit(service, second_dry_run, source)

    assert first.inserted == sum(first_counts) == 574
    assert first.unchanged == 0
    assert second.inserted == 0
    assert second.unchanged == 574
    assert _table_counts(database) == first_counts == (14, 280, 280)


def test_conflicting_existing_value_blocks_without_overwrite(
    tmp_path, database, service, masters, migration_api
):
    source = build_legacy_fixture(tmp_path / "legacy.xlsx")
    alpha_id = masters[1]["Alpha"].id
    with database.connection() as connection:
        connection.execute(
            "INSERT INTO monthly_actual_quantities"
            "(report_month, destination_id, quantity_ea_text, source_note) "
            "VALUES ('2026-08', ?, '999', 'manual')",
            (alpha_id,),
        )
        connection.commit()

    dry_run = service.dry_run(source, report_month="2026-08")

    conflict = next(item for item in dry_run.issues if item.code == "EXISTING_VALUE_CONFLICT")
    assert conflict.expected == "999"
    assert conflict.actual == "2681"
    with pytest.raises(migration_api["blocked"], match="차이"):
        _commit(service, dry_run, source)
    with database.connection() as connection:
        stored = connection.execute(
            "SELECT quantity_ea_text FROM monthly_actual_quantities"
        ).fetchone()[0]
    assert stored == "999"


def test_failure_or_new_lock_rolls_back_every_write(
    tmp_path, database, service, migration_api
):
    source = build_legacy_fixture(tmp_path / "legacy.xlsx")
    dry_run = service.dry_run(source, report_month="2026-08")
    with database.connection() as connection:
        connection.execute(
            "CREATE TRIGGER fail_legacy_cost BEFORE INSERT ON monthly_actual_costs "
            "BEGIN SELECT RAISE(ABORT, 'forced migration failure'); END"
        )
        connection.commit()
    with pytest.raises(migration_api["error"], match="롤백"):
        _commit(service, dry_run, source)
    assert _table_counts(database) == (0, 0, 0)

    with database.connection() as connection:
        connection.execute("DROP TRIGGER fail_legacy_cost")
        connection.commit()
    dry_run = service.dry_run(source, report_month="2026-08")
    with database.connection() as connection:
        connection.execute(
            "INSERT INTO month_locks"
            "(report_month, is_locked, locked_at, input_revision) "
            "VALUES ('2025-01', 1, '2026-09-18T00:00:00Z', 'later-lock')"
        )
        connection.commit()
    with pytest.raises(migration_api["revision"], match="상태"):
        _commit(service, dry_run, source)
    assert _table_counts(database) == (0, 0, 0)


def test_zero_is_staged_but_missing_history_makes_windows_incomplete(
    tmp_path, service
):
    zero_source = build_legacy_fixture(tmp_path / "zero.xlsx", zero_current=True)
    zero_result = service.dry_run(zero_source, report_month="2026-08")
    zero = next(
        record
        for record in zero_result.actuals
        if record.destination_alias == "Alpha" and record.report_month == "2026-08"
    )
    assert zero.quantity_ea == 0
    assert zero.cost_won == 0

    missing_source = build_legacy_fixture(
        tmp_path / "missing.xlsx", omit_history=("Alpha", 2026, 7)
    )
    missing_result = service.dry_run(missing_source, report_month="2026-08")
    assert not missing_result.reconciliation.actual_quantity_history.three_month.complete
    assert missing_result.reconciliation.actual_quantity_history.three_month.value is None
    assert missing_result.reconciliation.actual_quantity_history.three_month.missing_months == (
        "2026-07",
    )
    assert not missing_result.can_commit
    assert any(item.code == "INCOMPLETE_HISTORY" for item in missing_result.issues)


def test_missing_actual_quantity_preserves_cost_without_inventing_zero(
    tmp_path, database, service, masters
):
    source = build_legacy_fixture(
        tmp_path / "cost-only.xlsx",
        missing_actual_quantity=("당진", 2026, 8),
    )

    dry_run = service.dry_run(source, report_month="2026-08")
    record = next(
        item
        for item in dry_run.actuals
        if item.destination_alias == "당진" and item.report_month == "2026-08"
    )
    assert record.quantity_ea is None
    assert record.cost_won == 268_200
    assert dry_run.reconciliation.current_total.actual_quantity == sum(
        (Decimal(2680 + offset) for offset in range(1, 15) if offset != 2),
        Decimal(0),
    )
    assert dry_run.reconciliation.current_total.actual_cost_won == (
        sum(range(2681, 2695)) * 100
    )
    assert dry_run.can_commit

    _commit(service, dry_run, source)
    dangjin_id = masters[1]["당진"].id
    with database.connection() as connection:
        quantity = connection.execute(
            "SELECT quantity_ea_text FROM monthly_actual_quantities "
            "WHERE report_month = '2026-08' AND destination_id = ?",
            (dangjin_id,),
        ).fetchone()
        cost = connection.execute(
            "SELECT cost_won FROM monthly_actual_costs "
            "WHERE report_month = '2026-08' AND destination_id = ?",
            (dangjin_id,),
        ).fetchone()
    assert quantity is None
    assert cost[0] == 268_200


def test_reconciliation_uses_total_rules_groups_and_complete_history(
    tmp_path, service, masters
):
    source = build_legacy_fixture(tmp_path / "legacy.xlsx")

    result = service.dry_run(source, report_month="2026-08")

    total = result.reconciliation.current_total
    assert total.planned_quantity == sum(
        (Decimal(80 + row) for row in range(6, 20) if row != 7),
        Decimal(0),
    )
    assert total.planned_cost_won == sum((80 + row) * 1000 for row in range(6, 20))
    assert total.actual_quantity == sum(
        (Decimal(2680 + offset) for offset in range(1, 15) if offset != 2),
        Decimal(0),
    )
    assert total.actual_cost_won == sum(range(2681, 2695)) * 100
    south = next(item for item in result.reconciliation.groups if item.name == "South")
    assert south.result.planned_quantity == Decimal("177")
    assert south.result.planned_cost_won == 177_000
    history = result.reconciliation.actual_cost_history
    assert history.three_month.complete
    assert history.six_month.complete
    assert history.twelve_month.complete
    assert history.comparison_year.complete
    assert all(control.matches for control in result.reconciliation.controls)


def test_reconciliation_control_mismatch_is_exact_and_blocking(tmp_path, service):
    source = build_legacy_fixture(tmp_path / "mismatch.xlsx", total_mismatch=True)

    result = service.dry_run(source, report_month="2026-08")

    issue = next(item for item in result.issues if item.code == "RECONCILIATION_MISMATCH")
    assert issue.source_locator == "26년 월계획!Q20"
    assert issue.expected == "1209"
    assert issue.actual == "1208"
    assert not result.can_commit


@pytest.mark.parametrize(
    ("field", "value", "code", "message"),
    [
        ("year", "20X6년", "INVALID_YEAR", "연도"),
        ("month", "13월", "INVALID_MONTH", "월"),
        ("quantity", "많음", "INVALID_QUANTITY", "수량"),
        ("cost", 1.5, "INVALID_COST", "원 단위"),
    ],
)
def test_malformed_source_values_have_deterministic_korean_messages(
    tmp_path, service, field, value, code, message
):
    source = build_legacy_fixture(
        tmp_path / f"bad-{field}.xlsx", malformed=(field, value)
    )

    result = service.dry_run(source, report_month="2026-08")

    issue = next(item for item in result.issues if item.code == code)
    expected_cell = {
        "year": "A2",
        "month": "B2",
        "quantity": "E2",
        "cost": "F2",
    }[field]
    assert issue.source_locator.endswith(expected_cell)
    assert message in issue.message
    assert not result.can_commit


def test_wrong_headers_are_blocking_and_actionable(tmp_path, service):
    source = build_legacy_fixture(tmp_path / "headers.xlsx")
    from openpyxl import load_workbook

    workbook = load_workbook(source)
    workbook["누적 데이터"]["E1"] = "수량"
    workbook.save(source)
    workbook.close()

    result = service.dry_run(source, report_month="2026-08")

    issue = next(item for item in result.issues if item.code == "HEADER_MISMATCH")
    assert issue.source_locator == "누적 데이터!E1"
    assert "실적수량" in issue.message
    assert not result.can_commit


def test_internal_history_gap_is_a_structure_blocker_but_helper_total_is_ignored(
    tmp_path, service
):
    source = build_legacy_fixture(tmp_path / "structure.xlsx")
    from openpyxl import load_workbook

    workbook = load_workbook(source)
    history = workbook["누적 데이터"]
    for column in range(1, 7):
        history.cell(10, column).value = None
    workbook["26년 월계획"]["B21"] = "합계"
    workbook.save(source)
    workbook.close()

    result = service.dry_run(source, report_month="2026-08")

    codes = {item.code for item in result.issues}
    assert "UNEXPECTED_BLANK_HISTORY_ROW" in codes
    assert "MULTIPLE_PLAN_TOTALS" not in codes
    assert not result.can_commit


def test_moved_sole_plan_total_is_a_structure_blocker(tmp_path, service):
    source = build_legacy_fixture(tmp_path / "moved-total.xlsx")
    from openpyxl import load_workbook

    workbook = load_workbook(source)
    plan = workbook["26년 월계획"]
    plan["B20"] = None
    plan["B26"] = "합계"
    workbook.save(source)
    workbook.close()

    result = service.dry_run(source, report_month="2026-08")

    issue = next(item for item in result.issues if item.code == "PLAN_TOTAL_ROW_MISMATCH")
    assert issue.source_locator == "26년 월계획!B20"
    assert not result.can_commit


def test_extra_populated_history_row_is_blocked_and_not_staged(tmp_path, service):
    source = build_legacy_fixture(tmp_path / "extra-history.xlsx")
    from openpyxl import load_workbook

    workbook = load_workbook(source)
    history = workbook["누적 데이터"]
    history["A282"] = "26년"
    history["B282"] = "8월"
    history["D282"] = "Alpha"
    history["E282"] = 1
    history["F282"] = 100
    workbook.save(source)
    workbook.close()

    result = service.dry_run(source, report_month="2026-08")

    issue = next(item for item in result.issues if item.code == "HISTORY_ROW_OUT_OF_RANGE")
    assert issue.source_locator == "누적 데이터!A282:F282"
    assert len(result.actuals) == 280
    assert not result.can_commit


def test_identical_repeat_is_no_write_even_when_month_is_locked(
    tmp_path, database, service
):
    source = build_legacy_fixture(tmp_path / "locked-repeat.xlsx")
    first = service.dry_run(source, report_month="2026-08")
    _commit(service, first, source)
    with database.connection() as connection:
        connection.execute(
            "INSERT INTO month_locks"
            "(report_month, is_locked, locked_at, input_revision) "
            "VALUES ('2026-08', 1, '2026-09-18T00:00:00Z', 'locked-repeat')"
        )
        connection.commit()
    after_lock = database.path.read_bytes()

    repeat = service.dry_run(source, report_month="2026-08")

    assert repeat.can_commit
    assert not any(item.code == "MONTH_LOCKED" for item in repeat.issues)
    result = _commit(service, repeat, source)
    assert result.inserted == 0
    assert result.unchanged == 574
    assert _table_counts(database) == (14, 280, 280)
    assert database.path.read_bytes() == after_lock


def test_dry_run_exposes_frozen_database_snapshot_and_revision(tmp_path, service):
    source = build_legacy_fixture(tmp_path / "snapshot-fields.xlsx")
    changed_source = build_legacy_fixture(
        tmp_path / "snapshot-other-source.xlsx", zero_current=True
    )

    result = service.dry_run(source, report_month="2026-08")
    changed = service.dry_run(changed_source, report_month="2026-08")

    assert len(result.database_revision) == 64
    assert result.database_snapshot.month_locks[0].report_month == "2025-01"
    assert result.database_snapshot.month_locks[0].is_locked is False
    actions = {target.action for target in result.database_snapshot.targets}
    assert actions == {"insert"}
    assert len(result.database_snapshot.targets) == 574
    assert changed.database_revision == result.database_revision
    assert changed.revision != result.revision
    with pytest.raises(FrozenInstanceError):
        result.database_snapshot.month_locks[0].is_locked = True
