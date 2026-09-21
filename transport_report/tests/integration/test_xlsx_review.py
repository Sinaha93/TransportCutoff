"""Fictional integration coverage for the review-workbook exporter."""
from __future__ import annotations

from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook


ERROR_TOKENS = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
    "#SPILL!",
    "#CALC!",
}


def test_ppt_report_preflight_is_public():
    from app.reporting.pptx_report import validate_ppt_report

    assert callable(validate_ppt_report)


def _report_bundle():
    from app.domain.calculations import (
        calculate_destination,
        calculate_group,
        calculate_total,
        historical_averages,
    )
    from app.domain.models import Destination, DestinationAlias, GroupMember, ReportGroup
    from app.domain.validation import (
        DestinationValidationInput,
        ImportBatchInput,
        NextMonthPlanInput,
        ReportMonthSources,
        SalesInput,
        ValidationContext,
    )
    from app.reporting.charts import ChartReport
    from app.reporting.pptx_report import PlanReport, PptReport, ReportRow, SalesReport
    from app.reporting.xlsx_review import (
        ImportBatchEvidence,
        OperationEvidence,
        ReviewWorkbookData,
    )

    destinations = tuple(
        Destination(
            destination_id,
            f"가상납품처{destination_id:02d}",
            destination_id,
            True,
            True,
            f"가상품 {destination_id:02d}" if destination_id == 1 else None,
            True,
            True,
            True,
        )
        for destination_id in range(1, 13)
    )
    aliases = tuple(
        DestinationAlias(
            destination_id,
            destination_id,
            f"원천납품처{destination_id:02d}",
            "transport",
        )
        for destination_id in range(1, 13)
    )
    groups = (
        ReportGroup(10, "가상권역", 1, True),
        ReportGroup(11, "빈 가상권역", 2, True),
    )
    members = (
        GroupMember(10, 1, "가상납품처01", 1, True, True),
        GroupMember(10, 2, "가상납품처02", 2, True, False),
    )
    calculations = {
        destination.id: calculate_destination(
            Decimal(100 + destination.id),
            (100 + destination.id) * 100,
            Decimal(0 if destination.id == 2 else 120 + destination.id),
            0 if destination.id == 2 else (120 + destination.id) * 110,
            destination_id=destination.id,
        )
        for destination in destinations
    }
    group_calculation = calculate_group(calculations, members, group_id=10)
    empty_group_calculation = calculate_group(calculations, (), group_id=11)
    total_calculation = calculate_total(calculations, destinations)
    months = tuple(f"2025-{month:02d}" for month in range(1, 13)) + tuple(
        f"2026-{month:02d}" for month in range(1, 9)
    )

    def row(key, label, calculation):
        quantities = {month: Decimal("90") for month in months}
        costs = {month: 9_000 for month in months}
        quantities["2026-08"] = calculation.actual_quantity
        costs["2026-08"] = calculation.actual_cost_won
        return ReportRow(key, label, calculation, quantities, costs)

    direct_rows = tuple(
        row(str(destination.id), destination.name, calculations[destination.id])
        for destination in destinations
    )
    group_row = row("group:10", "가상권역", group_calculation)
    empty_group_row = row("group:11", "빈 가상권역", empty_group_calculation)
    total_row = row("total", "합계", total_calculation)
    nonregular = row(
        "nonregular",
        "비정규 운반비",
        calculate_destination(Decimal("0"), 0, Decimal("0"), 0),
    )
    sales_history = {month: 1_000_000 for month in months}
    sales_history["2026-08"] = 1_200_000
    sales = SalesReport(1_100_000, 1_200_000, sales_history)
    next_rows = tuple(
        replace(
            direct_rows[destination.id - 1],
            calculation=calculate_destination(
                Decimal(130 + destination.id),
                (130 + destination.id) * 110,
                None,
                None,
                destination_id=destination.id,
            ),
        )
        for destination in destinations
    )
    next_total = calculate_total(
        {destination.id: next_rows[destination.id - 1].calculation for destination in destinations},
        destinations,
    )
    chart = ChartReport(
        "2026-08",
        {
            month: (
                total_calculation.planned_quantity
                if month == "2026-08"
                else Decimal("150")
            )
            for month in months
        },
        {month: total_row.quantity_by_month[month] for month in months},
        {
            month: (
                total_calculation.planned_cost_won if month == "2026-08" else 15_000
            )
            for month in months
        },
        {month: total_row.cost_won_by_month[month] for month in months},
    )
    report = PptReport(
        "2026-08",
        date(2026, 9, 2),
        (*direct_rows, group_row, empty_group_row),
        nonregular,
        total_row,
        sales,
        PlanReport(
            "2026-09",
            next_rows,
            replace(nonregular, calculation=calculate_destination(Decimal("0"), 0, None, None)),
            replace(total_row, calculation=next_total),
            SalesReport(1_300_000, None, sales_history),
        ),
        chart,
    )
    batch = ImportBatchEvidence(
        batch_id=42,
        provenance_id="transport:2026-08:42",
        report_month="2026-08",
        source_type="transport",
        source_filename=r"D:\\fictional\\화성 가상자료.xlsx",
        file_sha256="ab" * 32,
        imported_at=datetime(2026, 9, 1, 8, 30, 15),
    )
    operations = (
        *(
            OperationEvidence(
                record_kind="destination",
                value_role="actual",
                input_kind="manual" if destination.id == 2 else "imported",
                batch_id=None if destination.id == 2 else 42,
                provenance_id=(
                    "monthly-actual:2026-08:2"
                    if destination.id == 2
                    else "transport:2026-08:42"
                ),
                report_month="2026-08",
                raw_destination=(
                    None if destination.id == 2 else f"원천납품처{destination.id:02d}"
                ),
                destination_id=destination.id,
                normalized_destination=destination.name,
                quantity_ea=calculations[destination.id].actual_quantity,
                cost_won=calculations[destination.id].actual_cost_won,
                source_locator=(
                    "당월 실적 입력"
                    if destination.id == 2
                    else rf"D:\fictional\화성운반비내역(8월)!{11 + destination.id}"
                ),
            )
            for destination in destinations
        ),
        OperationEvidence(
            record_kind="nonregular",
            value_role="plan",
            input_kind="manual",
            batch_id=None,
            provenance_id="monthly-plan:2026-08:nonregular",
            report_month="2026-08",
            raw_destination=None,
            destination_id=None,
            normalized_destination="비정규 운반비",
            quantity_ea=Decimal("0"),
            cost_won=0,
            source_locator="당월 계획 입력",
        ),
        OperationEvidence(
            record_kind="nonregular",
            value_role="actual",
            input_kind="imported",
            batch_id=42,
            provenance_id="transport:2026-08:42",
            report_month="2026-08",
            raw_destination="비정규 원천",
            destination_id=None,
            normalized_destination="비정규 운반비",
            quantity_ea=Decimal("0"),
            cost_won=0,
            source_locator=r"D:\fictional\화성운반비내역(8월)!99",
        ),
        *(
            OperationEvidence(
                record_kind="destination",
                value_role="plan",
                input_kind="manual",
                batch_id=None,
                provenance_id=f"monthly-plan:2026-08:{destination.id}",
                report_month="2026-08",
                raw_destination=None,
                destination_id=destination.id,
                normalized_destination=destination.name,
                quantity_ea=calculations[destination.id].planned_quantity,
                cost_won=calculations[destination.id].planned_cost_won,
                source_locator="당월 계획 입력",
            )
            for destination in destinations
        ),
        OperationEvidence(
            record_kind="sales",
            value_role="plan",
            input_kind="manual",
            batch_id=None,
            provenance_id="monthly-sales-plan:2026-08",
            report_month="2026-08",
            raw_destination=None,
            destination_id=None,
            normalized_destination="매출액",
            quantity_ea=None,
            cost_won=sales.planned_won,
            source_locator="당월 매출 계획 입력",
        ),
        OperationEvidence(
            record_kind="sales",
            value_role="actual",
            input_kind="manual",
            batch_id=None,
            provenance_id="monthly-sales-actual:2026-08",
            report_month="2026-08",
            raw_destination=None,
            destination_id=None,
            normalized_destination="매출액",
            quantity_ea=None,
            cost_won=sales.actual_won,
            source_locator="월 매출 입력",
        ),
        *(
            OperationEvidence(
                record_kind="destination",
                value_role="plan",
                input_kind="manual",
                batch_id=None,
                provenance_id=f"monthly-plan:2026-09:{destination.id}",
                report_month="2026-09",
                raw_destination=None,
                destination_id=destination.id,
                normalized_destination=destination.name,
                quantity_ea=next_rows[destination.id - 1].calculation.planned_quantity,
                cost_won=next_rows[destination.id - 1].calculation.planned_cost_won,
                source_locator="다음 달 계획 입력",
            )
            for destination in destinations
        ),
        OperationEvidence(
            record_kind="nonregular",
            value_role="plan",
            input_kind="manual",
            batch_id=None,
            provenance_id="monthly-plan:2026-09:nonregular",
            report_month="2026-09",
            raw_destination=None,
            destination_id=None,
            normalized_destination="비정규 운반비",
            quantity_ea=Decimal("0"),
            cost_won=0,
            source_locator="다음 달 계획 입력",
        ),
        OperationEvidence(
            record_kind="sales",
            value_role="plan",
            input_kind="manual",
            batch_id=None,
            provenance_id="monthly-sales-plan:2026-09",
            report_month="2026-09",
            raw_destination=None,
            destination_id=None,
            normalized_destination="매출액",
            quantity_ea=None,
            cost_won=1_300_000,
            source_locator="다음 달 매출 계획 입력",
        ),
    )
    validation_context = ValidationContext(
        report_month="2026-08",
        sales=SalesInput(1_200_000, "2026-09-01", "월 매출 입력"),
        prior_year_history=historical_averages(
            "2026-08", total_row.quantity_by_month
        ).comparison_year,
        month_sources=ReportMonthSources(
            "2026-08", "2026-08", "2026-08", "2026-08", "2026-08"
        ),
        destinations=tuple(
            DestinationValidationInput(
                destination.id,
                destination.name,
                destination.display_order,
                destination.required_for_report,
                calculations[destination.id].actual_quantity,
                NextMonthPlanInput(
                    "2026-09", next_rows[destination.id - 1].calculation.planned_quantity
                ),
            )
            for destination in destinations
        ),
        current_import_batches=(
            ImportBatchInput(
                batch.batch_id,
                batch.report_month,
                batch.source_type,
                batch.file_sha256,
                batch.source_filename,
                True,
            ),
        ),
    )
    return ReviewWorkbookData(
        report=report,
        validation_context=validation_context,
        destinations=destinations,
        aliases=aliases,
        groups=groups,
        group_members=members,
        import_batches=(batch,),
        operations=operations,
    )


def _find_row(sheet, label: str) -> int:
    for row in sheet.iter_rows():
        if row[0].value == label:
            return row[0].row
    raise AssertionError(f"{label!r} row not found in {sheet.title}")


def _find_rows(sheet, label: str) -> list[int]:
    return [row[0].row for row in sheet.iter_rows() if row[0].value == label]


def _as_date(value):
    return value.date() if isinstance(value, datetime) else value


def _with_missing_first_actual(bundle):
    """Return a canonical report/context whose first required actual is missing."""
    from app.domain.calculations import calculate_destination, calculate_group, calculate_total

    first = bundle.report.rows[0]
    missing_calculation = calculate_destination(
        first.calculation.planned_quantity,
        first.calculation.planned_cost_won,
        None,
        None,
        destination_id=1,
    )
    first = replace(
        first,
        calculation=missing_calculation,
        quantity_by_month={**first.quantity_by_month, "2026-08": None},
        cost_won_by_month={**first.cost_won_by_month, "2026-08": None},
    )
    direct_rows = (first, *bundle.report.rows[1:12])
    calculations = {
        row.calculation.provenance.destination_id: row.calculation
        for row in direct_rows
    }
    group_calculation = calculate_group(
        calculations, bundle.group_members, group_id=10
    )
    group = replace(
        bundle.report.rows[12],
        calculation=group_calculation,
        quantity_by_month={
            **bundle.report.rows[12].quantity_by_month,
            "2026-08": group_calculation.actual_quantity,
        },
        cost_won_by_month={
            **bundle.report.rows[12].cost_won_by_month,
            "2026-08": group_calculation.actual_cost_won,
        },
    )
    total_calculation = calculate_total(calculations, bundle.destinations)
    total = replace(
        bundle.report.total,
        calculation=total_calculation,
        quantity_by_month={
            **bundle.report.total.quantity_by_month,
            "2026-08": total_calculation.actual_quantity,
        },
        cost_won_by_month={
            **bundle.report.total.cost_won_by_month,
            "2026-08": total_calculation.actual_cost_won,
        },
    )
    next_first = replace(
        bundle.report.next_month.rows[0],
        quantity_by_month={
            **bundle.report.next_month.rows[0].quantity_by_month,
            "2026-08": None,
        },
        cost_won_by_month={
            **bundle.report.next_month.rows[0].cost_won_by_month,
            "2026-08": None,
        },
    )
    next_total = replace(
        bundle.report.next_month.total,
        quantity_by_month={
            **bundle.report.next_month.total.quantity_by_month,
            "2026-08": None,
        },
        cost_won_by_month={
            **bundle.report.next_month.total.cost_won_by_month,
            "2026-08": None,
        },
    )
    report = replace(
        bundle.report,
        rows=(*direct_rows, group, bundle.report.rows[13]),
        total=total,
        next_month=replace(
            bundle.report.next_month,
            rows=(next_first, *bundle.report.next_month.rows[1:]),
            total=next_total,
        ),
        charts=replace(
            bundle.report.charts,
            actual_quantity_by_month={
                **bundle.report.charts.actual_quantity_by_month,
                "2026-08": None,
            },
            actual_cost_won_by_month={
                **bundle.report.charts.actual_cost_won_by_month,
                "2026-08": None,
            },
        ),
    )
    context = replace(
        bundle.validation_context,
        destinations=(
            replace(bundle.validation_context.destinations[0], actual_quantity=None),
            *bundle.validation_context.destinations[1:],
        ),
    )
    operations = (
        replace(bundle.operations[0], quantity_ea=None, cost_won=None),
        *bundle.operations[1:],
    )
    return replace(
        bundle,
        report=report,
        validation_context=context,
        operations=operations,
    )


def test_export_review_workbook_reconciles_typed_values_and_review_evidence(tmp_path):
    from app.domain.calculations import HistoricalValueKind, historical_averages
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    output = tmp_path / "review.xlsx"
    export_review_workbook(bundle, output)

    workbook = load_workbook(output, data_only=False, keep_links=True)
    assert workbook.sheetnames == ["월간 종합", "운행실적", "검증 결과", "마스터 기준"]
    assert workbook.vba_archive is None
    assert not workbook._external_links

    summary = workbook["월간 종합"]
    assert summary.sheet_view.showGridLines is False
    assert summary.freeze_panes == "C9"
    assert summary.auto_filter.ref
    assert _as_date(summary["B3"].value) == date(2026, 8, 1)
    assert summary["B3"].is_date
    total_row = _find_row(summary, "합계")
    total = bundle.report.total.calculation
    assert summary.cell(total_row, 3).value == float(total.planned_quantity)
    assert summary.cell(total_row, 4).value == float(total.actual_quantity)
    assert summary.cell(total_row, 7).value == total.planned_cost_won
    assert summary.cell(total_row, 8).value == total.actual_cost_won
    assert summary.cell(total_row, 6).data_type == "n"
    assert summary.cell(total_row, 6).number_format == "0.0%"
    zero_row = _find_row(summary, "가상납품처02")
    assert summary.cell(zero_row, 4).value == 0
    assert summary.cell(zero_row, 12).value is None
    assert summary.cell(zero_row, 14).value is None
    assert summary.cell(zero_row, 4).fill.fgColor.rgb == summary.cell(zero_row, 3).fill.fgColor.rgb
    assert summary.cell(zero_row, 8).fill.fgColor.rgb == summary.cell(zero_row, 7).fill.fgColor.rgb
    assert summary.cell(zero_row, 8).fill.fgColor.rgb != summary.cell(9, 8).fill.fgColor.rgb

    sales_row = _find_row(summary, "매출액")
    assert summary.cell(sales_row, 3).value == 1_100_000
    assert summary.cell(sales_row, 4).value == 1_200_000
    ratio_row = _find_row(summary, "운반비/매출")
    assert isinstance(summary.cell(ratio_row, 4).value, (int, float))
    assert summary.cell(ratio_row, 4).number_format == "0.0%"
    averages_row = _find_row(summary, "수량(EA)")
    assert isinstance(summary.cell(averages_row, 2).value, (int, float))
    assert summary.cell(averages_row, 3).value == "완전"

    current_first = _find_rows(summary, "가상납품처01")[0]
    current_averages = historical_averages(
        bundle.report.report_month,
        bundle.report.rows[0].cost_won_by_month,
        value_kind=HistoricalValueKind.MONEY,
    )
    assert summary.cell(current_first, 15).value == float(
        current_averages.three_month.value
    )
    assert summary.cell(current_first, 18).value == float(
        current_averages.comparison_year.value
    )

    plan_title = _find_row(summary, "2026년 9월 운반비 계획 검토")
    assert summary.cell(plan_title + 1, 1).value == "납품처"
    first_plan = plan_title + 2
    assert summary.cell(first_plan, 1).value == "가상납품처01"
    assert summary.cell(first_plan, 3).value == 131
    assert summary.cell(first_plan, 4).value == 14_410
    assert summary.cell(first_plan, 5).value == 110
    next_nonregular = first_plan + 12
    assert summary.cell(next_nonregular, 1).value == "비정규 운반비"
    next_total = next_nonregular + 1
    assert summary.cell(next_total, 1).value == "합계"
    assert summary.cell(next_total, 3).value == float(
        bundle.report.next_month.total.calculation.planned_quantity
    )
    assert summary.cell(next_total, 4).value == bundle.report.next_month.total.calculation.planned_cost_won
    assert summary.cell(next_total + 2, 1).value == "매출액"
    assert summary.cell(next_total + 2, 4).value == 1_300_000

    evidence = workbook["운행실적"]
    assert evidence.freeze_panes == "A10"
    assert evidence.auto_filter.ref
    assert evidence["B4"].value == "transport"
    assert evidence["D4"].value == "화성 가상자료.xlsx"
    assert evidence["E4"].value == "ab" * 32
    assert evidence["F4"].value == datetime(2026, 9, 1, 8, 30, 15)
    assert evidence["F4"].is_date
    assert evidence["A4"].value == 42
    assert _as_date(evidence["F10"].value) == date(2026, 8, 1)
    assert evidence["F10"].is_date
    assert evidence["A10"].value == "가져오기"
    assert evidence["B10"].value == "납품처"
    assert evidence["C10"].value == "실적"
    assert evidence["K10"].value == "화성운반비내역(8월)!12"
    assert "D:\\fictional" not in evidence["K10"].value
    assert evidence["A11"].value == "수기 입력"
    assert evidence["I11"].value == 0
    assert evidence["J11"].value == 0
    assert evidence["B22"].value == "비정규"
    assert evidence["C22"].value == "계획"
    assert evidence["I22"].value == 0
    assert evidence["J22"].value == 0
    assert evidence["B23"].value == "비정규"
    assert evidence["C23"].value == "실적"
    assert evidence["A10"].fill.fgColor.rgb != evidence["A11"].fill.fgColor.rgb

    checks = workbook["검증 결과"]
    assert checks["B3"].value == "가능"
    assert checks["B4"].value == 0
    assert checks.freeze_panes == "A8"
    assert checks.auto_filter.ref
    assert checks["B8"].value == "정상"
    assert checks["B8"].fill.fgColor.rgb in {"00000000", "00FFFFFF"}

    masters = workbook["마스터 기준"]
    assert masters.protection.sheet is True
    assert masters.freeze_panes == "A6"
    assert masters.auto_filter.ref
    assert masters["B6"].value == "가상납품처01"
    assert "원천납품처01 (transport)" in masters["C6"].value
    group_header = _find_row(masters, "그룹명")
    assert masters.cell(group_header + 1, 1).value == "가상권역"
    assert masters.cell(group_header + 3, 1).value == "빈 가상권역"
    assert masters.cell(group_header + 3, 4).value is None

    for sheet in workbook.worksheets:
        assert any(cell.style_id for row in sheet.iter_rows() for cell in row)
        assert all(dimension.width and dimension.width < 60 for dimension in sheet.column_dimensions.values())
        for row in sheet.iter_rows():
            for cell in row:
                assert not (isinstance(cell.value, str) and cell.value in ERROR_TOKENS)
                assert cell.data_type != "e"
                assert not (isinstance(cell.value, str) and cell.value.startswith("="))


def test_export_review_workbook_shows_neutral_valid_state(tmp_path):
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    output = tmp_path / "valid.xlsx"
    export_review_workbook(bundle, output)

    checks = load_workbook(output)["검증 결과"]
    assert checks["B3"].value == "가능"
    assert checks["B4"].value == 0
    assert checks["A8"].value == "VALID"
    assert checks["B8"].value == "정상"
    assert checks["A8"].fill.fgColor.rgb in {"00000000", "00FFFFFF"}


def test_export_review_workbook_shows_blocking_errors_distinctly(tmp_path):
    from app.domain.validation import UnresolvedDestinationAlias
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    bundle = replace(
        bundle,
        validation_context=replace(
            bundle.validation_context,
            unresolved_aliases=(
                UnresolvedDestinationAlias(
                    "미등록 원천", r"D:\secret\source.xlsx!A2"
                ),
            ),
        ),
    )
    output = tmp_path / "blocked.xlsx"
    export_review_workbook(bundle, output)

    checks = load_workbook(output)["검증 결과"]
    assert checks["B3"].value == "불가"
    assert checks["B4"].value == 1
    assert checks["B8"].value == "오류"
    assert checks["B8"].fill.fgColor.rgb != "00000000"
    all_text = "\n".join(
        cell.value
        for sheet in load_workbook(output).worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    )
    assert r"D:\secret" not in all_text


def test_export_review_workbook_rejects_forged_validation_result(tmp_path):
    from app.domain.validation import UnresolvedDestinationAlias, ValidationResult
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    invalid = replace(
        bundle,
        validation_context=replace(
            bundle.validation_context,
            unresolved_aliases=(UnresolvedDestinationAlias("누락", "가상.xlsx!A2"),),
        ),
        validation_result=ValidationResult(()),
    )
    with pytest.raises(ValueError, match="supplied validation result"):
        export_review_workbook(invalid, tmp_path / "forged.xlsx")


@pytest.mark.parametrize("missing", ["sales", "actual", "next-plan", "history"])
def test_export_review_workbook_rejects_validation_snapshot_conflicting_with_report(
    tmp_path, missing
):
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    context = bundle.validation_context
    if missing == "sales":
        context = replace(context, sales=None)
    elif missing == "actual":
        item = replace(context.destinations[0], actual_quantity=None)
        context = replace(context, destinations=(item, *context.destinations[1:]))
    elif missing == "next-plan":
        item = replace(context.destinations[0], next_month_plan=None)
        context = replace(context, destinations=(item, *context.destinations[1:]))
    else:
        context = replace(context, prior_year_history=None)
    output = tmp_path / f"missing-{missing}.xlsx"
    output.write_bytes(b"preserve")

    with pytest.raises(ValueError, match="validation snapshot"):
        export_review_workbook(replace(bundle, validation_context=context), output)

    assert output.read_bytes() == b"preserve"


def test_export_review_workbook_exports_missing_actual_as_blocking_blank(tmp_path):
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _with_missing_first_actual(_report_bundle())
    output = tmp_path / "missing-actual.xlsx"

    export_review_workbook(bundle, output)

    workbook = load_workbook(output, data_only=False)
    summary = workbook["월간 종합"]
    row = _find_row(summary, "가상납품처01")
    assert summary.cell(row, 4).value is None
    assert summary.cell(row, 8).value is None
    evidence = workbook["운행실적"]
    assert evidence["I10"].value is None
    assert evidence["J10"].value is None
    checks = workbook["검증 결과"]
    assert checks["B3"].value == "불가"
    assert checks["B4"].value == 1
    assert checks["A8"].value == "MISSING_ACTUAL_QUANTITY"


@pytest.mark.parametrize("case", ["fabricated-zero", "missing-zero-evidence"])
def test_export_review_workbook_requires_exact_evidence_for_zero(tmp_path, case):
    from app.reporting.xlsx_review import export_review_workbook

    if case == "fabricated-zero":
        bundle = _with_missing_first_actual(_report_bundle())
        operations = (
            replace(bundle.operations[0], quantity_ea=Decimal("0"), cost_won=0),
            *bundle.operations[1:],
        )
    else:
        bundle = _report_bundle()
        operations = (
            bundle.operations[0],
            replace(bundle.operations[1], quantity_ea=None, cost_won=None),
            *bundle.operations[2:],
        )

    with pytest.raises(ValueError, match="operation evidence does not reconcile"):
        export_review_workbook(
            replace(bundle, operations=operations), tmp_path / f"{case}.xlsx"
        )


@pytest.mark.parametrize("case", ["duplicate", "missing"])
def test_export_review_workbook_rejects_duplicate_or_missing_report_identity(tmp_path, case):
    from app.domain.calculations import calculate_destination
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    first = bundle.report.rows[0]
    if case == "duplicate":
        rows = (first, first, *bundle.report.rows[2:])
        message = "duplicate"
    else:
        rows = (
            replace(first, calculation=calculate_destination(
                first.calculation.planned_quantity,
                first.calculation.planned_cost_won,
                first.calculation.actual_quantity,
                first.calculation.actual_cost_won,
            )),
            *bundle.report.rows[1:],
        )
        message = "missing destination identity"
    invalid = replace(bundle, report=replace(bundle.report, rows=rows))
    output = tmp_path / f"{case}.xlsx"

    with pytest.raises(ValueError, match=message):
        export_review_workbook(invalid, output)

    assert not output.exists()


@pytest.mark.parametrize(
    ("case", "message"),
    [
        ("alias", "alias identity"),
        ("evidence", "evidence does not reconcile"),
        ("duplicate-evidence", "duplicate identity"),
        ("batch-month", "validation snapshot import provenance"),
        ("group", "group calculation"),
        ("total", "total calculation"),
    ],
)
def test_export_review_workbook_rejects_inconsistent_audit_provenance(
    tmp_path, case, message
):
    from app.domain.calculations import calculate_group, calculate_total
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    if case == "alias":
        first = replace(bundle.operations[0], raw_destination="미등록 원천")
        invalid = replace(bundle, operations=(first, *bundle.operations[1:]))
    elif case == "evidence":
        first = replace(bundle.operations[0], quantity_ea=Decimal("119"))
        invalid = replace(bundle, operations=(first, *bundle.operations[1:]))
    elif case == "duplicate-evidence":
        invalid = replace(bundle, operations=(*bundle.operations, bundle.operations[0]))
    elif case == "batch-month":
        invalid = replace(
            bundle,
            import_batches=(replace(bundle.import_batches[0], report_month="2026-07"),),
        )
    elif case == "group":
        calculations = {
            row.calculation.provenance.destination_id: row.calculation
            for row in bundle.report.rows[:12]
            if row is not None
        }
        wrong = calculate_group(calculations, bundle.group_members[:1], group_id=10)
        rows = (*bundle.report.rows[:12], replace(bundle.report.rows[12], calculation=wrong), bundle.report.rows[13])
        invalid = replace(bundle, report=replace(bundle.report, rows=rows))
    else:
        direct_rows = tuple(row for row in bundle.report.rows if row is not None)[:12]
        calculations = {
            row.calculation.provenance.destination_id: row.calculation
            for row in direct_rows
        }
        wrong = calculate_total(calculations, bundle.destinations[:-1])
        total = replace(
            bundle.report.total,
            calculation=wrong,
            quantity_by_month={
                **bundle.report.total.quantity_by_month,
                "2026-08": wrong.actual_quantity,
            },
            cost_won_by_month={
                **bundle.report.total.cost_won_by_month,
                "2026-08": wrong.actual_cost_won,
            },
        )
        next_total = replace(
            bundle.report.next_month.total,
            quantity_by_month={
                **bundle.report.next_month.total.quantity_by_month,
                "2026-08": wrong.actual_quantity,
            },
            cost_won_by_month={
                **bundle.report.next_month.total.cost_won_by_month,
                "2026-08": wrong.actual_cost_won,
            },
        )
        charts = replace(
            bundle.report.charts,
            planned_quantity_by_month={
                **bundle.report.charts.planned_quantity_by_month,
                "2026-08": wrong.planned_quantity,
            },
            actual_quantity_by_month={
                **bundle.report.charts.actual_quantity_by_month,
                "2026-08": wrong.actual_quantity,
            },
            planned_cost_won_by_month={
                **bundle.report.charts.planned_cost_won_by_month,
                "2026-08": wrong.planned_cost_won,
            },
            actual_cost_won_by_month={
                **bundle.report.charts.actual_cost_won_by_month,
                "2026-08": wrong.actual_cost_won,
            },
        )
        invalid = replace(
            bundle,
            report=replace(
                bundle.report,
                total=total,
                next_month=replace(bundle.report.next_month, total=next_total),
                charts=charts,
            ),
        )
    output = tmp_path / f"{case}.xlsx"

    with pytest.raises(ValueError, match=message):
        export_review_workbook(invalid, output)

    assert not output.exists()


def test_export_review_workbook_canonical_preflight_preserves_existing_output(tmp_path):
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    chart_values = dict(bundle.report.charts.actual_quantity_by_month)
    chart_values["2026-07"] += Decimal("1")
    invalid = replace(
        bundle,
        report=replace(
            bundle.report,
            charts=replace(
                bundle.report.charts, actual_quantity_by_month=chart_values
            ),
        ),
    )
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"preserve")

    with pytest.raises(ValueError, match="actual history conflict"):
        export_review_workbook(invalid, output)

    assert output.read_bytes() == b"preserve"


def test_export_review_workbook_missing_canonical_sales_preserves_existing_output(tmp_path):
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    invalid = replace(
        bundle,
        report=replace(
            bundle.report,
            sales=replace(bundle.report.sales, actual_won=None),
        ),
    )
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"preserve")

    with pytest.raises(ValueError, match="current actual sales"):
        export_review_workbook(invalid, output)

    assert output.read_bytes() == b"preserve"


@pytest.mark.parametrize("case", ["missing", "conflict"])
def test_export_review_workbook_rejects_nonregular_without_exact_evidence(
    tmp_path, case
):
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    if case == "missing":
        operations = (*bundle.operations[:13], *bundle.operations[14:])
    else:
        operations = (
            *bundle.operations[:13],
            replace(bundle.operations[13], cost_won=1),
            *bundle.operations[14:],
        )
    with pytest.raises(ValueError, match="nonregular evidence"):
        export_review_workbook(
            replace(bundle, operations=operations), tmp_path / f"{case}.xlsx"
        )


def test_export_review_workbook_rejects_missing_nonregular_without_matching_blocker(tmp_path):
    from app.domain.calculations import calculate_destination
    from app.domain.validation import UnresolvedDestinationAlias
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    nonregular = replace(
        bundle.report.nonregular,
        calculation=calculate_destination(Decimal("0"), 0, None, None),
        quantity_by_month={**bundle.report.nonregular.quantity_by_month, "2026-08": None},
        cost_won_by_month={**bundle.report.nonregular.cost_won_by_month, "2026-08": None},
    )
    next_nonregular = replace(
        bundle.report.next_month.nonregular,
        quantity_by_month={
            **bundle.report.next_month.nonregular.quantity_by_month,
            "2026-08": None,
        },
        cost_won_by_month={
            **bundle.report.next_month.nonregular.cost_won_by_month,
            "2026-08": None,
        },
    )
    report = replace(
        bundle.report,
        nonregular=nonregular,
        next_month=replace(bundle.report.next_month, nonregular=next_nonregular),
    )
    operations = (
        *bundle.operations[:13],
        replace(bundle.operations[13], quantity_ea=None, cost_won=None),
        *bundle.operations[14:],
    )
    output = tmp_path / "missing-nonregular.xlsx"
    context = replace(
        bundle.validation_context,
        unresolved_aliases=(
            UnresolvedDestinationAlias("미등록 비정규", "비정규 실적 입력"),
        ),
    )

    with pytest.raises(ValueError, match="no blocking validation issue"):
        export_review_workbook(
            replace(
                bundle,
                report=report,
                operations=operations,
                validation_context=context,
            ),
            output,
        )


def test_export_review_workbook_requires_one_evidence_record_per_present_input(tmp_path):
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    current_plan_index = next(
        index
        for index, item in enumerate(bundle.operations)
        if item.record_kind == "destination"
        and item.value_role == "plan"
        and item.report_month == "2026-08"
        and item.destination_id == 1
    )
    missing = (
        *bundle.operations[:current_plan_index],
        *bundle.operations[current_plan_index + 1 :],
    )
    with pytest.raises(ValueError, match="evidence does not reconcile"):
        export_review_workbook(
            replace(bundle, operations=missing), tmp_path / "missing-input.xlsx"
        )

    duplicate = replace(
        bundle.operations[current_plan_index],
        provenance_id="duplicate-plan-provenance",
    )
    with pytest.raises(ValueError, match="exactly one evidence"):
        export_review_workbook(
            replace(bundle, operations=(*bundle.operations, duplicate)),
            tmp_path / "duplicate-input.xlsx",
        )


def test_unrelated_global_blocker_does_not_excuse_missing_input_evidence(tmp_path):
    from app.domain.validation import UnresolvedDestinationAlias
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    current_plan_index = next(
        index
        for index, item in enumerate(bundle.operations)
        if item.record_kind == "destination"
        and item.value_role == "plan"
        and item.report_month == "2026-08"
        and item.destination_id == 1
    )
    operations = (
        *bundle.operations[:current_plan_index],
        *bundle.operations[current_plan_index + 1 :],
    )
    context = replace(
        bundle.validation_context,
        unresolved_aliases=(
            UnresolvedDestinationAlias("무관한 원천", "별칭 검토"),
        ),
    )

    with pytest.raises(ValueError, match="evidence does not reconcile"):
        export_review_workbook(
            replace(bundle, operations=operations, validation_context=context),
            tmp_path / "unrelated-blocker.xlsx",
        )


def test_xlsx_recomputes_next_month_total_without_trusting_ppt_validator(
    tmp_path, monkeypatch
):
    from app.domain.calculations import calculate_destination
    from app.reporting import xlsx_review

    bundle = _report_bundle()
    first = bundle.report.next_month.rows[0]
    changed = replace(
        first,
        calculation=calculate_destination(
            first.calculation.planned_quantity,
            first.calculation.planned_cost_won + 1,
            None,
            None,
            destination_id=first.calculation.destination_id,
        ),
    )
    report = replace(
        bundle.report,
        next_month=replace(
            bundle.report.next_month,
            rows=(changed, *bundle.report.next_month.rows[1:]),
        ),
    )
    monkeypatch.setattr(xlsx_review, "validate_ppt_report", lambda report: None)

    with pytest.raises(ValueError, match="next-month total calculation"):
        xlsx_review.export_review_workbook(
            replace(bundle, report=report), tmp_path / "stale-next-total.xlsx"
        )


def test_export_review_workbook_exports_one_current_batch_per_source_type(tmp_path):
    from app.domain.models import DestinationAlias
    from app.domain.validation import ImportBatchInput
    from app.reporting.xlsx_review import ImportBatchEvidence, export_review_workbook

    bundle = _report_bundle()
    second = ImportBatchEvidence(
        batch_id=43,
        provenance_id="erp:2026-08:43",
        report_month="2026-08",
        source_type="erp",
        source_filename=r"D:\fictional\ERP 가상자료.xlsx",
        file_sha256="cd" * 32,
        imported_at=datetime(2026, 9, 1, 9, 45),
    )
    first_operation = replace(
        bundle.operations[0],
        batch_id=43,
        provenance_id=second.provenance_id,
        raw_destination="ERP원천01",
        source_locator=r"D:\fictional\ERP 가상자료.xlsx!A12",
    )
    context = replace(
        bundle.validation_context,
        current_import_batches=(
            *bundle.validation_context.current_import_batches,
            ImportBatchInput(
                second.batch_id,
                second.report_month,
                second.source_type,
                second.file_sha256,
                second.source_filename,
                True,
            ),
        ),
    )
    output = tmp_path / "two-sources.xlsx"
    export_review_workbook(
        replace(
            bundle,
            validation_context=context,
            aliases=(
                *bundle.aliases,
                DestinationAlias(99, 1, "ERP원천01", "erp"),
            ),
            import_batches=(*bundle.import_batches, second),
            operations=(first_operation, *bundle.operations[1:]),
        ),
        output,
    )

    evidence = load_workbook(output)["운행실적"]
    assert evidence["B4"].value == "transport"
    assert evidence["B5"].value == "erp"
    assert evidence["D4"].value == "화성 가상자료.xlsx"
    assert evidence["D5"].value == "ERP 가상자료.xlsx"


def test_export_review_workbook_rejects_duplicate_batch_source_type(tmp_path):
    from app.domain.validation import ImportBatchInput
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    duplicate = replace(
        bundle.import_batches[0],
        batch_id=99,
        provenance_id="transport:2026-08:99",
    )
    context = replace(
        bundle.validation_context,
        current_import_batches=(
            *bundle.validation_context.current_import_batches,
            ImportBatchInput(
                duplicate.batch_id,
                duplicate.report_month,
                duplicate.source_type,
                duplicate.file_sha256,
                duplicate.source_filename,
                True,
            ),
        ),
    )
    with pytest.raises(ValueError, match="source type.*duplicate"):
        export_review_workbook(
            replace(
                bundle,
                validation_context=context,
                import_batches=(*bundle.import_batches, duplicate),
            ),
            tmp_path / "duplicate-source.xlsx",
        )


@pytest.mark.parametrize(
    ("value", "safe_fragment", "secret_fragments"),
    [
        (
            '원본 "D:\\secret folder\\private data\\manifest!A2", 확인',
            "manifest!A2",
            (r"D:\secret folder", "private data"),
        ),
        (
            r"원본 D:\secret folder\private data\source.parquet!A2, 확인",
            "source.parquet!A2",
            (r"D:\secret folder", "private data"),
        ),
        (
            r"원본 \\server\share name\secret folder\audit.log!B4; 확인",
            "audit.log!B4",
            (r"\\server\share name", "secret folder"),
        ),
        (
            "원본 (/home/user folder/private data/source file), 확인",
            "source file",
            ("/home/user folder", "private data"),
        ),
        (
            "원본 '/var/log/private folder/import.log!C7'\n확인",
            "import.log!C7",
            ("/var/log", "private folder"),
        ),
    ],
)
def test_safe_text_redacts_extension_independent_absolute_paths(
    value, safe_fragment, secret_fragments
):
    from app.reporting.xlsx_review import _contains_absolute_path, _safe_display_text

    assert _contains_absolute_path(value) is True
    safe = _safe_display_text(value)

    assert safe_fragment in safe
    assert _contains_absolute_path(safe) is False
    for fragment in secret_fragments:
        assert fragment not in safe


def test_safe_text_preserves_plain_comparison_slash_and_masks_root_only_paths():
    from app.reporting.xlsx_review import _contains_absolute_path, _safe_display_text

    assert _safe_display_text("계획 / 실적 비교") == "계획 / 실적 비교"
    assert _safe_display_text('원본 "D:\\"') == '원본 "경로 숨김"'
    for value in (
        "https://example.com/a",
        "http://example.com/a/b",
        "시간 12:30",
        "계획:실적 비율 1:2",
    ):
        assert _safe_display_text(value) == value
        assert _contains_absolute_path(value) is False


@pytest.mark.parametrize(
    "value",
    [
        "https://example.com/a/b?next=/c/d#frag/path",
        "https://example.com?next=/a/b#section/path",
        "ftp://example.com/a/b?next=/c/d#frag/path",
        "ftps://example.com/a/b?next=/c/d#frag/path",
        "계획 /실적 비교",
        "계획/실적",
        "/실적",
    ],
)
def test_safe_text_preserves_full_urls_and_ordinary_slash_phrases(value):
    from app.reporting.xlsx_review import _contains_absolute_path, _safe_display_text

    assert _safe_display_text(value) == value
    assert _contains_absolute_path(value) is False


@pytest.mark.parametrize(
    ("value", "safe"),
    [
        ("/home/user/file", "file"),
        ("/home/user folder/source.parquet", "source.parquet"),
        ("source:/home/user/file", "source:file"),
        ("/source.xlsx!A2", "source.xlsx!A2"),
    ],
)
def test_safe_text_redacts_only_credible_posix_absolute_paths(value, safe):
    from app.reporting.xlsx_review import _contains_absolute_path, _safe_display_text

    assert _contains_absolute_path(value) is True
    assert _safe_display_text(value) == safe
    assert _contains_absolute_path(safe) is False


@pytest.mark.parametrize(
    ("value", "safe_fragment"),
    [
        (r"원본:D:\secret folder\private\drive.parquet!A2, 확인", "drive.parquet!A2"),
        (r"source:\\server\share\private\unc.log!B3; 확인", "unc.log!B3"),
        ("source:/home/user/private/posix!C4, 확인", "posix!C4"),
    ],
)
def test_safe_text_redacts_colon_adjacent_local_paths(value, safe_fragment):
    from app.reporting.xlsx_review import _contains_absolute_path, _safe_display_text

    assert _contains_absolute_path(value) is True
    safe = _safe_display_text(value)
    assert safe_fragment in safe
    assert _contains_absolute_path(safe) is False


def test_validation_message_rows_expand_within_visual_bounds(tmp_path):
    from app.domain.validation import UnresolvedDestinationAlias
    from app.reporting.xlsx_review import _wrapped_line_count, export_review_workbook

    bundle = _report_bundle()
    context = replace(
        bundle.validation_context,
        unresolved_aliases=(
            UnresolvedDestinationAlias(
                "매우 긴 가상 원천 납품처 이름 " * 3,
                "계획 / 실적 비교 화면에서 확인해야 하는 매우 긴 수기 위치 설명 " * 2,
            ),
        ),
    )
    output = tmp_path / "long-validation.xlsx"
    export_review_workbook(replace(bundle, validation_context=context), output)

    checks = load_workbook(output)["검증 결과"]
    height = checks.row_dimensions[8].height
    message = checks["C8"]
    locator = checks["F8"]
    assert 36 < height <= 120
    assert message.alignment.wrap_text is True
    assert locator.alignment.wrap_text is True
    assert message.font.sz >= 9
    assert locator.font.sz >= 9
    required_lines = max(
        _wrapped_line_count(
            message.value, checks.column_dimensions["C"].width, message.font.sz
        ),
        _wrapped_line_count(
            locator.value, checks.column_dimensions["F"].width, locator.font.sz
        ),
    )
    line_capacity = int((height - 6) // (message.font.sz * 1.5))
    assert required_lines >= 3
    assert required_lines <= line_capacity


def test_validation_layout_uses_font_metrics_for_wide_latin_and_long_text(tmp_path):
    from app.domain.validation import UnresolvedDestinationAlias
    from app.reporting.xlsx_review import (
        _validation_line_capacity,
        _wrapped_line_count,
        export_review_workbook,
    )

    assert _wrapped_line_count("W" * 200, 59, 9) >= 6
    bundle = _report_bundle()
    alias = "W" * 200
    locator = "M" * 300
    context = replace(
        bundle.validation_context,
        unresolved_aliases=(UnresolvedDestinationAlias(alias, locator),),
    )
    output = tmp_path / "wide-latin-validation.xlsx"
    export_review_workbook(replace(bundle, validation_context=context), output)

    checks = load_workbook(output)["검증 결과"]
    message_width = checks.column_dimensions["C"].width
    locator_width = checks.column_dimensions["F"].width
    assert 59 < max(message_width, locator_width) <= 80
    message_text = "".join(
        str(checks.cell(row, 3).value or "").replace("\n", "")
        for row in range(8, checks.max_row + 1)
    )
    locator_text = "".join(
        str(checks.cell(row, 6).value or "").replace("\n", "")
        for row in range(8, checks.max_row + 1)
    )
    assert alias in message_text
    assert locator in message_text
    assert locator in locator_text
    for row in range(8, checks.max_row + 1):
        height = checks.row_dimensions[row].height
        font_size = checks.cell(row, 3).font.sz
        assert font_size >= 9
        assert 36 <= height <= 120
        required = max(
            _wrapped_line_count(
                str(checks.cell(row, 3).value or ""), message_width, font_size
            ),
            _wrapped_line_count(
                str(checks.cell(row, 6).value or ""), locator_width, font_size
            ),
        )
        assert required <= _validation_line_capacity(height, font_size)


def test_validation_layout_uses_continuation_rows_before_clipping(tmp_path):
    from app.domain.validation import UnresolvedDestinationAlias
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    alias = "W" * 2_000
    locator = "M" * 2_000
    context = replace(
        bundle.validation_context,
        unresolved_aliases=(UnresolvedDestinationAlias(alias, locator),),
    )
    output = tmp_path / "continuation-validation.xlsx"
    export_review_workbook(replace(bundle, validation_context=context), output)

    checks = load_workbook(output)["검증 결과"]
    assert checks.max_row > 8
    assert any(
        "계속" in str(checks.cell(row, 1).value or "")
        for row in range(9, checks.max_row + 1)
    )
    assert all(
        36 <= checks.row_dimensions[row].height <= 120
        for row in range(8, checks.max_row + 1)
    )


def test_saved_workbook_validation_detects_path_leak_independently(
    tmp_path, monkeypatch
):
    from app.reporting import xlsx_review

    bundle = _report_bundle()
    output = tmp_path / "leaked-path.xlsx"
    xlsx_review.export_review_workbook(bundle, output)
    workbook = load_workbook(output)
    leaked = r"원본:D:\secret folder\private data\payload.parquet!A2, 확인"
    workbook["검증 결과"]["C8"] = leaked
    workbook.save(output)
    original_sanitizer = xlsx_review._safe_display_text
    monkeypatch.setattr(
        xlsx_review,
        "_safe_display_text",
        lambda value: value if value == leaked else original_sanitizer(value),
    )

    with pytest.raises(ValueError, match="exposes an absolute path"):
        xlsx_review._validate_saved_workbook(output, bundle)


@pytest.mark.parametrize(
    "value",
    [
        "https://example.com/a/b?next=/c/d#frag/path",
        "https://example.com?next=/a/b#section/path",
        "ftp://example.com/a/b?next=/c/d#frag/path",
        "ftps://example.com/a/b?next=/c/d#frag/path",
        "계획 /실적 비교",
        "계획/실적",
        "/실적",
    ],
)
def test_export_review_workbook_preserves_urls_and_ordinary_slash_phrases(
    tmp_path, value
):
    from app.domain.validation import UnresolvedDestinationAlias
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    context = replace(
        bundle.validation_context,
        unresolved_aliases=(UnresolvedDestinationAlias(value, value),),
    )
    output = tmp_path / "preserved-text.xlsx"
    export_review_workbook(replace(bundle, validation_context=context), output)

    text = "\n".join(
        cell.value
        for sheet in load_workbook(output).worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    )
    assert value in text


@pytest.mark.parametrize(
    ("locator", "safe_fragment", "secret_fragments"),
    [
        (
            r'D:\secret folder\private data\source file.xlsx!A2',
            "source file.xlsx!A2",
            (r"D:\secret folder", "private data"),
        ),
        (
            r"\\server\share name\secret folder\source file.xlsx!A2",
            "source file.xlsx!A2",
            (r"\\server\share name", "secret folder"),
        ),
        (
            "/home/user folder/private data/source file.xlsx!A2",
            "source file.xlsx!A2",
            ("/home/user folder", "private data"),
        ),
        (
            r"원본:D:\secret folder\private data\source.parquet!A2, 확인",
            "source.parquet!A2",
            (r"D:\secret folder", "private data"),
        ),
        (
            r"source:\\server\share name\secret folder\source.log!A2; 확인",
            "source.log!A2",
            (r"\\server\share name", "secret folder"),
        ),
        (
            "source:/home/user folder/private data/source file!A2, 확인",
            "source file!A2",
            ("/home/user folder", "private data"),
        ),
        (
            "/source.xlsx!A2",
            "source.xlsx!A2",
            ("/source.xlsx!A2",),
        ),
    ],
)
def test_export_review_workbook_sanitizes_paths_in_all_visible_text(
    tmp_path, locator, safe_fragment, secret_fragments
):
    from app.domain.validation import UnresolvedDestinationAlias
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    context = replace(
        bundle.validation_context,
        unresolved_aliases=(
            UnresolvedDestinationAlias(f'"{locator}"', f"({locator}), 원본"),
        ),
    )
    output = tmp_path / "safe.xlsx"
    export_review_workbook(replace(bundle, validation_context=context), output)

    text = "\n".join(
        cell.value
        for sheet in load_workbook(output).worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    )
    for fragment in secret_fragments:
        assert fragment not in text
    assert safe_fragment in text


@pytest.mark.parametrize(
    "case", ["stale-name", "duplicate-order", "fabricated", "inclusion"]
)
def test_export_review_workbook_rejects_inconsistent_group_member_snapshot(
    tmp_path, case
):
    from app.domain.models import GroupMember
    from app.reporting.xlsx_review import export_review_workbook

    bundle = _report_bundle()
    if case == "stale-name":
        members = (replace(bundle.group_members[0], name="오래된 이름"), *bundle.group_members[1:])
    elif case == "duplicate-order":
        members = (bundle.group_members[0], replace(bundle.group_members[1], display_order=1))
    elif case == "fabricated":
        members = (*bundle.group_members, GroupMember(10, 999, "조작", 3, True, True))
    else:
        members = (replace(bundle.group_members[0], include_cost=False), *bundle.group_members[1:])

    message = "group calculation" if case == "inclusion" else "group member"
    with pytest.raises(ValueError, match=message):
        export_review_workbook(
            replace(bundle, group_members=members), tmp_path / f"{case}.xlsx"
        )


def test_export_review_workbook_is_atomic_when_final_replace_fails(tmp_path, monkeypatch):
    from app.reporting import xlsx_review

    output = tmp_path / "existing.xlsx"
    original = b"existing review workbook"
    output.write_bytes(original)

    def fail_replace(source: Path, destination: Path) -> None:
        raise OSError("synthetic replace failure")

    monkeypatch.setattr(xlsx_review.os, "replace", fail_replace)
    with pytest.raises(OSError, match="synthetic replace failure"):
        xlsx_review.export_review_workbook(_report_bundle(), output)

    assert output.read_bytes() == original
    assert list(tmp_path.glob(".*.tmp.xlsx")) == []
