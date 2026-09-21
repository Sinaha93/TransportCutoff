"""Fictional data only. The business runtime asset is never a test fixture."""
from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
from datetime import date
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def _text(slide, name, text, x=.3, y=.2, w=9, h=.4):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        paragraph.runs[0].font.name = 'Malgun Gothic'
        paragraph.runs[0].font.size = Pt(16)
    return shape


def _table(slide, name, rows, cols, y, h):
    shape = slide.shapes.add_table(rows, cols, Inches(.3), Inches(y), Inches(10.1), Inches(h))
    shape.name = name
    for row in shape.table.rows:
        for cell in row.cells:
            cell.text = '0'
            cell.margin_left = cell.margin_right = Inches(.015)
            cell.margin_top = cell.margin_bottom = 0
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.RIGHT
            run = paragraph.runs[0]
            run.font.name = 'Malgun Gothic'
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor(10, 30, 90)
    return shape.table


def build_template(path):
    """Fresh six-slide design with matching semantic roles and merge topology."""
    p = Presentation()
    p.slide_width, p.slide_height = 9906000, 6858000
    for i in range(1, 7):
        slide = p.slides.add_slide(p.slide_layouts[6])
        _text(slide, 'report.title', '가상 보고서' if i != 1 else '가상 표지')
        _text(slide, 'report.date', '2000.1.1\n가상 물류팀' if i == 1 else '2000.1.1', 8.6, .65, 2, .5)
        if i in (3, 4, 5):
            blob = BytesIO()
            Image.new('RGB', (80, 50), (180, 220, 245)).save(blob, format='PNG')
            blob.seek(0)
            chart = slide.shapes.add_picture(blob, Inches(.4), Inches(1.3), Inches(10), Inches(5))
            chart.name = 'report.chart'
            chart.crop_left = .01
            chart.crop_right = .02
            blob.seek(0)
            slide.shapes.add_picture(blob, Inches(9.5), Inches(7), Inches(.3), Inches(.2)).name = 'unrelated.logo'
            if i in (3, 5):
                _text(slide, 'report.quantity_average', '100', .2, 1, 1, .3)
            if i in (4, 5):
                _text(slide, 'report.cost_average', '200', 8.8, 1, 1.3, .3)
    t = _table(p.slides[1], 'report.monthly_table', 21, 18, 1.1, 4.2)
    t.cell(0, 0).merge(t.cell(1, 1))
    for a, b in ((2, 4), (5, 7), (8, 13), (14, 17)):
        t.cell(0, a).merge(t.cell(0, b))
    for a in (8, 10, 12):
        t.cell(1, a).merge(t.cell(1, a + 1))
    for r in (*range(2, 11), *range(16, 21)):
        t.cell(r, 0).merge(t.cell(r, 1))
    t.cell(11, 0).merge(t.cell(14, 0))
    for r in (15, 16):
        t.cell(r, 2).merge(t.cell(r, 4))
        t.cell(r, 8).merge(t.cell(r, 9))
        t.cell(r, 12).merge(t.cell(r, 13))
    t.cell(16, 5).merge(t.cell(16, 7))
    t.cell(17, 12).merge(t.cell(17, 13))
    for r in (18, 19, 20):
        for a, b in ((2, 4), (5, 7), (8, 10), (11, 13)):
            t.cell(r, a).merge(t.cell(r, b))
    for r in range(2, 21):
        for cell in t.rows[r].cells:
            cell.text = '' if cell.is_spanned else '0'
            if not cell.is_spanned:
                paragraph = cell.text_frame.paragraphs[0]
                paragraph.alignment = PP_ALIGN.RIGHT
                paragraph.runs[0].font.name = 'Malgun Gothic'
                paragraph.runs[0].font.size = Pt(8)
                paragraph.runs[0].font.bold = True
                paragraph.runs[0].font.color.rgb = RGBColor(10, 30, 90)
    # A truly empty cell with explicit paragraph/end-paragraph style.
    cell = t.cell(2, 6)
    cell.text = ''
    cell.text_frame.paragraphs[0].font.name = 'Malgun Gothic'
    cell.text_frame.paragraphs[0].font.size = Pt(8)
    cell.text_frame.paragraphs[0].font.bold = True
    review = _table(p.slides[1], 'report.review_table', 7, 8, 5.5, 1.5)
    review.cell(0, 0).merge(review.cell(1, 0))
    review.cell(0, 1).merge(review.cell(0, 7))
    review.columns[7].width = Inches(3)
    for column in list(review.columns)[:7]:
        column.width = Inches(7.1 / 7)
    plan = _table(p.slides[5], 'report.plan_table', 19, 8, 1.1, 5.4)
    plan.cell(0, 0).merge(plan.cell(1, 0))
    plan.cell(0, 1).merge(plan.cell(0, 3))
    plan.cell(0, 4).merge(plan.cell(0, 7))
    plan.cell(14, 1).merge(plan.cell(14, 2))
    plan.cell(14, 3).merge(plan.cell(15, 3))
    for r in (16, 17, 18):
        plan.cell(r, 1).merge(plan.cell(r, 3))
    for row in list(plan.rows)[2:]:
        for cell in row.cells:
            cell.text = '' if cell.is_spanned else '0'
            if not cell.is_spanned:
                run = cell.text_frame.paragraphs[0].runs[0]
                run.font.name, run.font.size = 'Malgun Gothic', Pt(8)
    # Replace merge-concatenated placeholder paragraphs with fictional headers.
    for table in (t, review, plan):
        for r in (0, 1):
            for c, cell in enumerate(table.rows[r].cells):
                cell.text = '' if cell.is_spanned else ('항목' if c == 0 else '구분')
                if not cell.is_spanned:
                    run = cell.text_frame.paragraphs[0].runs[0]
                    run.font.name, run.font.size = 'Malgun Gothic', Pt(8)
                    run.font.bold = True
    p.save(path)
    return path


@pytest.fixture
def template(tmp_path):
    return build_template(tmp_path / 'synthetic-template.pptx')


def make_report():
    from app.domain.calculations import calculate_destination, calculate_total
    from app.domain.models import Destination
    from app.reporting.charts import ChartReport
    from app.reporting.pptx_report import PptReport, ReportRow, PlanReport, SalesReport, ReviewDetail
    months = [*(f'2025-{m:02}' for m in range(1, 13)), *(f'2026-{m:02}' for m in range(1, 10))]
    quantities = {m: Decimal(20000) for m in months}
    costs = {m: 10000000 for m in months}
    rows, next_rows, calculations, next_calculations, destinations = [], [], {}, {}, []
    for index in range(1, 13):
        dest = Destination(index, f'가상{index:02}', index, True, True, None, True, True, False)
        calculation = calculate_destination(Decimal(5000), 2000000, Decimal(5000), 2279583 if index < 12 else 2279587, destination_id=index)
        next_calc = calculate_destination(Decimal(6000), 2400000, None, None, destination_id=index)
        destinations.append(dest)
        calculations[index], next_calculations[index] = calculation, next_calc
        row_quantities = {**quantities, '2026-08': calculation.actual_quantity}
        row_costs = {**costs, '2026-08': calculation.actual_cost_won}
        rows.append(ReportRow(str(index), dest.name, calculation, row_quantities, row_costs))
        next_rows.append(ReportRow(str(index), dest.name, next_calc, row_quantities, row_costs))
    total = calculate_total(calculations, destinations)
    next_total = calculate_total(next_calculations, destinations)
    actual_quantities = {**quantities, '2026-08': total.actual_quantity}
    actual_costs = {**costs, '2026-08': total.actual_cost_won}
    total_row = ReportRow('total', '합계', total, actual_quantities, actual_costs)
    nonregular = ReportRow('extra', '비정규 운반비', calculate_destination(Decimal(0), 0, Decimal(0), 0), {**quantities, '2026-08': Decimal(0)}, {**costs, '2026-08': 0})
    chart = ChartReport('2026-08', {**quantities, '2026-08': total.planned_quantity}, {**quantities, '2026-08': total.actual_quantity}, {**costs, '2026-08': total.planned_cost_won}, {**costs, '2026-08': total.actual_cost_won})
    sales = SalesReport(300000000, 310000000, {**{m: 300000000 for m in months}, '2026-08': 310000000})
    return PptReport('2026-08', date(2026, 9, 1), tuple(rows) + (None, None), nonregular, total_row, sales,
        PlanReport('2026-09', tuple(next_rows), replace(nonregular, calculation=calculate_destination(Decimal(0), 0, None, None)), ReportRow('total', '합계', next_total, actual_quantities, actual_costs), SalesReport(400000000, None, sales.actual_won_by_month)), chart, ())


def with_actual_rows(report, *rows):
    """Update fictional actuals and the same identity's shared history together."""
    from app.domain.calculations import calculate_total
    from app.domain.models import Destination

    updates = {row.key: replace(row, quantity_by_month={**row.quantity_by_month, '2026-08': row.calculation.actual_quantity}, cost_won_by_month={**row.cost_won_by_month, '2026-08': row.calculation.actual_cost_won}) for row in rows}
    current = tuple(updates.get(row.key, row) if row else None for row in report.rows)
    next_rows = tuple(replace(row, quantity_by_month=updates[row.key].quantity_by_month, cost_won_by_month=updates[row.key].cost_won_by_month) if row and row.key in updates else row for row in report.next_month.rows)
    provenance = report.total.calculation.provenance
    rules = tuple(
        Destination(
            index,
            row.label,
            order,
            True,
            True,
            None,
            index in provenance.quantity_destination_ids,
            index in provenance.cost_destination_ids,
            False,
        )
        for order, row in enumerate(current[:12], start=1)
        for index in (row.calculation.destination_id,)
    )
    total = calculate_total(
        {row.calculation.destination_id: row.calculation for row in current[:12]},
        rules,
    )
    quantity_history = {
        **report.total.quantity_by_month,
        '2026-08': total.actual_quantity,
    }
    cost_history = {
        **report.total.cost_won_by_month,
        '2026-08': total.actual_cost_won,
    }
    return replace(
        report,
        rows=current,
        total=replace(
            report.total,
            calculation=total,
            quantity_by_month=quantity_history,
            cost_won_by_month=cost_history,
        ),
        next_month=replace(
            report.next_month,
            rows=next_rows,
            total=replace(
                report.next_month.total,
                quantity_by_month=quantity_history,
                cost_won_by_month=cost_history,
            ),
        ),
        charts=replace(
            report.charts,
            actual_quantity_by_month=quantity_history,
            actual_cost_won_by_month=cost_history,
        ),
    )


def reviewed_report(report, reason='가상 원인'):
    from app.domain.calculations import calculate_destination
    from app.reporting.pptx_report import ReviewDetail
    first = replace(report.rows[0], calculation=calculate_destination(Decimal(5000), 2000000, Decimal(4000), 2400000, destination_id=1))
    return replace(with_actual_rows(report, first), reviews=(ReviewDetail('1', 5, 4, Decimal(1200), reason),))


@pytest.fixture
def report():
    return make_report()


@pytest.fixture
def fast_charts(monkeypatch):
    """Validation regressions do not need to exercise Task10 rendering again."""
    from app.reporting import pptx_report
    def render(report, path):
        Image.new('RGB', (80, 50), 'white').save(path)
    for name in ('render_quantity_chart', 'render_cost_chart', 'render_combined_chart'):
        monkeypatch.setattr(pptx_report, name, render)


def named(p, slide, role):
    return next(s for s in p.slides[slide - 1].shapes if s.name == role)


def test_ppt_generation_updates_values_and_preserves_layout(template, report, tmp_path):
    from app.reporting.pptx_report import generate_pptx
    from tools.verify_ppt_layout import verify_ppt_layout
    before = sha256(template.read_bytes()).hexdigest()
    output = tmp_path / '26년 8월 운반비 보고.pptx'
    generate_pptx(template, report, output)
    p = Presentation(output)
    assert len(p.slides) == 6
    assert named(p, 1, 'report.title').text == '26년 8월 화성공장 운반비 보고'
    assert named(p, 6, 'report.title').text == '▣ 26년 9월 운반비 계획'
    assert named(p, 1, 'report.date').text == '2026.9.1\n가상 물류팀'
    t = named(p, 2, 'report.monthly_table').table
    assert t.cell(17, 0).text == '합계'
    assert t.cell(17, 6).text == '27,355'
    assert t.cell(2, 7).text == '456'
    assert t.cell(2, 14).text == '10,000'
    assert t.cell(18, 5).text == '456'
    assert t.cell(20, 5).text == '8.82%'
    plan = named(p, 6, 'report.plan_table').table
    assert plan.cell(2, 1).text == '6,000'
    assert plan.cell(15, 2).text == '28,800'
    assert plan.cell(17, 1).text == '400,000'
    assert named(p, 3, 'report.quantity_average').text == '20,000'
    assert named(p, 4, 'report.cost_average').text == '10,000'
    assert verify_ppt_layout(template, output) == []
    assert sha256(template.read_bytes()).hexdigest() == before


def test_ppt_preflight_rejects_mutated_next_month_row_with_stale_total(report):
    from app.domain.calculations import calculate_destination
    from app.reporting.pptx_report import validate_ppt_report

    first = report.next_month.rows[0]
    mutated = replace(
        first,
        calculation=calculate_destination(
            first.calculation.planned_quantity,
            first.calculation.planned_cost_won + 1,
            None,
            None,
            destination_id=first.calculation.destination_id,
        ),
    )
    invalid = replace(
        report,
        next_month=replace(
            report.next_month, rows=(mutated, *report.next_month.rows[1:])
        ),
    )

    with pytest.raises(ValueError, match="next-month total"):
        validate_ppt_report(invalid)


def test_ppt_preflight_rejects_non_destination_next_month_row_cleanly(report):
    from app.reporting.pptx_report import validate_ppt_report

    invalid_row = replace(
        report.next_month.rows[0], calculation=report.total.calculation
    )
    invalid = replace(
        report,
        next_month=replace(
            report.next_month,
            rows=(invalid_row, *report.next_month.rows[1:]),
        ),
    )

    with pytest.raises(ValueError, match="next-month destination identity"):
        validate_ppt_report(invalid)


def test_ppt_preflight_uses_split_quantity_and_cost_total_provenance(report):
    from app.domain.calculations import calculate_destination, calculate_total
    from app.domain.models import Destination
    from app.reporting.pptx_report import validate_ppt_report

    rules = tuple(
        Destination(
            index,
            row.label,
            index,
            True,
            True,
            None,
            index != 12,
            index != 11,
            False,
        )
        for index, row in enumerate(report.rows[:12], start=1)
    )
    current = {row.calculation.destination_id: row.calculation for row in report.rows[:12]}
    following = {
        row.calculation.destination_id: row.calculation
        for row in report.next_month.rows
    }
    current_total = calculate_total(current, rules)
    next_total = calculate_total(following, rules)
    total_quantity_history = {
        **report.total.quantity_by_month,
        "2026-08": current_total.actual_quantity,
    }
    total_cost_history = {
        **report.total.cost_won_by_month,
        "2026-08": current_total.actual_cost_won,
    }
    split_report = replace(
        report,
        total=replace(
            report.total,
            calculation=current_total,
            quantity_by_month=total_quantity_history,
            cost_won_by_month=total_cost_history,
        ),
        next_month=replace(
            report.next_month,
            total=replace(
                report.next_month.total,
                calculation=next_total,
                quantity_by_month=total_quantity_history,
                cost_won_by_month=total_cost_history,
            ),
        ),
        charts=replace(
            report.charts,
            planned_quantity_by_month={
                **report.charts.planned_quantity_by_month,
                "2026-08": current_total.planned_quantity,
            },
            actual_quantity_by_month={
                **report.charts.actual_quantity_by_month,
                "2026-08": current_total.actual_quantity,
            },
            planned_cost_won_by_month={
                **report.charts.planned_cost_won_by_month,
                "2026-08": current_total.planned_cost_won,
            },
            actual_cost_won_by_month={
                **report.charts.actual_cost_won_by_month,
                "2026-08": current_total.actual_cost_won,
            },
        ),
    )
    validate_ppt_report(split_report)

    first = split_report.next_month.rows[0]
    changed = replace(
        first,
        calculation=calculate_destination(
            first.calculation.planned_quantity,
            first.calculation.planned_cost_won + 1,
            None,
            None,
            destination_id=first.calculation.destination_id,
        ),
    )
    invalid = replace(
        split_report,
        next_month=replace(
            split_report.next_month,
            rows=(changed, *split_report.next_month.rows[1:]),
        ),
    )

    with pytest.raises(ValueError, match="next-month total"):
        validate_ppt_report(invalid)


@pytest.mark.parametrize('mutation', ['missing', 'duplicate', 'wrong_kind', 'rows', 'columns', 'missing_cell', 'merge', 'slides', 'picture'])
def test_structure_failures_are_clear_and_atomic(template, report, tmp_path, mutation):
    from app.reporting.pptx_report import generate_pptx, TemplateStructureChanged
    p = Presentation(template)
    shape = named(p, 2, 'report.monthly_table')
    if mutation == 'missing':
        shape.name = 'gone'
    elif mutation == 'duplicate':
        p.slides[1].shapes._spTree.append(deepcopy(shape._element))
    elif mutation == 'wrong_kind':
        shape.name = 'gone'
        _text(p.slides[1], 'report.monthly_table', 'wrong')
    elif mutation == 'rows':
        shape.table._tbl.remove(shape.table._tbl.tr_lst[-1])
    elif mutation == 'columns':
        shape.table._tbl.tblGrid.remove(shape.table._tbl.tblGrid.gridCol_lst[-1])
    elif mutation == 'missing_cell':
        row = shape.table._tbl.tr_lst[2]
        row.remove(row.tc_lst[-1])
    elif mutation == 'merge':
        shape.table.cell(2, 2).merge(shape.table.cell(2, 3))
    elif mutation == 'slides':
        p.slides.add_slide(p.slide_layouts[6])
    else:
        named(p, 3, 'report.chart').name = 'gone'
    broken = tmp_path / 'broken.pptx'
    p.save(broken)
    output = tmp_path / 'existing.pptx'
    output.write_bytes(b'previous successful output')
    with pytest.raises(TemplateStructureChanged, match='TEMPLATE_STRUCTURE_CHANGED'):
        generate_pptx(broken, report, output)
    assert output.read_bytes() == b'previous successful output'
    assert not list(tmp_path.glob('.existing.pptx.*'))


def test_style_empty_run_crop_and_unrelated_picture_preserved(template, report, tmp_path):
    from app.reporting.pptx_report import generate_pptx
    output = tmp_path / 'result.pptx'
    generate_pptx(template, report, output)
    before, after = Presentation(template), Presentation(output)
    for slide in (3, 4, 5):
        old, new = named(before, slide, 'report.chart'), named(after, slide, 'report.chart')
        assert old.image.blob != new.image.blob
        assert [s.name for s in before.slides[slide-1].shapes] == [s.name for s in after.slides[slide-1].shapes]
        assert old.crop_left == new.crop_left and old.crop_right == new.crop_right
        assert named(before, slide, 'unrelated.logo').image.blob == named(after, slide, 'unrelated.logo').image.blob
        assert Image.open(BytesIO(new.image.blob)).size == {3: (2125, 1441), 4: (2060, 1145), 5: (2091, 933)}[slide]
    old = named(before, 2, 'report.monthly_table').table.cell(2, 7)
    new = named(after, 2, 'report.monthly_table').table.cell(2, 7)
    assert old._tc.tcPr.xml == new._tc.tcPr.xml
    assert old.text_frame.paragraphs[0]._p.pPr.xml == new.text_frame.paragraphs[0]._p.pPr.xml
    assert old.text_frame.paragraphs[0].runs[0]._r.rPr.xml == new.text_frame.paragraphs[0].runs[0]._r.rPr.xml
    empty = named(after, 2, 'report.monthly_table').table.cell(2, 6)
    assert empty.text == '2,280'
    assert empty.text_frame.paragraphs[0].runs[0].font.size == Pt(8)


def test_missing_zero_and_selected_reviews(template, report, tmp_path):
    from app.domain.calculations import calculate_destination
    from app.reporting.pptx_report import generate_pptx, ReviewDetail
    first = replace(report.rows[0], calculation=calculate_destination(Decimal(5000), 2000000, Decimal(0), 0, destination_id=1))
    second = replace(report.rows[1], calculation=calculate_destination(Decimal(5000), 2000000, Decimal(4000), 2400000, destination_id=2))
    report = replace(with_actual_rows(report, first, second), reviews=(ReviewDetail('2', 5, 4, Decimal(1200), '가상 원인'),))
    output = tmp_path / 'review.pptx'
    generate_pptx(template, report, output)
    t = named(Presentation(output), 2, 'report.monthly_table').table
    assert t.cell(2, 5).text == '0'
    assert t.cell(2, 7).text == ''
    assert t.cell(15, 6).text == ''
    review = named(Presentation(output), 2, 'report.review_table').table
    assert [review.cell(2, c).text for c in range(8)] == ['가상02', '5', '4', '4,000', '1,000', '1,200', '83.3%', '가상 원인']
    assert all(review.cell(3, c).text == '' for c in range(8))


@pytest.mark.parametrize('case', ['same_path', 'wrong_month', 'missing_plan', 'duplicate_row', 'chart_conflict', 'review_not_selected', 'overflow'])
def test_invalid_report_does_not_write(template, report, tmp_path, case):
    from app.reporting.pptx_report import generate_pptx, ReviewDetail
    output = tmp_path / 'result.pptx'
    if case == 'same_path':
        output = template
    elif case == 'wrong_month':
        report = replace(report, report_month='2026-07')
    elif case == 'missing_plan':
        report = replace(report, next_month=replace(report.next_month, month='2026-08'))
    elif case == 'duplicate_row':
        report = replace(report, rows=(report.rows[0], report.rows[0], *report.rows[2:]))
    elif case == 'chart_conflict':
        report = replace(report, charts=replace(report.charts, actual_cost_won_by_month={**report.charts.actual_cost_won_by_month, '2026-08': 1}))
    elif case == 'review_not_selected':
        report = replace(report, reviews=(ReviewDetail('1', 1, 1, Decimal(1), ''),))
    else:
        report = replace(report, rows=(*report.rows, report.rows[0]))
    before = template.read_bytes()
    with pytest.raises(ValueError):
        generate_pptx(template, report, output)
    assert template.read_bytes() == before
    if output != template:
        assert not output.exists()


def test_verifier_detects_geometry_crop_and_unchanged_logo(template, tmp_path):
    from tools.verify_ppt_layout import verify_ppt_layout
    for change in ('geometry', 'crop', 'logo', 'fill'):
        p = Presentation(template)
        shape = named(p, 3, 'report.chart')
        if change == 'geometry':
            shape.left += 1
        elif change == 'crop':
            shape.crop_top = .3
        elif change == 'logo':
            named(p, 3, 'unrelated.logo').name = 'changed'
        else:
            box = named(p, 3, 'report.quantity_average')
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor(200, 0, 0)
        output = tmp_path / f'{change}.pptx'
        p.save(output)
        assert verify_ppt_layout(template, output)


def test_preparation_rejects_unrecognized_unnamed_template(template, tmp_path):
    from app.reporting.pptx_report import prepare_template, TemplateStructureChanged
    p = Presentation(template)
    named(p, 1, 'report.title').name = 'Text Box 5'
    p.save(template)
    with pytest.raises(TemplateStructureChanged, match='TEMPLATE_STRUCTURE_CHANGED'):
        prepare_template(template, tmp_path / 'normalized.pptx')


def test_preparation_preserves_legacy_master_bytes(template, tmp_path):
    from app.reporting.pptx_report import prepare_template
    blob = BytesIO()
    with ZipFile(template) as old, ZipFile(blob, 'w') as new:
        for part in old.infolist():
            value = old.read(part.filename)
            if part.filename.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/')):
                value = value.replace(b"<?xml version='1.0' encoding='UTF-8' standalone='yes'?>", b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n')
            new.writestr(part, value)
    template.write_bytes(blob.getvalue())
    output = tmp_path / 'prepared.pptx'
    prepare_template(template, output)
    with ZipFile(template) as old, ZipFile(output) as new:
        for name in old.namelist():
            if name.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/theme/')):
                assert old.read(name) == new.read(name)


def test_blank_cells_keep_explicit_paragraph_font_and_remove_old_breaks(template, report, tmp_path):
    from pptx.oxml.xmlchemy import OxmlElement
    from app.reporting.pptx_report import generate_pptx
    p = Presentation(template)
    cell = named(p, 2, 'report.monthly_table').table.cell(14, 1)
    paragraph = cell.text_frame.paragraphs[0]
    paragraph._p.append(OxmlElement('a:br'))
    p.save(template)
    output = tmp_path / 'blank.pptx'
    generate_pptx(template, report, output)
    paragraph = named(Presentation(output), 2, 'report.monthly_table').table.cell(14, 1).text_frame.paragraphs[0]
    assert not paragraph._p.xpath('./a:br')
    end = paragraph._p.xpath('./a:endParaRPr')
    assert end and end[0].get('sz') == '800'


def test_chart_pixels_are_contained_in_original_crop(template, report, tmp_path):
    from app.reporting.pptx_report import generate_pptx
    p = Presentation(template)
    picture = named(p, 3, 'report.chart')
    picture.crop_left, picture.crop_bottom = .03, .12
    p.save(template)
    output = tmp_path / 'cropped.pptx'
    generate_pptx(template, report, output)
    image = Image.open(BytesIO(named(Presentation(output), 3, 'report.chart').image.blob)).convert('RGB')
    # Safe padding below the visible viewport, while the full chart remains inside.
    assert image.crop((0, int(image.height * .89), image.width, image.height)).getextrema() == ((255, 255),) * 3
    assert image.crop((0, 0, int(image.width * .02), image.height)).getextrema() == ((255, 255),) * 3


def test_render_failure_preserves_existing_output(template, report, tmp_path, monkeypatch):
    from app.reporting import pptx_report
    output = tmp_path / 'output.pptx'
    output.write_bytes(b'old output')
    def fail(*args):
        raise RuntimeError('renderer failed')
    monkeypatch.setattr(pptx_report, 'render_cost_chart', fail)
    with pytest.raises(RuntimeError, match='renderer failed'):
        pptx_report.generate_pptx(template, report, output)
    assert output.read_bytes() == b'old output'
    assert not list(tmp_path.glob('.output.pptx.*'))


def test_next_month_averages_include_august_while_report_excludes_it(template, report, tmp_path):
    from app.reporting.pptx_report import generate_pptx
    output = tmp_path / 'averages.pptx'
    generate_pptx(template, report, output)
    p = Presentation(output)
    assert named(p, 2, 'report.monthly_table').table.cell(2, 14).text == '10,000'
    assert named(p, 6, 'report.plan_table').table.cell(2, 4).text == '7,427'


def test_mixed_run_formatting_keeps_visible_run_boundaries(template, report, tmp_path):
    from app.reporting.pptx_report import generate_pptx
    p = Presentation(template)
    paragraph = named(p, 2, 'report.monthly_table').table.cell(2, 2).text_frame.paragraphs[0]
    paragraph.runs[0].text = '12'
    second = paragraph.add_run()
    second.text = '345'
    second.font.size = Pt(8)
    second.font.bold = False
    second.font.color.rgb = RGBColor(255, 0, 0)
    p.save(template)
    output = tmp_path / 'runs.pptx'
    generate_pptx(template, report, output)
    runs = named(Presentation(output), 2, 'report.monthly_table').table.cell(2, 2).text_frame.paragraphs[0].runs
    assert [run.text for run in runs] == ['5,', '000']
    assert runs[0].font.bold is True
    assert runs[1].font.bold is False
    assert runs[1].font.color.rgb == RGBColor(255, 0, 0)


def test_hash_guarded_normalization_names_exact_synthetic_shapes(template, tmp_path, monkeypatch):
    from app.reporting import pptx_report
    p = Presentation(template)
    mapping = {}
    for number, roles in pptx_report._SOURCE_ROLES.items():
        mapping[number] = {}
        for role in roles:
            shape = named(p, number, role)
            shape.name = f'Legacy {shape.shape_id}'
            mapping[number][role] = shape.shape_id, shape.name
    p.save(template)
    before = template.read_bytes()
    monkeypatch.setattr(pptx_report, 'SOURCE_TEMPLATE_SHA256', sha256(before).hexdigest())
    monkeypatch.setattr(pptx_report, '_SOURCE_ROLES', mapping)
    output = tmp_path / 'named.pptx'
    pptx_report.prepare_template(template, output)
    prepared = Presentation(output)
    for number, roles in mapping.items():
        for role in roles:
            assert named(prepared, number, role).name == role
    assert template.read_bytes() == before


def test_postserialization_verification_failure_preserves_output(template, tmp_path, monkeypatch):
    from app.reporting.pptx_report import prepare_template, TemplateStructureChanged
    from tools import verify_ppt_layout
    output = tmp_path / 'previous.pptx'
    output.write_bytes(b'last successful report')
    monkeypatch.setattr(verify_ppt_layout, 'verify_ppt_layout', lambda *args: ['simulated layout failure'])
    with pytest.raises(TemplateStructureChanged, match='simulated layout failure'):
        prepare_template(template, output)
    assert output.read_bytes() == b'last successful report'
    assert not list(tmp_path.glob('.previous.pptx.*'))


def test_empty_review_retains_existing_paragraphs_that_set_native_row_height(template, report, tmp_path):
    from app.reporting.pptx_report import generate_pptx
    p = Presentation(template)
    frame = named(p, 2, 'report.review_table').table.cell(4, 7).text_frame
    second = frame.add_paragraph()
    run = second.add_run()
    run.text = '가상 추가 원인'
    run.font.name, run.font.size = 'Malgun Gothic', Pt(8)
    p.save(template)
    output = tmp_path / 'paragraphs.pptx'
    generate_pptx(template, report, output)
    frame = named(Presentation(output), 2, 'report.review_table').table.cell(4, 7).text_frame
    assert len(frame.paragraphs) == 2
    assert all(not p.text for p in frame.paragraphs)


def test_empty_cell_preserves_inherited_font_without_direct_overrides(template, report, tmp_path):
    from app.reporting.pptx_report import generate_pptx
    p = Presentation(template)
    cell = named(p, 2, 'report.monthly_table').table.cell(2, 6)
    paragraph = cell.text_frame.paragraphs[0]
    for element in paragraph._p.xpath('./a:r | ./a:endParaRPr | ./a:pPr/a:defRPr'):
        element.getparent().remove(element)
    # The synthetic deck's default/minor theme supplies Calibri at 18pt.
    inherited = p._element.xpath('./p:defaultTextStyle/a:lvl1pPr/a:defRPr')[0]
    assert inherited.get('sz') == '1800'
    assert inherited.xpath('./a:latin')[0].get('typeface') == '+mn-lt'
    assert not paragraph._p.xpath('./a:rPr | ./a:pPr/a:defRPr | ./a:endParaRPr')
    paragraph_style, cell_style = paragraph._p.pPr.xml, cell._tc.tcPr.xml
    p.save(template)
    output = tmp_path / 'inherited-font.pptx'
    generate_pptx(template, report, output)
    after = named(Presentation(output), 2, 'report.monthly_table').table.cell(2, 6)
    assert after.text == '2,280'
    assert len(after.text_frame.paragraphs) == 1
    assert not after._tc.xpath('.//a:rPr | .//a:defRPr | .//a:endParaRPr')
    assert after.text_frame.paragraphs[0]._p.pPr.xml == paragraph_style
    assert after._tc.tcPr.xml == cell_style


@pytest.mark.parametrize('reason', [
    'SP3 납품 수량 증가로 적재율 개선\nNQ5 합짐 운영 확대\n배차 조정 및 운행 횟수 감소',
    '납품 수량 증가와 합짐 운영 확대로 적재율을 개선하고 배차 조정 및 운행 횟수를 감소함 ' * 8,
])
def test_review_overflow_rejected_before_mutation(template, report, tmp_path, monkeypatch, reason):
    from app.reporting import pptx_report
    output = tmp_path / 'previous.pptx'
    output.write_bytes(b'previous')
    def unexpected(*args, **kwargs):
        pytest.fail('preflight must finish before any text mutation or rendering')
    monkeypatch.setattr(pptx_report, '_write', unexpected)
    with pytest.raises(ValueError, match='2번 슬라이드.*검토표.*3행.*8열'):
        pptx_report.generate_pptx(template, reviewed_report(report, reason), output)
    assert output.read_bytes() == b'previous'
    assert not list(tmp_path.glob('.previous.pptx.*'))


@pytest.mark.parametrize('case', ['margins', 'font', 'spacing', 'line_height', 'inherited'])
def test_review_fit_uses_destination_metrics(template, report, tmp_path, case, fast_charts):
    from app.reporting.pptx_report import generate_pptx
    p = Presentation(template)
    cell = named(p, 2, 'report.review_table').table.cell(2, 7)
    paragraph = cell.text_frame.paragraphs[0]
    if case == 'margins':
        cell.margin_left = Inches(2.95)
    elif case == 'font':
        paragraph.runs[0].font.size = Pt(30)
    elif case == 'spacing':
        paragraph.space_after = Pt(20)
    elif case == 'line_height':
        paragraph.line_spacing = Pt(30)
    else:
        for element in paragraph._p.xpath('./a:r | ./a:endParaRPr | ./a:pPr/a:defRPr'):
            element.getparent().remove(element)
    p.save(template)
    with pytest.raises(ValueError, match='검토표.*3행.*8열'):
        generate_pptx(template, reviewed_report(report), tmp_path / 'rejected.pptx')
    assert not (tmp_path / 'rejected.pptx').exists()


@pytest.mark.parametrize('measure', ['quantity_by_month', 'cost_won_by_month'])
@pytest.mark.parametrize('location', ['chart', 'next_total', 'next_row', 'current_row', 'historical_row'])
def test_conflicting_actual_lineages_rejected_atomically(template, report, tmp_path, measure, location, fast_charts):
    from app.reporting.pptx_report import generate_pptx
    month = '2026-07' if location in ('chart', 'historical_row') else '2026-08'
    value = Decimal(33333) if measure == 'quantity_by_month' else 20000000
    def changed(row):
        return replace(row, **{measure: {**getattr(row, measure), month: value}})
    if location == 'chart':
        report = replace(report, total=changed(report.total))
    elif location == 'next_total':
        report = replace(report, next_month=replace(report.next_month, total=changed(report.next_month.total)))
    elif location in ('next_row', 'historical_row'):
        report = replace(report, next_month=replace(report.next_month, rows=(changed(report.next_month.rows[0]), *report.next_month.rows[1:])))
    else:
        report = replace(report, rows=(changed(report.rows[0]), *report.rows[1:]))
    output = tmp_path / 'existing.pptx'
    output.write_bytes(b'previous')
    with pytest.raises(ValueError, match='actual history conflict'):
        generate_pptx(template, report, output)
    assert output.read_bytes() == b'previous'


@pytest.mark.parametrize('actual', [None, -1, True, 1.5, '1', 2**63])
def test_required_current_sales_rejected_atomically(template, report, tmp_path, actual, fast_charts):
    from app.reporting.pptx_report import generate_pptx
    output = tmp_path / 'existing.pptx'
    output.write_bytes(b'previous')
    with pytest.raises(ValueError, match='actual sales'):
        generate_pptx(template, replace(report, sales=replace(report.sales, actual_won=actual)), output)
    assert output.read_bytes() == b'previous'


def test_zero_sales_is_present_but_ratio_is_unavailable(template, report, tmp_path, fast_charts):
    from app.reporting.pptx_report import generate_pptx
    output = tmp_path / 'zero-sales.pptx'
    history = {**report.sales.actual_won_by_month, '2026-08': 0}
    report = replace(report, sales=replace(report.sales, actual_won=0, actual_won_by_month=history), next_month=replace(report.next_month, sales=replace(report.next_month.sales, actual_won_by_month=history)))
    generate_pptx(template, report, output)
    table = named(Presentation(output), 2, 'report.monthly_table').table
    assert table.cell(19, 5).text == '0'
    assert table.cell(20, 5).text == ''


@pytest.mark.parametrize('case', ['prior_overlap', 'current_august', 'next_august', 'next_missing_august', 'zero_conflict'])
def test_sales_history_conflict_rejected_before_mutation(template, report, tmp_path, monkeypatch, case):
    from app.reporting import pptx_report
    current = dict(report.sales.actual_won_by_month)
    following = dict(report.next_month.sales.actual_won_by_month)
    actual = report.sales.actual_won
    if case == 'prior_overlap':
        current['2025-01'], following['2025-01'] = 300000, 600000
    elif case == 'current_august':
        current['2026-08'] = 300000000
    elif case == 'next_august':
        following['2026-08'] = 300000000
    elif case == 'next_missing_august':
        following.pop('2026-08')
    else:
        actual, current['2026-08'] = 0, 0
    report = replace(report, sales=replace(report.sales, actual_won=actual, actual_won_by_month=current), next_month=replace(report.next_month, sales=replace(report.next_month.sales, actual_won_by_month=following)))
    output = tmp_path / 'previous.pptx'
    output.write_bytes(b'previous report')
    def unexpected(*args, **kwargs):
        pytest.fail('sales history validation must precede text mutation')
    monkeypatch.setattr(pptx_report, '_write', unexpected)
    with pytest.raises(ValueError, match='actual history conflict.*sales'):
        pptx_report.generate_pptx(template, report, output)
    assert output.read_bytes() == b'previous report'
    assert not list(tmp_path.glob('.previous.pptx.*'))


def test_current_sales_history_may_omit_august_when_next_history_matches(template, report, tmp_path, fast_charts):
    from app.reporting.pptx_report import generate_pptx
    history = dict(report.sales.actual_won_by_month)
    history.pop('2026-08')
    report = replace(report, sales=replace(report.sales, actual_won_by_month=history))
    output = tmp_path / 'sparse-current-sales.pptx'
    generate_pptx(template, report, output)
    table = named(Presentation(output), 6, 'report.plan_table').table
    assert table.cell(17, 4).text == '303,333'


def test_review_one_line_boundary_preserves_layout(template, report, tmp_path, fast_charts):
    from app.reporting.pptx_report import generate_pptx
    from tools.verify_ppt_layout import verify_ppt_layout
    output = tmp_path / 'accepted-review.pptx'
    generate_pptx(template, reviewed_report(report, '납품 수량 증가로 적재율 개선'), output)
    assert verify_ppt_layout(template, output) == []
    assert named(Presentation(output), 2, 'report.review_table').table.cell(2, 7).text == '납품 수량 증가로 적재율 개선'


def test_distinct_shared_history_drives_table_and_both_callouts(template, report, tmp_path, fast_charts):
    from app.reporting.pptx_report import generate_pptx
    quantities = {**report.total.quantity_by_month, '2025-01': Decimal(25000)}
    costs = {**report.total.cost_won_by_month, '2025-01': 20000000}
    report = replace(report,
        total=replace(report.total, quantity_by_month=quantities, cost_won_by_month=costs),
        next_month=replace(report.next_month, total=replace(report.next_month.total, quantity_by_month=quantities, cost_won_by_month=costs)),
        charts=replace(report.charts, actual_quantity_by_month=quantities, actual_cost_won_by_month=costs))
    output = tmp_path / 'shared-history.pptx'
    generate_pptx(template, report, output)
    p = Presentation(output)
    assert named(p, 2, 'report.monthly_table').table.cell(17, 17).text == '10,833'
    for slide in (3, 5):
        assert named(p, slide, 'report.quantity_average').text == '20,417'
    for slide in (4, 5):
        assert named(p, slide, 'report.cost_average').text == '10,833'
